"""离线全链路测试 — 纯构造数据，不依赖网络/邮箱/工单系统/LLM，秒级跑完

覆盖: M2过滤 / M3字段提取(含规则项目提取) / M4标准化拆分 / M5按行比对(串行污染回归) /
      M6 Excel输出 / OCR新旧返回格式兼容
运行: python test_offline.py  (exit 0 = 全部通过)
"""
import os
import sys
import types
import shutil
import json
import tempfile
import asyncio
import time
from datetime import datetime

# 解析失败会抓结果页现场快照；测试统一丢到临时目录，避免污染真实 output/debug
os.environ.setdefault("MAIL_AUDIT_DEBUG_DIR", tempfile.mkdtemp(prefix="mail_audit_debug_"))

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

PASS = 0
FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  [PASS] {name}")
    else:
        FAIL += 1
        print(f"  [FAIL] {name} {detail}")


def test_mail_filter():
    print("== M2 邮件过滤 ==")
    from modules.mail_filter import MailFilter
    filt = MailFilter({"boss@eu-helper.com"})

    mails = [
        {"subject": "乐天 德国WEEE注册", "body_text": "请处理", "sender_email": "a@x.com",
         "attachments": [], "skip_reason": None},
        {"subject": "7月账单", "body_text": "见附件", "sender_email": "b@x.com",
         "attachments": [], "skip_reason": None},
        {"subject": "周报", "body_text": "本周情况", "sender_email": "boss@eu-helper.com",
         "attachments": [], "skip_reason": None},
        {"subject": "闲聊", "body_text": "中午吃什么", "sender_email": "c@x.com",
         "attachments": [], "skip_reason": None},
        {"subject": "回复: 某某", "body_text": "", "sender_email": "me@eu-helper.com",
         "attachments": [], "skip_reason": "self_sent"},
        {"subject": "客户A 德国WEEE证书", "body_text": "证书已下发", "sender_email": "d@x.com",
         "attachments": [], "skip_reason": None},
        {"subject": "客户B 德国WEEE注册已完成", "body_text": "证书请查收", "sender_email": "d2@x.com",
         "attachments": [], "skip_reason": None},
        {"subject": "客户C 德国WEEE证书申请", "body_text": "请申请新注册证书，不是下证通知",
         "sender_email": "d3@x.com", "attachments": [], "skip_reason": None},
        {"subject": "Germany WEEE registration completed", "body_text": "certificate attached",
         "sender_email": "d4@x.com", "attachments": [], "skip_reason": None},
        {"subject": "请处理", "body_text": "", "sender_email": "e2@x.com",
         "attachments": [{"filename": "德国WEEE注册申请表.xlsx", "text_content": ""}],
         "skip_reason": None},
        {"subject": "德国WEEE资料", "body_text": "请处理资料", "sender_email": "e@x.com",
         "attachments": [], "skip_reason": None},
        {"subject": "内部协同", "body_text": "德国WEEE注册", "sender_email": "f@x.com",
         "recipient": "colleague@eu-helper.com", "attachments": [], "skip_reason": None},
    ]
    valid, filtered = filt.filter_mails(mails)

    subj_valid = {m["subject"] for m in valid}
    subj_filtered = {m["filter_reason"] for m in filtered}
    check("业务+注册词 → 有效", "乐天 德国WEEE注册" in subj_valid)
    check("账单 → 过滤", any("无关" in r for r in subj_filtered))
    check("内部无关 → 过滤", any("内部无关" in m["filter_reason"] for m in filtered
                                  if m["subject"] == "周报"))
    check("外部闲聊 → 过滤", any(m["subject"] == "闲聊" for m in filtered))
    check("自身发送 → 过滤", any(m["filter_reason"] == "自身发送" for m in filtered))
    cert = next(m for m in filtered if m["subject"] == "客户A 德国WEEE证书")
    check("证书通知 → 过滤且不调LLM", cert["llm_eligible"] is False)
    completed_cert = next(m for m in filtered if m["subject"] == "客户B 德国WEEE注册已完成")
    check("完成态证书含注册词 → 仍过滤", completed_cert["llm_eligible"] is False)
    request_cert = next(m for m in valid if m["subject"] == "客户C 德国WEEE证书申请")
    check("证书词同时含申请 → 保留新申请", request_cert["filter_status"] == "valid")
    english_completed = next(m for m in filtered if m["subject"] == "Germany WEEE registration completed")
    check("英文完成态证书 → 仍过滤", english_completed["llm_eligible"] is False)
    filename_request = next(m for m in valid if m["subject"] == "请处理")
    check("附件文件名含项目/申请 → 保留待查", filename_request["filter_status"] == "valid")
    business_only = next(m for m in valid if m["subject"] == "德国WEEE资料")
    check("仅业务词 → 保留人工复核", business_only["filter_status"] == "uncertain")
    recipient_internal = next(m for m in filtered if m["subject"] == "内部协同")
    check("内部收件人域 → 直接过滤", recipient_internal["filter_reason"] == "ECOPV系统内部收件人邮箱")
    audit_filter = MailFilter({"audit@eu-helper.com"}, audit_mailbox="audit@eu-helper.com")
    audit_valid, audit_filtered = audit_filter.filter_mails([{
        "subject": "客户德国WEEE注册", "body_text": "请办理", "sender_email": "customer@example.com",
        "recipient": "audit@eu-helper.com", "attachments": [], "skip_reason": None,
    }])
    check("审计邮箱本身不触发内部收件人过滤", len(audit_valid) == 1 and not audit_filtered,
          {"valid": audit_valid, "filtered": audit_filtered})
    alias_valid, alias_filtered = MailFilter(
        set(), audit_mailbox="huiyan.song@eu-helper.com"
    ).filter_mails([{
        "subject": "客户 德国WEEE注册", "body_text": "请办理",
        "sender_email": "customer@example.com",
        "recipient": "report <report@eu-helper.com>",
        "recipient_emails": ["report@eu-helper.com"],
        "attachments": [], "skip_reason": None,
    }])
    check("公共 report 审计地址不误判为内部收件人", len(alias_valid) == 1 and not alias_filtered,
          {"valid": alias_valid, "filtered": alias_filtered})


def test_mail_reader_uses_real_uids():
    """缓存键和 IMAP 调用必须以 UID 而非易变的消息序号为准。"""
    print("== M1 IMAP UID 缓存契约 ==")
    from modules.mail_reader import MailReader

    class FakeConnection:
        def __init__(self):
            self.uid_calls = []
            self.fetches = 0

        def response(self, key):
            return (key, [b"77"])

        def uid(self, command, *args):
            self.uid_calls.append((command, args))
            if command == "SEARCH":
                return "OK", [b"501"]
            if command == "FETCH":
                self.fetches += 1
                raw = (
                    b"From: sender@example.com\r\n"
                    b"To: receiver@example.com\r\n"
                    b"Subject: UID contract\r\n"
                    b"Date: Tue, 01 Sep 2026 10:00:00 +0000\r\n\r\nbody\n"
                    + "谢谢".encode("utf-8")
                )
                return "OK", [(b"501 (RFC822 {1}", raw)]
            raise AssertionError(f"unexpected UID command: {command}")

        def close(self):
            pass

        def logout(self):
            pass

    with tempfile.TemporaryDirectory() as cache_dir:
        conn = FakeConnection()
        reader = MailReader({
            "imap_server": "unused", "imap_port": 993,
            "address": "self@example.com", "password": "unused",
            "cache_dir": cache_dir,
        })
        reader._connect = lambda ctx: conn
        mails = reader.fetch_mails("2026-09-01", "2026-09-01", "self@example.com")
        cache_files = os.listdir(cache_dir)
        check("真实 UID SEARCH/FETCH 调用",
              [call[0] for call in conn.uid_calls] == ["SEARCH", "FETCH"],
              f"got {conn.uid_calls}")
        check("缓存名含 UIDVALIDITY 和 UID",
              len(cache_files) == 1 and "uidv77_501" in cache_files[0],
              f"got {cache_files}")
        check("UID 拉取邮件可解析", len(mails) == 1 and mails[0]["uid"] == "501")
        check("邮件收件人随原始邮件保留", mails[0].get("recipient") == "receiver@example.com", mails)
        check("完整正文与抽取清洗正文分离",
              mails[0].get("body_original") == "body\n谢谢"
              and mails[0].get("body_text") == "body",
              mails[0])


def test_mail_reader_invalid_date_is_tolerated():
    """缺失/损坏 Date 头不能中断整批邮件解析。"""
    print("== M1 异常邮件日期容错 ==")
    from modules.mail_reader import MailReader

    class FakeConnection:
        def response(self, key):
            return key, [b"88"]

        def uid(self, command, *args):
            if command == "SEARCH":
                return "OK", [b"601"]
            if command == "FETCH":
                raw = (
                    b"From: sender@example.com\r\n"
                    b"Subject: missing date header\r\n"
                    b"Date: \r\n\r\nbody"
                )
                return "OK", [(b"601 (RFC822 {1}", raw)]
            raise AssertionError(f"unexpected UID command: {command}")

        def close(self):
            pass

        def logout(self):
            pass

    with tempfile.TemporaryDirectory() as cache_dir:
        reader = MailReader({
            "imap_server": "unused", "imap_port": 993,
            "address": "self@example.com", "password": "unused",
            "cache_dir": cache_dir,
        })
        reader._connect = lambda ctx: FakeConnection()
        mails = reader.fetch_mails("2026-09-01", "2026-09-01", "self@example.com")
        check("异常 Date 不终止阶段一", len(mails) == 1 and mails[0]["date"] is None)


def test_field_extractor():
    print("== M3 字段提取 (无附件四, 规则兜底) ==")
    from modules.field_extractor import FieldExtractor

    agent_map = {
        "agent1@lekotest.com": {"代理": "乐天", "代理简称": "乐天", "收件人邮箱": "agent1@lekotest.com"},
    }
    fx = FieldExtractor(agent_map, project_names=[])  # 附件四为空

    mail = {
        "sender_email": "agent1@lekotest.com",
        "sender_name": "乐天",
        "recipient": "audit@eu-helper.com",
        "date": datetime(2026, 9, 1, 10, 0),
        "subject": "乐天+888+深圳市甲科技有限公司+法国WEEE注册",
        "body_text": "",
        "attachments": [],
    }
    rows = fx.extract_fields(mail)
    r = rows[0]
    check("代理精确匹配", r["代理"] == "乐天" and r["代理匹配方式"] == "邮箱精确匹配", f"got {r['代理']}")
    check("客户从主题提取", r["客户"] == "深圳市甲科技有限公司", f"got {r['客户']}")
    check("项目规则提取(法国WEEE)", r["项目"] == "法国WEEE", f"got {r['项目']}")
    check("需求=注册", r["需求"] == "注册")
    check("置信度high", r["置信度"] == "high", f"got {r['置信度']}")
    check("字段输出保留收件人", r.get("recipient") == "audit@eu-helper.com", r)

    body_customer = fx._extract_customer(
        "主题里的错误公司名：深圳市甲科技有限公司",
        "申请客户：深圳市乙贸易有限公司\n请注册德国WEEE",
        [],
    )
    check(
        "客户名称冲突时正文优先",
        body_customer == {"customer": "深圳市乙贸易有限公司", "source": "正文"},
        f"got {body_customer}",
    )

    # 组合项目规则提取
    mail2 = dict(mail, subject="乐天+深圳市甲科技有限公司+荷兰WEEE+包装法 新增")
    rows2 = fx.extract_fields(mail2)
    projs = {r["项目"] for r in rows2}
    check("共享前缀组合拆分(荷兰WEEE+包装法)", projs == {"荷兰WEEE", "荷兰包装法"}, f"got {projs}")

    # 真实邮箱样本：多个国家以斜杠连接，共享末尾业务类型。
    slash_projects = {
        p["standard_name"]
        for p in fx._extract_projects_by_rules("比利时/波兰/荷兰/丹麦包装法")
    }
    check(
        "多国斜杠共享业务全部展开",
        slash_projects == {"比利时包装法", "波兰包装法", "荷兰包装法", "丹麦包装法"},
        f"got {slash_projects}",
    )

    # 真实主题格式：国家用“+”串起来，业务只写在最后一个国家后面。
    chain_rows = fx.extract_fields(dict(
        mail,
        subject="苏州贝瓦科技有限公司-比利时+波兰+荷兰+丹麦包装法-严立",
    ))
    check(
        "单公司国家链共享末尾业务全部展开",
        {r["项目"] for r in chain_rows}
        == {"比利时包装法", "波兰包装法", "荷兰包装法", "丹麦包装法"},
        f"got {chain_rows}",
    )

    # 真实邮箱样本：正文/主题用顿号列出两家公司，共享一个项目；两家公司
    # 都必须保留，关联不确定时标记人工确认而不是静默丢掉第二家公司。
    pair_text = "永康市稳雅贸易有限公司、义乌市聚洽贸易有限公司申请注册爱尔兰包装法"
    pair_groups = fx._structured_groups(pair_text, "")
    check(
        "顿号分隔的多家公司全部保留",
        {g["customer"] for g in pair_groups}
        == {"永康市稳雅贸易有限公司", "义乌市聚洽贸易有限公司"},
        f"got {pair_groups}",
    )
    check(
        "多家公司共享项目进入人工复核",
        all(g.get("needs_review") for g in pair_groups)
        and all(g["projects"] == ["爱尔兰包装法"] for g in pair_groups),
        f"got {pair_groups}",
    )

    # 真实邮箱样本：& 分隔的“项目+编号+公司”对，项目必须和对应公司配对。
    paired_subject = (
        "德国电池法 BG2002 深圳市甲科技有限公司 & "
        "德国WEEE EG3205 深圳市乙贸易有限公司"
    )
    paired_groups = fx._structured_groups(paired_subject, "")
    check(
        "& 分隔的公司项目按窗口配对",
        {(g["customer"], tuple(g["projects"])) for g in paired_groups}
        == {
            ("深圳市甲科技有限公司", ("德国电池法",)),
            ("深圳市乙贸易有限公司", ("德国WEEE",)),
        },
        f"got {paired_groups}",
    )

    # 真实主题样式：项目/证书类型 + 客户编号 + 英文公司名。
    # WEEE 和 EG/BG 均不可写进客户公司名称，编号应与对应项目保留在同一行。
    entity_fx = FieldExtractor({
        "known@agent.com": {"代理": "美鸥", "代理简称": "美鸥"},
        "giant@agent.com": {"代理": "巨齿鲨", "代理简称": "巨齿鲨"},
        "gudao@agent.com": {"代理": "古道", "代理简称": "古道"},
    }, project_names=[])
    code_rows = entity_fx.extract_fields({
        "sender_email": "de-epr@seamew.de", "sender_name": "",
        "date": datetime(2026, 8, 21, 11, 11),
        "subject": (
            "美鸥跨境——2026新注册 德国WEEE EG3164 RONG FANG TECHNOLOGY LIMITED "
            "+电池BG1945 RONG FANG TECHNOLOGY LIMITED 小型非光伏设备+便携式电池+DESMAX"
        ),
        "body_text": "", "attachments": [],
    })
    check("主题代理别名命中美鸥",
          {r["代理"] for r in code_rows} == {"美鸥"}
          and {r["代理匹配方式"] for r in code_rows} == {"主题代理别名精确匹配"},
          f"got {code_rows}")
    check("证书类型和客户编号不混入英文公司名",
          {r["客户"] for r in code_rows} == {"RONG FANG TECHNOLOGY LIMITED"},
          f"got {code_rows}")
    check("编号和项目一一对应",
          {(r["项目"], r["客户编号"]) for r in code_rows}
          == {("德国WEEE", "EG3164"), ("德国电池法", "BG1945")},
          f"got {code_rows}")

    power_rows = entity_fx.extract_fields({
        "sender_email": "de-epr@seamew.de", "sender_name": "",
        "date": datetime(2026, 8, 21, 11, 12),
        "subject": "美鸥跨境——2026新注册 德国WEEE EG3157 POWERVOLT TECHNOLOGIES LIMITED 小型IT和通讯设备+TAIFU",
        "body_text": "", "attachments": [],
    })
    check("单项目英文公司清洗 EG3157",
          len(power_rows) == 1
          and power_rows[0]["客户"] == "POWERVOLT TECHNOLOGIES LIMITED"
          and power_rows[0]["客户编号"] == "EG3157"
          and power_rows[0]["项目"] == "德国WEEE",
          f"got {power_rows}")

    bracket_rows = entity_fx.extract_fields({
        "sender_email": "unknown@example.com", "sender_name": "",
        "date": datetime(2026, 8, 21, 16, 41),
        "subject": (
            "上海古道+香港美之乐电子商务有限公司+德国WEEE（大型设备+小型设备）"
            "+法国EEE（大型设备+小型设备）+法国包装法（TEMU平台）"
        ),
        "body_text": "", "attachments": [],
    })
    check("括号内加号不切断项目且 EEE 规范为 WEEE",
          {r["项目"] for r in bracket_rows} == {"德国WEEE", "法国WEEE", "法国包装法"},
          f"got {bracket_rows}")
    check("括号项目主题代理别名命中古道",
          {r["代理"] for r in bracket_rows} == {"古道"}, f"got {bracket_rows}")

    # 英文/波兰语法定后缀不能让客户候选退化成申请表说明文字。
    polish_rows = entity_fx.extract_fields({
        "sender_email": "giant@agent.com", "sender_name": "",
        "date": datetime(2026, 8, 21, 15, 54),
        "subject": (
            "三头六臂+HOMELET SPÓŁKA Z OGRANICZONĄ ODPOWIEDZIALNOŚCIĄ+"
            "比利时包装法"
        ),
        "body_text": "",
        "attachments": [{
            "filename": "HOMELET SPÓŁKA Z OGRANICZONĄ ODPOWIEDZIALNOŚCIĄ.xlsx"
        }],
    })
    check(
        "波兰语法定公司后缀正确提取",
        len(polish_rows) == 1
        and "HOMELET" in polish_rows[0]["客户"]
        and "SPÓŁKA" in polish_rows[0]["客户"]
        and "在其他欧盟" not in polish_rows[0]["客户"],
        f"got {polish_rows}",
    )

    # 正文只说“两家公司”、公司名写在两个压缩包名中时，不能只取第一个附件。
    attachment_rows = entity_fx.extract_fields({
        "sender_email": "giant@agent.com", "sender_name": "",
        "date": datetime(2026, 8, 21, 9, 38),
        "subject": "方信-2026.08.20 德国包装法注册",
        "body_text": "你好，附件为两家公司需注册德国包装法资料",
        "attachments": [
            {"filename": "德国包装法注册-甲公司有限公司.zip"},
            {"filename": "德国包装法注册-乙公司有限公司.zip"},
        ],
    })
    check(
        "多个附件公司全部展开",
        {r["客户"] for r in attachment_rows}
        == {"甲公司有限公司", "乙公司有限公司"}
        and {r["项目"] for r in attachment_rows} == {"德国包装法"},
        f"got {attachment_rows}",
    )

    giant_rows = entity_fx.extract_fields({
        "sender_email": "unknown@example.com", "sender_name": "",
        "date": datetime(2026, 8, 21, 16, 42),
        "subject": "巨齿鲨+深圳市辰森科技有限公司+比利时包装法撤单",
        "body_text": "", "attachments": [],
    })
    check("主题中的巨齿鲨直接提取为代理",
          giant_rows[0]["代理"] == "巨齿鲨"
          and giant_rows[0]["代理匹配方式"] == "主题代理别名精确匹配",
          f"got {giant_rows}")

    # 无项目 → 仍输出一行
    mail3 = dict(mail, subject="乐天 通知邮件")
    rows3 = fx.extract_fields(mail3)
    check("无项目输出一行", len(rows3) == 1 and rows3[0]["项目"] == "")

    malformed_customer = fx._sanitize_customer_result({
        "customer": "类目\tRe: 测试\t深圳市样例网络科技有限公司\t德国WEEE",
        "source": "主题",
    })
    check("制表符客户字段只保留公司名",
          malformed_customer["customer"] == "深圳市样例网络科技有限公司",
          f"got {malformed_customer['customer']}")

    check("英文批量客户去编号并支持连字符括号",
          fx._company_substring("BG1892 Sample-Tech (EU) Ltd.") == "Sample-Tech (EU) Ltd",
          f"got {fx._company_substring('BG1892 Sample-Tech (EU) Ltd.')}")

    # 编号逐行列出的批量客户，项目只在主题声明。回归：不能只取第一家公司。
    bulk_mail = dict(
        mail,
        subject="德国电池法申请加入",
        body_text="BG1892 深圳市样例甲科技有限公司\nBG1893 宁波市样例乙贸易有限公司",
    )
    bulk_rows = fx.extract_fields(bulk_mail)
    check("批量编号清单保留全部客户",
          {row["客户"] for row in bulk_rows} == {"深圳市样例甲科技有限公司", "宁波市样例乙贸易有限公司"},
          f"got {[row['客户'] for row in bulk_rows]}")
    check("批量编号清单继承主题项目",
          {row["项目"] for row in bulk_rows} == {"德国电池法"},
          f"got {[row['项目'] for row in bulk_rows]}")

    # 无法用法定后缀识别的品牌/商号不能被静默漏掉；必须保留并标成人工确认。
    nonstandard_bulk = dict(
        mail,
        subject="德国电池法申请加入",
        body_text="BG1892 深圳市样例甲科技有限公司\nBG1893 NorthstarCommerce",
    )
    nonstandard_rows = fx.extract_fields(nonstandard_bulk)
    nonstandard = next(row for row in nonstandard_rows if row["客户"] == "NorthstarCommerce")
    check("非标准英文商号保留而非漏掉",
          len(nonstandard_rows) == 2 and nonstandard["客户提取来源"] == "正文批量清单(名称待确认)",
          f"got {nonstandard_rows}")
    check("非标准英文商号强制人工确认", nonstandard["置信度"] == "需人工确认")

    # 即使结构化分组异常导致没有可展开的项目，也要写一条人工复核行，不能返回空列表。
    fallback_fx = FieldExtractor(agent_map, project_names=[])
    fallback_fx._structured_groups = lambda subject, body: [{
        "customer": "深圳市样例甲科技有限公司", "source": "测试", "projects": []
    }]
    fallback_fx._extract_projects = lambda subject, body, attachments, groups: {
        "projects": [], "epr_form": None, "need_review": False,
        "groups": [{"customer": "深圳市样例甲科技有限公司", "source": "测试", "projects": []}],
    }
    fallback_rows = fallback_fx.extract_fields(mail)
    check("结构化空行降级为人工复核",
          len(fallback_rows) == 1 and fallback_rows[0]["置信度"] == "需人工确认",
          f"got {fallback_rows}")


def test_project_normalizer():
    print("== M4 标准化拆分 ==")
    from modules.project_normalizer import ProjectNormalizer

    table = [
        {"项目编号": "1", "项目名称": "法国WEEE", "国家": "法国", "业务类型": "WEEE"},
        {"项目编号": "2", "项目名称": "德国WEEE", "国家": "德国", "业务类型": "WEEE"},
        {"项目编号": "3", "项目名称": "德国电池法", "国家": "德国", "业务类型": "电池法"},
    ]
    norm = ProjectNormalizer(table)

    r = norm.normalize_and_split("法国WEEE")
    check("标准名直过", len(r) == 1 and r[0]["standard_name"] == "法国WEEE" and r[0]["matched"])
    r = norm.normalize_and_split("德国WEEE+电池法")
    check("组合拆分", [x["standard_name"] for x in r] == ["德国WEEE", "德国电池法"],
          f"got {[x['standard_name'] for x in r]}")
    r = norm.normalize_and_split("电池法", subject="德国客户咨询", body="")
    check("纯业务+上下文补国家", any(x["standard_name"] == "德国电池法" for x in r),
          f"got {r}")
    # 附件四为空时: 原样返回但流程不崩
    norm0 = ProjectNormalizer([])
    r = norm0.normalize_and_split("瑞典WEEE")
    check("空表不崩", len(r) == 1 and r[0]["standard_name"] == "瑞典WEEE")


def test_country_of_project():
    print("== 项目名 → 国家 解析 (国家列不得含业务) ==")
    from modules.project_normalizer import country_of_project

    cases = {
        "奥地利WEEE": "奥地利",
        "德国WEEE": "德国",
        "德国 WEEE": "德国",
        "比利时包装法": "比利时",
        "德国一次性塑料": "德国",
        "罗马尼亚包装法": "罗马尼亚",
        "Sweden WEEE": "瑞典",
        "Ireland WEEE": "爱尔兰",
        "Battery": "",
        "": "",
    }
    for raw, expect in cases.items():
        got = country_of_project(raw)
        check(f"{raw or '(空)'} → {expect or '(空)'}", got == expect, f"got {got!r}")

    # 回归: 国家列绝不能再出现 "国家+业务" 拼在一起的写法
    for raw in ("奥地利WEEE", "德国电池法", "法国包装法", "荷兰一次性塑料"):
        got = country_of_project(raw)
        check(f"国家列不含业务词: {raw}", not any(
            token in got for token in ("WEEE", "电池", "包装", "EPR", "BAT", "塑料")
        ), f"got {got!r}")


def test_workorder_match():
    print("== M5 工单按行比对 (串行污染回归) ==")
    from modules.workorder_checker import WorkOrderChecker
    cfg = {"url": "u", "username": "a", "password": "b"}
    checker = WorkOrderChecker(cfg, None)

    date = datetime(2026, 9, 1, 10, 0)
    row_a = {"date": date, "代理": "乐天", "客户": "深圳市甲科技有限公司", "项目": "德国WEEE"}
    row_b = {"date": date, "代理": "乐天", "客户": "深圳市乙贸易有限公司", "项目": "德国WEEE"}

    wo_a = [{"代理": "乐天", "客户": "深圳市甲科技有限公司", "服务项目": "德国WEEE",
             "创建时间": "2026-09-03 10:00:00"}]
    check(
        "工单表‘下单日期’列可解析",
        checker._parse_wo_date({"下单日期": "2026-09-03 10:00:00"})
        == datetime(2026, 9, 3, 10, 0),
    )

    # 最终核对结果要带平台**原始**下单日期（照抄页面那一列，方便人工逐字对照）
    wo_platform = [{
        "状态": "待下证", "下单日期": "2026-09-03", "注册类型": "首次注册",
        "主工单编号/客户号": "WR117382357001 HW088", "证书号": "",
        "公司名称": "深圳市甲科技有限公司", "所属代理/直客": "乐天",
        "国家": "德国", "服务项目": "德国WEEE", "日期": "2026-09-03",
    }]
    out_p = checker.match_records([dict(row_a)], wo_platform)[0]
    check(
        "已录单行带平台原始下单日期（不格式化）",
        out_p["下单日期"] == "2026-09-03",
        f"got 下单日期={out_p.get('下单日期')!r}",
    )
    # 语义红线：工单日期 = 邮件发来的日期（发件日期），不是平台那列下单日期。
    # 两列同源会让这一列失去意义，人工也无法核对“客户哪天发的邮件”。
    check(
        "工单日期=邮件发来的日期（非平台下单日期）",
        out_p["工单日期"] == datetime(2026, 9, 1, 10, 0)
        and out_p["工单日期"] != datetime(2026, 9, 3),
        f"got 工单日期={out_p.get('工单日期')!r}",
    )

    # A 行与自己的工单 → 已录单
    checker.match_records([dict(row_a)], wo_a)
    out_a = checker.match_records([row_a], wo_a)[0]
    check("本行命中→已录单", out_a["是否已录单"] == "是" and out_a["匹配状态"] == "精确匹配",
          f"got {out_a['匹配状态']}")

    # 返回多条 → 条件不唯一，按约定交人工排查（不再自动收敛成 1 条）
    wo_multi = [
        dict(wo_a[0], 创建时间="2026-09-03 10:00:00"),
        dict(wo_a[0], 创建时间="2026-08-01 10:00:00"),
    ]
    out_multi = checker.match_records([dict(row_a)], wo_multi)[0]
    check(
        "多工单→待复核交人工排查",
        out_multi["是否已录单"] == "待复核"
        and out_multi["匹配状态"] == "多匹配-人工排查"
        and len(out_multi["工单记录"]) == 2
        and "人工排查" in str(out_multi.get("日期筛选说明", "")),
        f"got status={out_multi.get('匹配状态')}, "
        f"records={len(out_multi.get('工单记录', []))}",
    )
    check(
        "多匹配行取『日期最近』候选的下单日期",
        out_multi["下单日期"] == "2026-09-03 10:00:00",
        f"got {out_multi.get('下单日期')!r}",
    )

    # 规则变更（按查询返回条数判定）：返回 1 条即认定已录单，本地字段对不上时
    # 只写提示状态，不再推翻结论——否则本地字段写法差异会制造假漏单。
    # （行隔离由调用方保证：search_batch 每行只传自己的查询结果）
    out_b = checker.match_records([dict(row_b)], wo_a)[0]
    check(
        "返回1条但字段不一致→仍已录单并提示复核",
        out_b["是否已录单"] == "是"
        and out_b["匹配状态"] == "已录单-字段不一致需复核"
        and "建议人工确认" in str(out_b.get("日期筛选说明", "")),
        f"got {out_b['是否已录单']} / {out_b['匹配状态']}",
    )

    # 日期超容差但返回 1 条 → 仍认定已录单，只标注日期异常供复核
    wo_old = [dict(wo_a[0], 创建时间="2026-01-01 10:00:00")]
    out_c = checker.match_records([dict(row_a)], wo_old)[0]
    check(
        "日期超容差→仍认定已录单并标注异常",
        out_c["是否已录单"] == "是"
        and out_c["匹配状态"] == "已录单-日期异常需复核"
        and out_c["工单记录"][0]["日期异常"] is True,
        f"got {out_c['是否已录单']} / {out_c['匹配状态']}",
    )

    # 返回多条且字段全部对不上 → 仍交人工排查，不能落成漏单
    wo_other = [
        {"代理": "别家代理", "客户": "完全无关公司A", "服务项目": "法国包装法",
         "创建时间": "2026-09-03 10:00:00"},
        {"代理": "别家代理", "客户": "完全无关公司B", "服务项目": "法国包装法",
         "创建时间": "2026-09-02 10:00:00"},
    ]
    out_e = checker.match_records([dict(row_a)], wo_other)[0]
    check(
        "多条全不一致→待复核（不落漏单）",
        out_e["是否已录单"] == "待复核"
        and out_e["匹配状态"] == "多匹配-人工排查"
        and "查询条件是否真的生效" in str(out_e.get("日期筛选说明", "")),
        f"got {out_e['是否已录单']} / {out_e.get('日期筛选说明')}",
    )

    # 空查询结果 → 漏单
    out_d = checker.match_records([dict(row_a)], [])[0]
    check("无工单→漏单", out_d["是否已录单"] == "否" and out_d["匹配状态"] == "漏单")
    check("漏单行下单日期为空串（不留 None 脏值）",
          out_d.get("下单日期") == "", f"got {out_d.get('下单日期')!r}")
    check("漏单行工单日期仍=邮件发来的日期",
          out_d.get("工单日期") == datetime(2026, 9, 1, 10, 0),
          f"got {out_d.get('工单日期')!r}")
    check("多匹配行工单日期=邮件发来的日期",
          out_multi.get("工单日期") == datetime(2026, 9, 1, 10, 0),
          f"got {out_multi.get('工单日期')!r}")


def test_workorder_navigation_contract():
    """M5 导航必须走 工单管理 → 注册工单，失败时禁止继续查单。"""
    print("== M5 工单两级导航契约 ==")
    from modules.workorder_checker import DEFAULT_SELECTORS, WorkOrderChecker

    check("导航选择器含父菜单", "workorder_management" in DEFAULT_SELECTORS)
    check("不使用全局注册工单选择器", "register_order" not in DEFAULT_SELECTORS)
    check("移除宽泛 nav 入口", "nav" not in DEFAULT_SELECTORS)
    check("公司控件使用页面精确占位符",
          'input[placeholder="工单编号/客户号/公司名称"]' in DEFAULT_SELECTORS["company_input"])
    check("代理/国家/服务项目使用独立控件选择器",
          all(key in DEFAULT_SELECTORS for key in ("agent_input", "country_input", "service_item_input")))
    check("下拉兼容 Ant Design 容器占位文本",
          all(any(".ant-select:has-text" in sel for sel in DEFAULT_SELECTORS[key])
              for key in ("agent_input", "country_input", "service_item_input")))

    class FakePage:
        def __init__(self):
            self.goto_urls = []
            self.screenshots = []

        async def goto(self, url, wait_until):
            self.goto_urls.append((url, wait_until))

        async def screenshot(self, path, full_page):
            self.screenshots.append((path, full_page))

    async def standard_navigation():
        checker = WorkOrderChecker({"url": "https://example.test/workbench/main", "username": "a", "password": "b"})
        checker._page = FakePage()
        state = {"parent_open": False, "on_register_page": False}
        clicked = []

        async def is_registration_page():
            return state["on_register_page"]

        async def find_registration_under_parent():
            if not state["parent_open"]:
                return None, ""
            return object(), "测试定位"

        async def click_item(text, selector_key):
            clicked.append(text)
            if text == "工单管理":
                state["parent_open"] = True
            return True

        async def click_registration_under_parent():
            clicked.append("注册工单")
            state["on_register_page"] = True
            return True

        async def wait_for_registration_page():
            return state["on_register_page"]

        checker._is_registration_order_page = is_registration_page
        checker._find_visible_registration_order_under_parent = find_registration_under_parent
        checker._click_navigation_item = click_item
        checker._click_registration_order_under_parent = click_registration_under_parent
        checker._wait_for_registration_order_page = wait_for_registration_page
        return await checker._navigate_to_order_list(), clicked

    ok, clicked = asyncio.run(standard_navigation())
    check("按父菜单再子菜单导航", ok and clicked == ["工单管理", "注册工单"], f"got {clicked}")

    async def navigation_failure_stops_query():
        checker = WorkOrderChecker({"url": "https://example.test/workbench/main", "username": "a", "password": "b"})
        checker._logged_in = True
        search_called = False

        async def navigation_fails():
            return False

        async def search_should_not_run(**kwargs):
            nonlocal search_called
            search_called = True
            return []

        checker._navigate_to_order_list = navigation_fails
        checker.search_one = search_should_not_run
        result = await checker.search_batch([
            {"客户": "测试公司", "代理": "代理甲", "项目": "德国WEEE"}
        ])
        return result, search_called

    result, search_called = asyncio.run(navigation_failure_stops_query())
    check("导航失败不查询其他页面", result == [] and not search_called)

    clean, note = WorkOrderChecker._sanitize_customer_for_query(
        "类目\tRe: 测试\t深圳市样例网络科技有限公司\t德国WEEE"
    )
    check("阶段二异常客户字段自动修复",
          clean == "深圳市样例网络科技有限公司" and bool(note), f"got {clean}")
    malformed, note = WorkOrderChecker._sanitize_customer_for_query("类目\t德国WEEE")
    check("阶段二无法修复的客户字段拒绝查询", not malformed and bool(note))


def test_workorder_cancellation():
    """停止请求必须取消正在等待的单条 RPA 查询，且不能继续下一行。"""
    print("== M5 工单查询停止响应 ==")
    from modules.workorder_checker import WorkOrderChecker

    async def stop_during_inflight_query():
        stop_state = {"requested": False}
        checker = WorkOrderChecker(
            {"url": "u", "username": "a", "password": "b", "query_interval_seconds": 0},
            cancel_requested=lambda: stop_state["requested"],
        )
        checker._logged_in = True
        queried_companies = []

        async def navigation_ok():
            return True

        async def slow_search(**kwargs):
            queried_companies.append(kwargs["company"])
            await asyncio.sleep(10)
            return []

        checker._navigate_to_order_list = navigation_ok
        checker.search_one = slow_search
        task = asyncio.create_task(checker.search_batch([
            {"客户": "第一家公司", "代理": "代理甲", "项目": "德国WEEE"},
            {"客户": "第二家公司", "代理": "代理甲", "项目": "德国WEEE"},
        ]))
        await asyncio.sleep(0.05)
        start = time.monotonic()
        stop_state["requested"] = True
        result = await asyncio.wait_for(task, timeout=0.8)
        return result, checker.cancelled, queried_companies, time.monotonic() - start

    result, cancelled, queried_companies, elapsed = asyncio.run(stop_during_inflight_query())
    check("停止取消正在运行的RPA操作",
          result == [] and cancelled and queried_companies == ["第一家公司"] and elapsed < 0.8,
          f"result={result}, cancelled={cancelled}, queries={queried_companies}, elapsed={elapsed:.2f}s")

    from gui import WorkerThread
    worker = WorkerThread({}, mode="stage2")
    worker.stop()
    check("GUI停止信号可传入RPA线程", worker._stop and worker.is_stop_requested())


def test_workorder_row_failure_continues():
    """单条下拉值异常只标记当前行，不应中断后续查询。"""
    print("== M5 单条查询失败继续执行 ==")
    from modules.workorder_checker import WorkOrderChecker

    checker = WorkOrderChecker({
        "url": "u", "username": "a", "password": "b",
        "query_interval_seconds": 0,
        "max_consecutive_query_failures": 3,
    })
    checker._logged_in = True

    async def navigation_ok():
        return True

    async def fake_search_one(**kwargs):
        if kwargs["company"] == "坏代理公司":
            checker._last_query_status = "查询条件填写失败"
            return []
        checker._last_query_status = "实时查询"
        return [{"代理": kwargs["agent"], "客户": kwargs["company"], "服务项目": "德国WEEE"}]

    checker._navigate_to_order_list = navigation_ok
    checker.search_one = fake_search_one
    checker.match_records = lambda rows, orders: [
        (row.update({"是否已录单": "是", "匹配状态": "精确匹配"}) or row)
        for row in rows
    ]
    rows = [
        {"客户": "第一家公司", "代理": "代理甲", "项目": "德国WEEE"},
        {"客户": "坏代理公司", "代理": "友安通", "项目": "瑞典包装法"},
        {"客户": "第三家公司", "代理": "代理甲", "项目": "德国WEEE"},
    ]
    result = asyncio.run(checker.search_batch(rows))
    check("单条失败仍返回全部已处理行", len(result) == 3, f"got {len(result)}")
    check("失败行保留查询失败状态",
          result[1].get("RPA查询状态") == "查询条件填写失败")
    check("失败行写入可追踪原因", "已跳过本条" in result[1].get("_query_error", ""))
    check("失败后续行继续查询", result[2].get("匹配状态") == "精确匹配")


def test_workorder_agent_alias():
    """已知代理别名应只影响网页查询值，不改写原始字段。"""
    print("== M5 代理别名规范化 ==")
    from modules.workorder_checker import WorkOrderChecker
    checker = WorkOrderChecker({"url": "u", "username": "a", "password": "b"})
    check("友安通规范化为友岸通",
          checker._canonical_dropdown_value("所属代理", "友安通") == "友岸通")
    check("非代理字段不套用别名",
          checker._canonical_dropdown_value("国家", "友安通") == "友安通")


def test_excel_writer():
    print("== M6 Excel 输出 ==")
    from modules.excel_writer import ExcelWriter
    tmpdir = tempfile.mkdtemp(prefix="mail_audit_test_")
    try:
        w = ExcelWriter(tmpdir)
        rows = [
            {"sender_email": "a@x.com", "date": datetime(2026, 9, 1), "subject": "s",
             "body_text": "b", "代理": "乐天", "代理匹配方式": "邮箱精确匹配",
             "客户": "深圳市甲科技有限公司", "客户提取来源": "主题",
             "项目": "德国WEEE", "项目原始值": "德国WEEE", "需求": "注册",
             "是否已录单": "否", "匹配状态": "漏单", "查询时间戳": datetime.now()},
        ]
        f1 = w.write_missing_report(rows)
        check("漏单清单生成", os.path.exists(f1) and f1.endswith(".xlsx"))
        f2 = w.write_filtered_report([{"subject": "账单", "filter_reason": "外部无关邮件",
                                       "sender_email": "b@x.com", "date": datetime(2026, 9, 1)}])
        check("过滤清单生成", os.path.exists(f2))
        # 内容抽查
        from openpyxl import load_workbook
        wb = load_workbook(f1)
        ws = wb.active
        check("漏单行数=2(表头+1)", ws.max_row == 2, f"got {ws.max_row}")
        hdr = [c.value for c in ws[1]]
        check("客户列正确", ws.cell(row=2, column=hdr.index("客户") + 1).value == "深圳市甲科技有限公司")
        check("漏单清单保留客户编号列", "客户编号" in hdr, f"got {hdr}")
        check("漏单清单含『下单日期』列", "下单日期" in hdr, f"got {hdr}")
        check("下单日期紧跟工单日期(漏单清单)",
              hdr.index("下单日期") == hdr.index("工单日期") + 1,
              f"got {hdr.index('下单日期')} vs {hdr.index('工单日期')}")
        # 平台原始文本原样写入，不做日期格式化
        f3 = w.write_missing_report([
            {"客户": "已录单公司", "是否已录单": "是", "匹配状态": "精确匹配",
             "工单日期": datetime(2026, 9, 3), "下单日期": "2026-09-03"},
        ])
        ws3 = load_workbook(f3).active
        h3 = [c.value for c in ws3[1]]
        check("漏单清单写入平台原始下单日期（不被格式化）",
              ws3.cell(row=2, column=h3.index("下单日期") + 1).value == "2026-09-03",
              f"got {ws3.cell(row=2, column=h3.index('下单日期') + 1).value!r}")
        # 漏单/未查询行没有平台下单日期，但工单日期（=邮件发来的日期）必须有
        f4 = w.write_missing_report([
            {"客户": "漏单公司", "是否已录单": "否", "匹配状态": "漏单",
             "date": datetime(2026, 9, 1, 8, 15)},
        ])
        ws4 = load_workbook(f4).active
        h4 = [c.value for c in ws4[1]]
        check("漏单行工单日期回落到邮件发件日期",
              ws4.cell(row=2, column=h4.index("工单日期") + 1).value == "2026-09-01 08:15:00",
              f"got {ws4.cell(row=2, column=h4.index('工单日期') + 1).value!r}")
        check("漏单行下单日期留空",
              ws4.cell(row=2, column=h4.index("下单日期") + 1).value in ("", None),
              f"got {ws4.cell(row=2, column=h4.index('下单日期') + 1).value!r}")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def test_ocr_compat():
    print("== OCR 返回格式兼容 ==")
    from utils.attachment_parser import _extract_ocr_texts
    # PaddleOCR 3.x: [{'rec_texts': [...], ...}]
    r3 = [{"rec_texts": ["France WEEE 998877", ""], "rec_scores": [0.99, 0.1]}]
    check("3.x字典格式", _extract_ocr_texts(r3) == "France WEEE 998877")
    # PaddleOCR 2.x: [[page], ...] 其中 page = [[box, (text, score)], ...]
    r2 = [[
        [[[0, 0], [1, 1]], ("France WEEE 998877", 0.99)],
        [[[0, 0], [1, 1]], ("second line", 0.9)],
    ]]
    check("2.x列表格式", _extract_ocr_texts(r2) == "France WEEE 998877\nsecond line")
    check("空结果", _extract_ocr_texts([]) == "" and _extract_ocr_texts([None, []]) == "")


def test_archive_safety():
    """压缩包仅释放允许的成员，拒绝路径穿越。"""
    print("== 附件压缩包安全 ==")
    import zipfile
    from utils.attachment_parser import parse_attachment, _repair_legacy_zip_filename

    tmpdir = tempfile.mkdtemp(prefix="archive_safe_")
    try:
        safe_zip = os.path.join(tmpdir, "safe.zip")
        with zipfile.ZipFile(safe_zip, "w") as zf:
            zf.writestr("资料.txt", "德国WEEE注册")
        safe = parse_attachment(safe_zip, "safe.zip")
        check("安全ZIP正常解析", "德国WEEE注册" in safe["text_content"])

        original_name = "营业执照.jpg"
        cp437_mojibake = original_name.encode("gb18030").decode("cp437")
        check("GBK ZIP成员名恢复中文",
              _repair_legacy_zip_filename(cp437_mojibake) == original_name)
        check("正常英文ZIP成员名不改写",
              _repair_legacy_zip_filename("application_form.jpg") == "application_form.jpg")

        unsafe_zip = os.path.join(tmpdir, "unsafe.zip")
        with zipfile.ZipFile(unsafe_zip, "w") as zf:
            zf.writestr("../../escape.txt", "must-not-write")
        unsafe = parse_attachment(unsafe_zip, "unsafe.zip")
        check("路径穿越ZIP被拒绝", unsafe["text_content"] == "")

        # 当前 Windows 已安装 7-Zip 时，应能显式注册给 rarfile，不依赖旧进程 PATH。
        import rarfile
        from utils.attachment_parser import _configure_rar_tool
        tool = _configure_rar_tool(rarfile)
        if os.name == "nt" and os.path.exists(r"C:\Program Files\7-Zip\7z.exe"):
            check("RAR后端: 发现7-Zip", tool.lower().endswith("7z.exe"))
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def test_ocr_fallback():
    """E3: OCR 图片兜底 — 懒加载/触发条件/多项目拆行/来源标记/找不到返回空"""
    import utils.attachment_parser as ap
    from modules.field_extractor import FieldExtractor

    # 1. 常规解析阶段图片不触发 OCR（懒登记）
    tmp_png = os.path.join(tempfile.mkdtemp(), "截图.png")
    with open(tmp_png, "wb") as f:
        f.write(b"\x89PNG fake")
    called = []
    orig_parse_image = ap._parse_image
    ap._parse_image = lambda p: called.append(p) or "SHOULD_NOT_APPEAR"
    try:
        att = ap.parse_attachment(tmp_png, "截图.png")
        check("懒登记: 不OCR", att["text_content"] == "" and att.get("ocr_pending") is True)
        check("懒登记: 保留路径", att.get("filepath") == tmp_png)
        check("懒登记: OCR未被调用", not called)

        # 2. 字段齐全 → 兜底不触发
        fx = FieldExtractor({}, [])
        fx._ocr_image = lambda p: called.append(p) or ""
        mail_full = {
            "subject": "深圳市测试科技有限公司 法国WEEE 新增",
            "body_text": "", "sender_email": "",
            "attachments": [att],
        }
        rows = fx.extract_fields(mail_full)
        check("字段齐全: 不触发OCR", not called)
        # 客户提取自带主题整段行为(既有逻辑), 只验证关键内容存在 + 项目正确
        check("字段齐全: 正常提取", "深圳市测试科技有限公司" in rows[0]["客户"]
              and rows[0]["项目"] == "法国WEEE")

        # 3. 字段缺失 + OCR 命中 → 兜底提取 + 来源标记 + 置信度low
        att2 = ap.parse_attachment(tmp_png, "截图.png")
        fx2 = FieldExtractor({}, [])
        fx2._ocr_image = lambda p: "订单确认 深圳市某某科技有限公司 德国WEEE、德国电池法 请处理"
        mail_missing = {
            "subject": "请处理附件", "body_text": "", "sender_email": "",
            "attachments": [att2],
        }
        rows2 = fx2.extract_fields(mail_missing)
        check("兜底: 拆两行", len(rows2) == 2)
        check("兜底: 客户", rows2[0]["客户"] == "深圳市某某科技有限公司")
        check("兜底: 客户来源标记", rows2[0]["客户提取来源"] == "OCR图片兜底")
        check("兜底: 项目集合", {r["项目"] for r in rows2} == {"德国WEEE", "德国电池法"})
        check("兜底: 置信度low", all(r["置信度"] == "low" for r in rows2))
        check("兜底: ocr_used标记", all(r.get("ocr_used") for r in rows2))
        check("兜底: 不覆盖已有代理", rows2[0]["代理"] == "")

        # 4. OCR 找不到目标信息 → 不填, 保持待确认单行
        att3 = ap.parse_attachment(tmp_png, "截图.png")
        fx3 = FieldExtractor({}, [])
        fx3._ocr_image = lambda p: "广告水印 满减促销 优惠券 ABCDEF"
        rows3 = fx3.extract_fields({
            "subject": "请查收", "body_text": "", "sender_email": "",
            "attachments": [att3],
        })
        check("无信息: 单行空结果", len(rows3) == 1 and rows3[0]["客户"] == ""
              and rows3[0]["项目"] == "")
        check("无信息: 来源待确认", rows3[0]["客户提取来源"] == "待确认")
        check("无信息: ocr_used=False", not rows3[0].get("ocr_used"))

        # 5. 开关关闭 → 永不触发
        att4 = ap.parse_attachment(tmp_png, "截图.png")
        fx4 = FieldExtractor({}, [], ocr_fallback=False)
        fx4._ocr_image = lambda p: called.append(p) or "深圳市某公司 法国WEEE"
        rows4 = fx4.extract_fields({
            "subject": "请处理附件", "body_text": "", "sender_email": "",
            "attachments": [att4],
        })
        check("开关关闭: 不触发OCR", len(called) == 0)
        check("开关关闭: 字段为空", rows4[0]["客户"] == "" and rows4[0]["项目"] == "")
    finally:
        ap._parse_image = orig_parse_image
        shutil.rmtree(os.path.dirname(tmp_png), ignore_errors=True)


def test_stage2_preprocess():
    """阶段二前置校验: 跳过代理空+需人工确认; (公司,项目) 去重"""
    print("== 阶段二前置校验 ==")
    from modules.workorder_checker import WorkOrderChecker

    rows = [
        # 1. 代理单命中 + 置信度高 → 通过
        {"客户": "A公司", "项目": "德国WEEE", "代理": "代理甲", "置信度": "高"},
        # 2. 代理空 + 置信度=需人工确认 → 跳过
        {"客户": "B公司", "项目": "法国包装法", "代理": "", "置信度": "需人工确认"},
        # 3. (公司,项目) 与第1行重复 → 跳过
        {"客户": "A公司", "项目": "德国WEEE", "代理": "代理甲", "置信度": "高"},
        # 4. 代理命中但置信度=需人工确认 — 按规格"且代理空"才跳, 这里代理非空, 应放行
        {"客户": "C公司", "项目": "意大利WEEE", "代理": "代理乙", "置信度": "需人工确认"},
        # 5. 代理空但置信度高 — 不在跳过规则, 放行(由 RPA/人工兜底)
        {"客户": "D公司", "项目": "西班牙电池法", "代理": "", "置信度": "高"},
    ]
    to_query, skipped = WorkOrderChecker.preprocess_rows(rows)
    check("前置校验: 跳过代理空+置信度=需人工确认",
          any("代理空" in (s.get("_skip_reason") or "") for s in skipped))
    check("前置校验: (公司,项目) 重复被合并",
          any("重复" in (s.get("_skip_reason") or "") for s in skipped))
    check("前置校验: 通过3条(1单命中 + 1非空代理需确认 + 1空代理高置信)",
          len(to_query) == 3,
          f"got {len(to_query)}")
    # 通过的行 (公司,项目) 应唯一
    keys = [(r["客户"], r["项目"]) for r in to_query]
    check("前置校验: 通过行 (公司,项目) 无重复", len(set(keys)) == len(keys))


def test_query_cache_roundtrip():
    """阶段二查询缓存: 命中/过期/写盘"""
    print("== 阶段二查询缓存 ==")
    import tempfile, os, json
    from modules.workorder_checker import WorkOrderChecker

    tmpdir = tempfile.mkdtemp(prefix="qa_cache_")
    cache_path = os.path.join(tmpdir, "cache.json")
    try:
        cfg = {
            "url": "u", "username": "a", "password": "b",
            "query_cache_path": cache_path,
            "query_cache_ttl_days": 7,
        }
        c = WorkOrderChecker(cfg)
        check("缓存初始化为空", c.get_cached_result("代A", "公司X", "德国WEEE") is None)
        # 新增模糊候选后，历史“精确查询 0 条”缓存不能阻断本次模糊重查。
        old_key = c._cache_key("代旧", "旧公司", "德国WEEE")
        c._query_cache[old_key] = {
            "orders": [], "timestamp": datetime.now().isoformat(timespec="seconds")
        }
        check("旧版空缓存会失效以允许模糊重查",
              c.get_cached_result("代旧", "旧公司", "德国WEEE") is None
              and old_key not in c._query_cache)
        c.save_to_cache("代A", "公司X", "德国WEEE", [{"col1": "v1", "date": "2026-09-01"}])
        check("写盘后再读命中",
              c.get_cached_result("代A", "公司X", "德国WEEE") is not None)
        check("不同键不命中",
              c.get_cached_result("代B", "公司X", "德国WEEE") is None)
        window_a = (datetime(2026, 8, 1).date(), datetime(2026, 10, 1).date())
        window_b = (datetime(2026, 9, 1).date(), datetime(2026, 11, 1).date())
        c.save_to_cache("代A", "公司X", "德国WEEE", [{"col1": "v2"}], date_window=window_a)
        check("不同邮件日期窗口不复用缓存",
              c.get_cached_result("代A", "公司X", "德国WEEE", date_window=window_b) is None)
        # 文件落盘
        check("缓存文件已写盘", os.path.exists(cache_path))
        with open(cache_path, encoding="utf-8") as f:
            raw = json.load(f)
        check("缓存文件含无日期与日期窗口两条", len(raw) == 2)

        # 重新构造实例, 模拟重启, 命中落盘数据
        c2 = WorkOrderChecker(cfg)
        cached = c2.get_cached_result("代A", "公司X", "德国WEEE")
        check("新实例从文件命中缓存", cached is not None)
        check("缓存内容含原始 orders", len(cached) == 1 and cached[0].get("col1") == "v1")

        # 过期: 用 cache_ttl_days=7 但把内部 timestamp 改成 2000 年 → age.days > 7 → 过期
        c3 = WorkOrderChecker(cfg)  # ttl=7
        for k, v in list(c3._query_cache.items()):
            v["timestamp"] = "2000-01-01T00:00:00"
        c3._save_cache()
        check("过期条目不返回(返回None)",
              c3.get_cached_result("代A", "公司X", "德国WEEE") is None)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def test_cached_query_replays_visible_inputs():
    """缓存命中时仍应重放浏览器输入，不能静默跳过页面。"""
    print("== 阶段二缓存可视化输入 ==")
    import tempfile, os
    from modules.workorder_checker import WorkOrderChecker

    tmpdir = tempfile.mkdtemp(prefix="qa_cached_inputs_")
    try:
        cache_path = os.path.join(tmpdir, "cache.json")
        checker = WorkOrderChecker({
            "url": "u", "username": "a", "password": "b",
            "query_cache_path": cache_path,
            "show_cached_inputs": True,
            "input_step_delay_seconds": 0,
            "cached_input_hold_seconds": 0,
        })
        checker._logged_in = True
        checker.save_to_cache("代理甲", "公司X", "德国_WEEE", [{"id": "cached"}])
        replay = {}

        async def fake_fill_query_form(**kwargs):
            replay.update(kwargs)
            return True

        checker._fill_query_form = fake_fill_query_form
        result = asyncio.run(checker.search_one(
            company="公司X", agent="代理甲", country="德国", service_item="WEEE"
        ))
        check("缓存命中仍展示查询条件",
              result == [{"id": "cached"}] and replay == {
                  "company": "公司X", "agent": "代理甲", "country": "德国",
                  "service_item": "WEEE", "status": "",
              }, f"got replay={replay}, result={result}")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def _run_async(coro):
    """跑异步协程, 不依赖外部事件循环"""
    import asyncio
    return asyncio.get_event_loop().run_until_complete(coro)


def test_stage2_outputs():
    """阶段一二输出文件名稳定 + 时序副本"""
    print("== 输出文件稳定名 + 时序副本 ==")
    tmpdir = tempfile.mkdtemp(prefix="qa_xlsx_")
    try:
        from modules.excel_writer import ExcelWriter
        writer = ExcelWriter(output_dir=tmpdir)

        rows = [
            {"客户": "A公司", "项目": "德国WEEE", "代理": "甲代理", "代理匹配方式": "精确匹配",
             "置信度": "高", "需求": "注册", "sender_email": "a@x.com", "date": None,
             "subject": "德国WEEE注册", "body_text": "请处理", "项目原始值": "德国WEEE",
             "客户提取来源": "邮件标题", "数据来源": "邮件标题",
             "附件名称": "申请表.xlsx", "filter_reason": "注册关键字+业务词", "llm_used": True},
            {"客户": "B公司", "项目": "法国包装法", "代理": "", "代理匹配方式": "未匹配",
             "置信度": "需人工确认", "需求": "新增", "sender_email": "b@x.com", "date": None,
             "subject": "请帮忙", "body_text": "帮忙新增", "项目原始值": "法国包装法",
             "客户提取来源": "邮件正文", "数据来源": "邮件正文"},
        ]
        primary, secondary = writer.write_stage1_outputs(rows, [])
        # 稳定名存在
        check("阶段一待查清单 稳定名存在",
              os.path.exists(os.path.join(tmpdir, "to_workorder_list.xlsx")))
        check("阶段一漏单复查 稳定名存在",
              os.path.exists(os.path.join(tmpdir, "to_review_list.xlsx")))
        # 时序副本存在
        ts_files = [f for f in os.listdir(tmpdir)
                    if f.startswith("to_workorder_list_") and f.endswith(".xlsx")]
        check("阶段一待查清单 时序副本存在", len(ts_files) >= 1)
        # 分区: A公司 高置信 → to_work; B公司 代理空 → to_review
        from openpyxl import load_workbook
        wb_w = load_workbook(os.path.join(tmpdir, "to_workorder_list.xlsx"))
        wb_r = load_workbook(os.path.join(tmpdir, "to_review_list.xlsx"))
        check("待查清单 1 条", wb_w.active.max_row == 2)  # 1 表头 + 1 数据
        check("漏单复查 1 条", wb_r.active.max_row == 2)
        work_headers = [c.value for c in wb_w.active[1]]
        work_values = [c.value for c in wb_w.active[2]]
        check("待查清单含邮件审核字段",
              all(h in work_headers for h in ["邮件正文摘要(最多300字)", "附件名称", "筛选依据", "人工复核提示"]))
        check("待查清单保留邮件摘要和附件",
              "请处理" in work_values and "申请表.xlsx" in work_values)

        # 阶段二输出
        # 工单日期 = 邮件发来的日期；下单日期 = 平台返回的原始文本（两者不同源）
        rows[0]["date"] = datetime(2026, 9, 1, 9, 30)
        rows[0]["是否已录单"] = "是"
        rows[0]["工单日期"] = rows[0]["date"]
        rows[0]["下单日期"] = "2026-09-03"
        rows[0]["匹配状态"] = "精确匹配"
        all_rows = list(rows)
        to_query = [rows[0]]
        skipped = [rows[1]]
        stable, ts = writer.write_workorder_check_result(all_rows, to_query, skipped)
        check("阶段二 稳定名存在", os.path.exists(os.path.join(tmpdir, "workorder_check_result.xlsx")))
        check("阶段二 时序副本存在", os.path.exists(ts))
        wb_s2 = load_workbook(stable)
        check("阶段二 包含原始邮件+查询+跳过行(3行:1表头+2数据)",
              wb_s2.active.max_row == 3)  # 1 表头 + 1 已查 + 1 跳过
        s2_hdr = [c.value for c in wb_s2.active[1]]
        s2_row = [c.value for c in wb_s2.active[2]]
        check("阶段二结果含『下单日期』列", "下单日期" in s2_hdr, f"got {s2_hdr}")
        check("阶段二结果含查询方式审计列",
              all(h in s2_hdr for h in ["查询方式", "模糊查询词"]), f"got {s2_hdr}")
        check("阶段二写入平台原始下单日期（不被格式化）",
              s2_row[s2_hdr.index("下单日期")] == "2026-09-03",
              f"got {s2_row[s2_hdr.index('下单日期')]!r}")
        check("下单日期紧跟工单日期",
              s2_hdr.index("下单日期") == s2_hdr.index("工单日期") + 1,
              f"got {s2_hdr.index('下单日期')} vs {s2_hdr.index('工单日期')}")
        check("阶段二『工单日期』写邮件发来的日期",
              s2_row[s2_hdr.index("工单日期")] == "2026-09-01 09:30:00",
              f"got {s2_row[s2_hdr.index('工单日期')]!r}")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def test_non_registration_filter():
    """非注册业务词表: 收款/水单、报价/续费、信息变更、发票、签署类
    都不属于 注册/新增/撤单，直接过滤且不进 LLM、不进工作台"""
    print("== T12 非注册业务过滤词表 ==")
    from modules.mail_filter import MailFilter

    def run(subject, body=""):
        mail = {
            "subject": subject, "body_text": body, "attachments": [],
            "sender_email": "agent@example.com",
        }
        return MailFilter(set(), None).filter_mails([mail])

    for subject in [
        "Fw: Re: 美鸥线下回收zgys已收款999.38欧元",   # 收款
        "客户已到账水单请查收",                        # 水单/到账
        "华之星-德国WEEE续费通知-6家",                 # 续费+业务词
        "续费工单-待确认",                             # 续费
        "XX代理-报价单-2026.8",                        # 报价
        "广州筑梦佳园跨境电商有限公司-修改公司名称和地址",  # 信息变更
        "德国WEEE注册-请查收发票",                     # 财务类优先于"注册"
        "一次性塑料授权书-待签字6家",                  # 签署类且主题无注册动作
    ]:
        valid, filtered = run(subject)
        check(f"非注册业务被过滤: {subject[:20]}", len(filtered) == 1 and not valid)
        if filtered:
            reason = filtered[0].get("filter_reason", "")
            check(f"  原因标注非注册业务: {subject[:12]}", "非注册业务" in reason)
            check(f"  不交 LLM 二次分类: {subject[:12]}",
                  filtered[0].get("llm_eligible") is False)

    # 授权书只是注册的配套材料 → 不能被误杀
    valid, filtered = run("新增德国WEEE注册（附授权书）")
    check("带注册动作的授权书邮件仍有效", len(valid) == 1 and not filtered)

    # 只扫主题: 正文/签名档里的价格、发票、报价不能误杀
    valid, filtered = run("新增荷兰包装法注册", body="价格与发票事宜见附件报价单")
    check("正文含价格/发票不误杀", len(valid) == 1 and not filtered)


def test_llm_intent_mock():
    """LLM 邮件分类 + 字段提取: mock 验证协议解析, 默认零成本;
    设环境变量 MAIL_AUDIT_REAL_LLM=1 会跑 1 次真实请求冒烟 (代价 ~ 几分钱)"""
    print("== T11 LLM 邮件分类 + 字段提取 (mock 默认, 可选真实冒烟) ==")
    import os as _os
    from modules.llm_intent import LLMIntentClient

    client = LLMIntentClient({
        "api_key": "mock-key", "base_url": "http://mock", "model": "mock",
        "timeout": 5, "max_tokens": 100,
    })

    # 1. mock classify_email — 验证协议解析
    def fake_classify(system, user, purpose, **_kwargs):
        assert purpose == "intent_classification"
        return {
            "is_target": True, "intent": "注册",
            "email_type": "注册类", "confidence": "high",
            "reason": "客户发起WEEE注册",
        }
    client._call_api = fake_classify
    r = client.classify_email("客户咨询WEEE注册", "请尽快处理", [])
    check("classify: 协议字段齐全",
          r and r["is_target"] is True and r["intent"] == "注册"
          and r["email_type"] == "注册类")
    check("classify: reason 透传", r["reason"] == "客户发起WEEE注册")

    # 2. mock extract_fields_llm — 验证协议 + 多项目拆
    def fake_extract(system, user, purpose, **_kwargs):
        assert purpose == "field_extraction"
        return {
            "代理": "乐天", "客户": "深圳市甲科技有限公司",
            "项目": ["德国WEEE", "德国电池法"],
            "记录": [],
            "需求": "注册", "confidence": "high",
        }
    client._call_api = fake_extract
    f = client.extract_fields_llm("乐天+甲科技+WEEE注册", "", [], "agent@leko.com")
    check("extract: 字段映射",
          f["代理"] == "乐天" and f["客户"] == "深圳市甲科技有限公司"
          and f["项目"] == ["德国WEEE", "德国电池法"]
          and f["需求"] == "注册" and f["confidence"] == "high")

    # 3. mock 失败 (返回 None) → 上层走空结果, 不崩
    client._call_api = lambda *a, **k: None
    r2 = client.classify_email("x", "y", [])
    check("classify 失败: 返回 None", r2 is None)
    f2 = client.extract_fields_llm("x", "y", [])
    check("extract 失败: 返回 None", f2 is None)

    # 4. 真实冒烟 (可选)
    if _os.environ.get("MAIL_AUDIT_REAL_LLM") == "1":
        from openai import OpenAI
        import yaml as _yaml
        cfg_path = _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "config.yaml")
        with open(cfg_path, encoding="utf-8") as _f:
            llm_cfg = _yaml.safe_load(_f)["llm"]
        live = LLMIntentClient(llm_cfg)
        live_resp = live.classify_email("WEEE 注册咨询", "请尽快处理", [])
        check("真实冒烟: 邮件分类成功",
              live_resp is not None and live_resp.get("is_target") is True,
              f"got {live_resp}")
    else:
        check("真实冒烟: 默认跳过(设 MAIL_AUDIT_REAL_LLM=1 启用)", True)


def test_llm_audit_metadata():
    """本地 LLM 审计只记录调用元数据，不触网也不落邮件正文。"""
    print("== LLM 调用审计元数据 ==")
    from modules.llm_intent import LLMIntentClient

    class Usage:
        prompt_tokens = 11
        completion_tokens = 7
        total_tokens = 18

    class Message:
        content = '{"is_target": true}'

    class Choice:
        message = Message()

    class Response:
        choices = [Choice()]
        usage = Usage()
        model = "qwen-test"
        _request_id = "req-test-1"

    class Completions:
        @staticmethod
        def create(**_kwargs):
            return Response()

    class Chat:
        completions = Completions()

    class FakeClient:
        chat = Chat()

    tmpdir = tempfile.mkdtemp(prefix="llm_audit_")
    try:
        audit_path = os.path.join(tmpdir, "calls.jsonl")
        client = LLMIntentClient({
            "api_key": "mock-key", "base_url": "https://workspace.example/v1", "model": "qwen-test",
            "audit_log_path": audit_path,
        })
        client.client = FakeClient()
        result = client._call_api("system", "邮件正文绝不能写入审计日志", "test")
        with open(audit_path, encoding="utf-8") as f:
            record = f.readline()
        summary = client.get_usage_summary()
        check("LLM审计: 记录成功元数据", result == {"is_target": True}
              and summary["requests_succeeded"] == 1 and summary["total_tokens"] == 18)
        check("LLM审计: 不含邮件正文或密钥",
              "邮件正文" not in record and "mock-key" not in record and "req-test-1" in record)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def test_llm_fallback_in_field_extractor():
    """T12: 字段提取 E2 LLM 兜底路径 — 缺字段时调用 LLM, 回填+标记 llm_used"""
    print("== T12 字段提取 LLM 兜底路径 ==")
    from modules.field_extractor import FieldExtractor
    from modules.llm_intent import LLMIntentClient

    # 1. LLM 不可用 → 缺字段直接空结果
    fx1 = FieldExtractor({}, [], llm_client=None)
    rows1 = fx1.extract_fields({
        "subject": "请处理", "body_text": "", "sender_email": "u@x.com",
        "attachments": [],
    })
    check("无 LLM 兜底: llm_used=False", not rows1[0].get("llm_used"))
    check("无 LLM 兜底: 字段空", rows1[0]["客户"] == "" and rows1[0]["项目"] == "")

    # 2. LLM 可用 + 字段缺失 → 触发 LLM 补字段
    fake = LLMIntentClient({"api_key": "mock", "base_url": "x", "model": "x"})
    fake.classify_email = lambda *a, **k: None
    fake.extract_fields_llm = lambda *a, **k: {
        "代理": "乐天", "客户": "深圳市甲科技有限公司",
        "项目": ["法国WEEE"], "需求": "注册", "confidence": "high",
    }
    fx2 = FieldExtractor({}, [], llm_client=fake)
    rows2 = fx2.extract_fields({
        "subject": "请处理", "body_text": "", "sender_email": "u@x.com",
        "attachments": [],
    })
    check("LLM 兜底: llm_used=True", rows2[0].get("llm_used") is True)
    check("LLM 兜底: 客户回填", rows2[0]["客户"] == "深圳市甲科技有限公司")
    check("LLM 兜底: 项目回填", rows2[0]["项目"] == "法国WEEE")
    check("LLM 兜底: 代理回填", rows2[0]["代理"] == "乐天")
    check("LLM 兜底: 需求回填", rows2[0]["需求"] == "注册")

    # 3. 字段已由规则提取齐全 → 不调 LLM
    called = []
    fake.extract_fields_llm = lambda *a, **k: called.append(1) or {
        "代理": "", "客户": "", "项目": [], "需求": "", "confidence": "low",
    }
    fx3 = FieldExtractor(
        {"agent1@x.com": {"代理": "乐天", "代理简称": "乐天", "收件人邮箱": "agent1@x.com"}},
        [],
        llm_client=fake,
    )
    rows3 = fx3.extract_fields({
        "subject": "乐天+深圳市甲科技有限公司+德国WEEE注册",
        "body_text": "", "sender_email": "agent1@x.com",
        "attachments": [],
    })
    check("字段齐全: 不调 LLM", not called and len(rows3) > 0)
    check("字段齐全: 规则结果正确", rows3[0]["客户"] == "深圳市甲科技有限公司"
          and rows3[0]["项目"] == "德国WEEE" and rows3[0]["代理"] == "乐天")

    # 3.1 只有代理邮箱表未匹配时，不为代理猜测额外调用 LLM
    called = []
    fake.extract_fields_llm = lambda *a, **k: called.append(1) or {}
    fx_agent_only = FieldExtractor({}, [], llm_client=fake)
    agent_only_rows = fx_agent_only.extract_fields({
        "subject": "深圳市甲科技有限公司 德国WEEE注册", "body_text": "",
        "sender_email": "unknown@x.com", "attachments": [],
    })
    check("仅代理缺失: 不调LLM", not called and agent_only_rows[0]["代理"] == "")

    # 3.2 规则已识别到多家公司但关系不确定时，仍必须调用 LLM 的“记录”协议；
    # 返回的公司—项目关系应展开为多行，而不是继续只用第一家公司。
    record_called = []
    fake.extract_fields_llm = lambda *a, **k: record_called.append(1) or {
        "代理": "", "客户": "深圳市甲科技有限公司",
        "项目": ["德国电池法", "德国WEEE"],
        "记录": [
            {"客户": "深圳市甲科技有限公司", "项目": ["德国电池法"], "confidence": "high"},
            {"客户": "宁波市乙贸易有限公司", "项目": ["德国WEEE"], "confidence": "high"},
        ],
        "需求": "注册", "confidence": "high",
    }
    fx_records = FieldExtractor({}, [], llm_client=fake)
    record_rows = fx_records.extract_fields({
        "subject": "深圳市甲科技有限公司、宁波市乙贸易有限公司申请注册德国电池法和德国WEEE",
        "body_text": "", "sender_email": "u@x.com", "attachments": [],
    })
    check("多公司关系不确定: 调用LLM记录协议", bool(record_called))
    check("多公司关系不确定: 记录展开两家公司",
          {(r["客户"], r["项目"]) for r in record_rows} == {
              ("深圳市甲科技有限公司", "德国电池法"),
              ("宁波市乙贸易有限公司", "德国WEEE"),
          }, f"got {record_rows}")

    # 4. LLM 返回项目是列表 → 拆多行
    fake.extract_fields_llm = lambda *a, **k: {
        "代理": "乐天", "客户": "深圳市甲科技有限公司",
        "项目": ["德国WEEE", "德国电池法"], "需求": "注册", "confidence": "high",
    }
    fx4 = FieldExtractor({}, [], llm_client=fake)
    rows4 = fx4.extract_fields({
        "subject": "请处理", "body_text": "", "sender_email": "u@x.com",
        "attachments": [],
    })
    check("LLM 多项目: 拆 2 行", len(rows4) == 2)
    check("LLM 多项目: 项目集合正确",
          {r["项目"] for r in rows4} == {"德国WEEE", "德国电池法"})
    check("LLM 多项目: 客户相同", len({r["客户"] for r in rows4}) == 1)


def test_epr_form_parser():
    """EPR 申请表复选框解析: 真实样本勾选状态"""
    print("== EPR 申请表勾选解析 ==")
    import os
    from utils.epr_form_parser import parse_epr_form

    # 真实样本: 亚易 GEEKS TECH 邮件里的申请表 (西班牙包装法 + 荷兰电池法 勾选)
    samples = {
        "tests/fixtures/泛欧EPR申请表-荷兰电池法.xlsx": ["西班牙包装法", "荷兰电池法"],
        "tests/fixtures/8国新版.xlsx": None,  # 存在性以现场为准, 不硬编码结果
    }
    if not os.path.exists("tests/fixtures/泛欧EPR申请表-荷兰电池法.xlsx"):
        print("  [SKIP] 真实样本缺失, 跳过勾选解析断言")
        return
    r = parse_epr_form("tests/fixtures/泛欧EPR申请表-荷兰电池法.xlsx")
    check("勾选解析: 识别为勾选式申请表", r is not None)
    if r:
        check("勾选解析: 勾选项目 = 西班牙包装法+荷兰电池法",
              r.get("projects") == samples["tests/fixtures/泛欧EPR申请表-荷兰电池法.xlsx"],
              f"got {r.get('projects')}")
        check("勾选解析: 勾选完整(无转人工项)",
              not r.get("unmatched_countries") and not r.get("orphan_business"))

    # 非勾选式文件 → None (回落文本规则)
    non_form = os.path.join(tempfile.mkdtemp(prefix="epr_nf_"), "t.xlsx")
    from openpyxl import Workbook
    wb = Workbook(); wb.active["A1"] = "WEEE"; wb.save(non_form); wb.close()
    check("勾选解析: 无控件文件返回 None(回落文本规则)",
          parse_epr_form(non_form) is None)


def test_epr_checkbox_priority():
    """字段提取勾选优先: 勾选=唯一权威; 混合附件主题闸门补充; 勾选不完整转人工"""
    print("== EPR 勾选优先提取 ==")
    import os, tempfile
    from modules.field_extractor import FieldExtractor
    from utils.attachment_parser import parse_attachment

    sample = "tests/fixtures/泛欧EPR申请表-荷兰电池法.xlsx"
    if not os.path.exists(sample):
        print("  [SKIP] 真实样本缺失, 跳过勾选优先断言")
        return
    tmp = tempfile.mkdtemp(prefix="epr_pr_")
    import shutil
    dst = os.path.join(tmp, "申请表.xlsx")
    shutil.copy(sample, dst)
    fx = FieldExtractor(agent_email_map={}, project_names=[], logger=None,
                        llm_client=None, ocr_fallback=False)

    # 1) 纯勾选表: 勾了什么就只出什么, 主题里"荷兰电池法"以外的声明不放大
    att = parse_attachment(dst, "申请表.xlsx")
    mail = {"subject": "测试-注册荷兰电池法-某公司", "sender_email": "",
            "body_text": "", "attachments": [att]}
    rows = fx.extract_fields(mail)
    check("勾选优先: 只出勾选的 西班牙包装法+荷兰电池法",
          {r["项目"] for r in rows} == {"西班牙包装法", "荷兰电池法"},
          f"got {[r['项目'] for r in rows]}")

    # 2) 混合附件: 勾选表 + 旧式表(整张网格), 主题闸门补充
    from openpyxl import Workbook
    plain = os.path.join(tmp, "意大利EPR申请表.xlsx")
    wb = Workbook(); ws = wb.active
    ws["A1"] = "意大利WEEE产品信息"; ws["A2"] = "意大利电池法信息"; ws["A3"] = "意大利包装法信息"
    wb.save(plain); wb.close()
    att2 = parse_attachment(plain, "意大利EPR申请表.xlsx")
    mail2 = {"subject": "测试-意大利电池法-某公司", "sender_email": "",
             "body_text": "", "attachments": [att, att2]}
    rows2 = fx.extract_fields(mail2)
    check("混合附件: 主题声明的 意大利电池法 被补充",
          "意大利电池法" in {r["项目"] for r in rows2},
          f"got {[r['项目'] for r in rows2]}")
    check("混合附件: 主题未声明的 意大利WEEE/包装法 不被扫出",
          {"意大利WEEE", "意大利包装法"}.isdisjoint({r["项目"] for r in rows2}),
          f"got {[r['项目'] for r in rows2]}")


def test_attachment_xlsx_record_authority():
    """附件结构化表优先：正文 15 条不能覆盖 xlsx 的 18 个物理记录。"""
    print("== 附件 Excel 记录数优先 ==")
    from openpyxl import Workbook
    from modules.field_extractor import FieldExtractor
    from utils.attachment_parser import parse_attachment

    tmp = tempfile.mkdtemp(prefix="xlsx_record_authority_")
    path = os.path.join(tmp, "菲利迪欧洲EPR明细.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "注册明细"
    ws.append(["序号", "客户名称", "项目"])
    for index in range(1, 19):
        ws.append([index, f"附件客户{index}有限公司", "德国WEEE"])
    wb.save(path)
    wb.close()

    attachment = parse_attachment(path, "菲利迪欧洲EPR明细.xlsx")
    check("xlsx 结构化解析保留 18 条物理行",
          len(attachment.get("structured_records") or []) == 18,
          f"got {len(attachment.get('structured_records') or [])}")
    check("xlsx 结构化行保留附件来源",
          all(record.get("attachment_name") == "菲利迪欧洲EPR明细.xlsx"
              and record.get("sheet_name") == "注册明细"
              for record in attachment.get("structured_records") or []),
          f"got {attachment.get('structured_records')}")

    extractor = FieldExtractor(
        {"registration@example.com": {"代理": "汇通创"}},
        [{"项目名称": "德国WEEE", "国家": "德国", "业务类型": "WEEE"}],
        ocr_fallback=False,
    )
    rows = extractor.extract_fields({
        "sender_email": "registration@example.com",
        "sender_name": "",
        "date": datetime(2026, 9, 16, 9, 0),
        "subject": "汇通创—菲利迪欧洲 EPR 德国WEEE注册",
        # 正文故意只列 15 个主体，模拟用户报出的真实场景。
        "body_text": "正文列出 15 家主体，详见附件表格。",
        "attachments": [attachment],
    })
    check("附件 18 条优先于正文 15 条，输出 18 条",
          len(rows) == 18,
          f"got {len(rows)}")
    check("附件 18 条客户均被保留",
          {row["客户"] for row in rows} == {f"附件客户{i}有限公司" for i in range(1, 19)},
          f"got {[row['客户'] for row in rows]}")
    check("附件记录数与输出数写入可审计字段",
          all(row.get("附件表格记录数") == 18
              and row.get("附件表格输出数") == 18
              and row.get("附件表格数量校验") == "通过"
              and "第" in row.get("附件明细来源", "")
              for row in rows),
          f"got {rows[:1]}")
    import json
    evidence = json.loads(rows[0].get("附件证据") or "[]")
    check("xlsx 附件证据包含工作表预览",
          evidence and evidence[0].get("sheets") and evidence[0]["sheets"][0].get("rows"),
          evidence)
    check("xlsx 附件证据定位当前物理行",
          evidence and evidence[0].get("records")
          and evidence[0]["records"][0].get("row_number") == 2,
          evidence)


def test_daily_register_rows_survive_llm_project_completion():
    """每日登记表的多主体行不能因 LLM 补项目而被压成一条。"""
    print("== 每日登记表多主体与 LLM 项目补全 ==")
    from modules.field_extractor import FieldExtractor

    class StubLLM:
        enabled = True

        @staticmethod
        def extract_fields_llm(*args, **kwargs):
            return {
                "客户": "南宁前莱文化传媒有限公司",
                "项目": ["德国电池法"],
                "需求": "撤单",
                "confidence": "high",
            }

    attachment = {
        "filename": "8.24 德国项目每日登记表――Peng.xlsx",
        "structured_records": [
            {
                "sheet_name": "Sheet1", "row_number": 2,
                "customer": "南宁前莱文化传媒有限公司", "customer_code": "XS0016",
                "business": "设备电池", "record_type": "entity",
            },
            {
                "sheet_name": "Sheet1", "row_number": 3,
                "customer": "南宁市乾元广进商贸有限公司", "customer_code": "XS0017",
                "business": "设备电池", "record_type": "entity",
            },
        ],
        "sheets": [],
    }
    extractor = FieldExtractor(
        {}, [{"项目名称": "德国电池法", "国家": "德国", "业务类型": "电池法"}],
        llm_client=StubLLM(), ocr_fallback=False,
    )
    rows = extractor.extract_fields({
        "subject": "8.24 每日登记表――Peng",
        "body_text": "请查收每日登记表",
        "sender_email": "agent@example.com",
        "attachments": [attachment],
    })
    check("LLM补项目后仍保留两条公司明细",
          len(rows) == 2
          and {row.get("客户") for row in rows} == {
              "南宁前莱文化传媒有限公司", "南宁市乾元广进商贸有限公司"
          }
          and all(row.get("项目") == "德国电池法" for row in rows),
          rows)


def test_multi_project_company_count_and_generic_phrase():
    """一家公司申请多个项目时不把“1家公司/回收公司”当成第二家公司。"""
    print("== 多项目公司数量与通用词校验 ==")
    from modules.field_extractor import FieldExtractor
    fx = FieldExtractor(
        {},
        [{"项目名称": "比利时包装法"}, {"项目名称": "意大利包装法"},
         {"项目名称": "德国WEEE"}, {"项目名称": "波兰包装法"}],
        ocr_fallback=False,
    )
    body = (
        "1家公司 新注册：\n"
        "1.EPR-比利时包装法 —— 包含100欧元\n"
        "2.EPR-意大利包装法 —— 不含回收公司名称：\n"
        "义乌缪斯电子科技有限公司"
    )
    rows = fx.extract_fields({
        "subject": "易迈+EPR申请",
        "body_text": body,
        "sender_email": "agent@example.com",
        "attachments": [],
    })
    check("一家公司两项目输出两条明细",
          len(rows) == 2 and {row["项目"] for row in rows} == {"比利时包装法", "意大利包装法"},
          rows)
    check("一家公司两项目客户保持一致且不含数量说明",
          {row["客户"] for row in rows} == {"义乌缪斯电子科技有限公司"},
          rows)
    check("公司数量说明与回收公司通用词不进入候选",
          fx._company_substrings(body) == ["义乌缪斯电子科技有限公司"],
          fx._company_substrings(body))

    bulk = (
        "TBA+厦门汉印股份有限公司+德国WEEE\n"
        "TBA+深圳盈拓资讯科技有限公司+德国WEEE\n"
        "TBA+MOONWAY GLOBAL RETAIL LIMITED+波兰包装法\n"
        "TBA+NORZN BROTHERS GOODS S.L.+波兰包装法"
    )
    bulk_rows = fx.extract_fields({
        "subject": "TBA+Z6244182",
        "body_text": bulk,
        "sender_email": "agent@example.com",
        "attachments": [],
    })
    check("正文四家公司按行展开且不把编号当公司",
          len(bulk_rows) == 4
          and {row["客户"] for row in bulk_rows} == {
              "厦门汉印股份有限公司", "深圳盈拓资讯科技有限公司",
              "MOONWAY GLOBAL RETAIL LIMITED", "NORZN BROTHERS GOODS S.L",
          },
          [(row["客户"], row["项目"]) for row in bulk_rows])


def test_company_name_cross_source_consensus_and_code_prefix():
    """公司名需跨来源印证；编号前缀和“一家公司”不能进入客户字段。"""
    print("== 公司名跨来源印证与编号前缀清洗 ==")
    from modules.field_extractor import FieldExtractor
    from modules.business_validator import inspect_row

    fx = FieldExtractor(
        {},
        [{"项目名称": "德国WEEE", "国家": "德国", "业务类型": "WEEE"}],
        ocr_fallback=False,
    )
    cleaned = fx._sanitize_customer_result({
        "customer": "SED5632 xxxxxxxx B.V.", "source": "附件表格结构化",
    })
    check("英文公司编号前缀被剥离",
          cleaned.get("customer") == "xxxxxxxx B.V",
          cleaned)
    check("一家公司占位词被拒绝",
          not fx._sanitize_customer_result({"customer": "一家公司", "source": "正文"}).get("customer"))

    attachment = {
        "filename": "SED5632 xxxxxxxx B.V..xlsx",
        "text_content": "公司名称：xxxxxxxx B.V.\n德国WEEE",
        "structured_records": [{
            "customer": "SED5632 xxxxxxxx B.V.",
            "raw_text": "SED5632 xxxxxxxx B.V. | 德国WEEE",
        }],
    }
    consensus = fx._cross_source_company_consensus(
        "SED5632 xxxxxxxx B.V.+德国WEEE",
        "请为 xxxxxxxx B.V. 注册德国WEEE",
        [attachment],
    )
    check("主题正文附件对公司名形成交叉印证",
          consensus.get("customer") == "xxxxxxxx B.V"
          and len(consensus.get("sources") or []) >= 2,
          consensus)
    rows = fx.extract_fields({
        "subject": "SED5632 xxxxxxxx B.V.+德国WEEE",
        "body_text": "请为 xxxxxxxx B.V. 注册德国WEEE",
        "sender_email": "agent@example.com",
        "attachments": [attachment],
    })
    check("输出明细不再使用客户编号作为公司名",
          rows and all(row.get("客户") == "xxxxxxxx B.V" for row in rows),
          rows)
    issues = inspect_row({"客户": "一家公司-德国包装法", "项目": "德国包装法", "需求": "注册"})
    check("业务复检拦截带项目说明的公司数量占位词",
          any(item.get("code") == "COMPANY_PLACEHOLDER" for item in issues),
          issues)

    # 用户反馈样例：K-DED0943 是复合客户编号，不能留在 RETOTALLY SAS 前面。
    retotally = fx.extract_fields({
        "subject": "K-DED0943 RETOTALLY SAS德国电池法注册+申报（2026年度）",
        "body_text": "请查收 RETOTALLY SAS 德国电池法 2026新注册资料",
        "sender_email": "agent@example.com",
        "attachments": [{"filename": "RETOTALLY SAS_资料.zip"}],
    })
    check("K-DED复合编号从 RETOTALLY SAS 公司名中剥离",
          retotally and retotally[0].get("客户") == "RETOTALLY SAS",
          retotally)

    # 个体工商户没有“有限公司”后缀，但标题和附件都明确给出了主体名称，
    # 不能把申请表里的“非中国公司”说明文字当客户。
    sole_proprietors = fx.extract_fields({
        "subject": "向善XS48211--东莞市虎门钦泓阁五金经营部（个体工商户） "
                   "东莞市虎门鑫瑞铜工艺品店（个体工商户）--德国包装法",
        "body_text": "帮忙安排下德国包装法注册，2个主体，包含注册+申报+AR授权代表",
        "sender_email": "agent@example.com",
        "attachments": [
            {"filename": "东莞市虎门鑫瑞铜工艺品店（个体工商户） 德包.xlsx"},
            {"filename": "营业执照 钦泓阁五金经营部（个体工商户）.jpg"},
        ],
    })
    sole_names = {row.get("客户") for row in sole_proprietors}
    check("个体工商户主体按标题/附件展开两家公司",
          sole_names == {"东莞市虎门钦泓阁五金经营部（个体工商户）",
                         "东莞市虎门鑫瑞铜工艺品店（个体工商户）"},
          sole_proprietors)
    check("非中国公司模板说明不进入客户字段",
          all(row.get("客户") != "非中国公司" for row in sole_proprietors),
          sole_proprietors)

    # 主题中没有公司法定后缀时，也只能从客户编号到国家/项目边界之间
    # 取出明确的英文商号；表格说明“不能与其它公司”必须被拒绝。
    loose_subject = "CHZ482 GLOBALCOMMERCE FRANCE比利时包装法新注册"
    check("无公司后缀的主题商号按边界提取",
          fx._extract_company_from_subject(loose_subject) == "GLOBALCOMMERCE FRANCE",
          fx._extract_company_from_subject(loose_subject))
    chz_rows = fx.extract_fields({
        "subject": loose_subject,
        "body_text": "CHZ482 GLOBALCOMMERCE FRANCE\n比利时包装法新注册\n需要出授权书",
        "sender_email": "agent@example.com",
        "attachments": [{
            "filename": "CHZ482 GLOBALCOMMERCE FRANCE比利时包装法.zip",
            "text_content": "公司邮箱不能与其它公司申请EPR的邮箱一致",
        }],
    })
    check("主题英文商号优先于表单说明",
          chz_rows and all(row.get("客户") == "GLOBALCOMMERCE FRANCE" for row in chz_rows),
          chz_rows)
    check("国公司及表单说明值被拒绝",
          not fx._sanitize_customer_result({"customer": "国公司", "source": "附件内容"}).get("customer")
          and not fx._sanitize_customer_result({"customer": "不能与其它公司", "source": "附件内容"}).get("customer"),
          fx._sanitize_customer_result({"customer": "不能与其它公司", "source": "附件内容"}))


def test_epr_form_labels_are_not_customer_records():
    """EPR 申请表的法人/注册资本等标签不能被识别成客户明细。"""
    print("== EPR 表单标签拒绝式客户校验 ==")
    from modules.field_extractor import FieldExtractor
    from utils.attachment_parser import parse_attachment

    fixture = "tests/fixtures/8国新版.xlsx"
    if os.path.exists(fixture):
        attachment = parse_attachment(fixture, "8国新版.xlsx")
        check("EPR 表单说明页不生成结构化客户记录",
              len(attachment.get("structured_records") or []) == 0,
              attachment.get("structured_records"))

    # 非复选框旧模板也可能有“注册公司信息”小节：其中的姓名、证件号、
    # 电话等是同一主体的字段，不能因为旁边出现“服务内容”就按 3 行计数。
    from utils.attachment_parser import _xlsx_structured_records
    form_rows = [
        ["注册公司信息", "服务内容"],
        ["Huiming Wu", "包装法"],
        ["360781199508126624", "包装法"],
        ["8617758087239", "包装法"],
    ]
    check("注册公司信息区不被当成三条业务明细",
          _xlsx_structured_records(form_rows, "注册公司信息") == [])

    # EPR 资料清单/注意事项页也可能出现“公司”字样，但这些行只是
    # 资料说明，不能按客户明细展开，更不能把说明句当成公司名称。
    caution_rows = [
        ["资料列表", "适用国家/业务", "文件示例", "注意事项"],
        ["4 产品图片或说明书", "德国WEEE、爱尔兰WEEE、意大利WEEE。", "", "若不能提供，请客户确认好注册类别。"],
        ["5 公司名称", "德国包装法", "", "必须由正规翻译公司盖章，否则不接单。"],
    ]
    check("EPR资料说明页不生成客户明细",
          _xlsx_structured_records(caution_rows, "EPR注册资料清单") == [])

    registration_rows = [
        ["公司英文名", "城市", "代理", "编号", "公司中文名"],
        ["weitaZhenke (Hangzhou) Technology Co., Ltd", "Hangzhou", "朴诚源", "PCY009", "维塔臻科（杭州）互联网科技有限公司"],
    ]
    registration_records = _xlsx_structured_records(registration_rows, "注册表")
    check("注册表优先使用中文法定公司名",
          len(registration_records) == 1
          and registration_records[0]["customer"] == "维塔臻科（杭州）互联网科技有限公司"
          and registration_records[0]["customer_code"] == "PCY009"
          and registration_records[0]["agent"] == "朴诚源",
          registration_records)

    # 德国项目每日登记表没有“项目/服务”列，而是用“种类”列记录
    # 设备电池、工业电池等业务类型。此前因未识别“种类”，同一附件的
    # XS0016/XS0017 两家公司只有预览、没有结构化记录，最终邮件只生成一条。
    daily_register_rows = [
        ["日期", "客户简称", "编号", "公司中文名", "公司英文名", "种类", "品牌名"],
        ["2026-08-24", "向善", "XS0016", "南宁前莱文化传媒有限公司", "Nanning Qianlai Cultural Media Co., Ltd.", "设备电池", "Nanning Qianlai Cultural Media Co., Ltd."],
        ["2026-08-24", "向善", "XS0017", "南宁市乾元广进商贸有限公司", "Nanning Qianyuan Guangjin Trading Company Limited", "设备电池", "Nanning Qianlai Cultural Media Co., Ltd."],
    ]
    daily_records = _xlsx_structured_records(daily_register_rows, "Sheet1")
    check("每日登记表按两行拆分公司主体",
          len(daily_records) == 2
          and [item["customer"] for item in daily_records] == [
              "南宁前莱文化传媒有限公司", "南宁市乾元广进商贸有限公司"
          ]
          and [item["customer_code"] for item in daily_records] == ["XS0016", "XS0017"]
          and all(item["business"] == "设备电池" for item in daily_records),
          daily_records)

    extractor = FieldExtractor({}, [], ocr_fallback=False)
    rejected = [
        "POA法人职务/Legal positions",
        "公司注册资金/Registration Capital (必填)",
        "签字地点/Place of signature",
        "Name of legal person",
        "Jiangrong Lu",
        "330219196212027219",
    ]
    check("表单标签、联系人和证件号不进入客户字段",
          all(not extractor._accept_llm_customer(value) for value in rejected),
          rejected)
    accepted = ["石狮市百灿电子商务有限公司", "Exsamind GmbH"]
    check("明确公司主体可被字段 Agent 接受",
          all(extractor._accept_llm_customer(value) == value for value in accepted),
          accepted)

    import workbench_server as W
    check("工作台隐藏旧版人名客户脏记录",
          W._is_non_company_customer_value("Huiming Wu"))
    check("工作台隐藏旧版证件号客户脏记录",
          W._is_non_company_customer_value("360781199508126624"))
    check("工作台隐藏表单说明型公司值",
          W._is_non_company_customer_value("不能与其它公司")
          and W._is_non_company_customer_value("国公司"))

    legacy_subject_row = {
        "邮件主题": "CHZ482 GLOBALCOMMERCE FRANCE比利时包装法新注册",
        "客户公司名称": "不能与其它公司",
        "客户提取来源": "附件内容",
    }
    legacy_repaired = W._repair_attachment_fields(legacy_subject_row)
    check("工作台旧记录从主题恢复英文商号",
          legacy_repaired.get("客户公司名称") == "GLOBALCOMMERCE FRANCE"
          and legacy_repaired.get("客户提取来源") == "主题案件编号边界恢复",
          legacy_repaired)

    evidence_row = {
        "客户公司名称": "Hangzhou",
        "客户编号": "PCY009",
        "附件明细来源": "附件表格：注册表.xlsx / 注册表 第2行",
        "附件证据": json.dumps([{
            "filename": "注册表.xlsx",
            "records": [{
                "attachment_name": "注册表.xlsx", "sheet_name": "注册表", "row_number": 2,
                "raw_text": "weitaZhenke (Hangzhou) Technology Co., Ltd | Hangzhou | 朴诚源 | PCY009 | 维塔臻科（杭州）互联网科技有限公司",
            }],
        }], ensure_ascii=False),
    }
    repaired = W._repair_attachment_fields(evidence_row)
    check("工作台从附件证据修复城市公司名和行级代理",
          repaired.get("客户公司名称") == "维塔臻科（杭州）互联网科技有限公司"
          and repaired.get("代理") == "朴诚源",
          repaired)

    # 欧洲注册表中还会出现 B.V./AB 等法定后缀；这些不能被城市名覆盖。
    eu_evidence_row = {
        "客户公司名称": "Koksijde",
        "附件明细来源": "附件表格：注册表.xlsx / 注册表 第17行",
        "附件证据": json.dumps([{
            "filename": "注册表.xlsx",
            "records": [{
                "attachment_name": "注册表.xlsx", "sheet_name": "注册表", "row_number": 17,
                "raw_text": "KOLLEKTIV7680 B.V. | Koksijde | EUC | EUC31 | KOLLEKTIV7680 B.V.",
            }],
        }], ensure_ascii=False),
    }
    eu_repaired = W._repair_attachment_fields(eu_evidence_row)
    check("工作台识别 B.V. 欧洲法定公司名",
          eu_repaired.get("客户公司名称") == "KOLLEKTIV7680 B.V.",
          eu_repaired)

    from modules.business_validator import inspect_row
    check("规则复检识别人名客户",
          any(item.get("code") == "COMPANY_PERSON_NAME"
              for item in inspect_row({"客户": "Huiming Wu", "项目": "瑞典包装法", "需求": "注册"})))
    check("规则复检识别证件号客户",
          any(item.get("code") == "COMPANY_CONTACT_OR_ID"
              for item in inspect_row({"客户": "360781199508126624", "项目": "瑞典包装法", "需求": "注册"})))


def test_stage2_preprocess_xlsx_headers():
    """阶段二去重键兼容 xlsx 表头键名(客户公司名称/标准化项目名称)"""
    print("== 阶段二去重键表头兼容 ==")
    from modules.workorder_checker import WorkOrderChecker

    rows = [
        {"客户公司名称": "A公司", "标准化项目名称": "德国WEEE", "代理": "代理甲", "置信度": "高"},
        {"客户公司名称": "A公司", "标准化项目名称": "德国WEEE", "代理": "代理乙", "置信度": "高"},
        {"客户公司名称": "B公司", "标准化项目名称": "法国包装法", "代理": "代理乙", "置信度": "高"},
        {"客户公司名称": "A公司", "标准化项目名称": "德国WEEE", "代理": "代理甲", "置信度": "高"},
    ]
    to_query, skipped = WorkOrderChecker.preprocess_rows(rows)
    check("表头兼容: 不同代理不能被错误合并",
          len(to_query) == 3, f"got {len(to_query)}")
    check("表头兼容: 重复 1 条被合并",
          any("重复" in (s.get("_skip_reason") or "") for s in skipped))


def test_stage1_semantic_routing_and_placeholder_dedupe():
    """阶段一回归：低置信意图进入语义 Agent，真实公司覆盖占位候选。"""
    print("== 阶段一语义路由与占位行去重 ==")
    from gui import WorkerThread

    row = {
        "filter_status": "uncertain", "代理": "古道", "客户": "香港美之樂電子商務有限公司",
        "项目": "法国包装法", "需求": "注册", "业务规则校验状态": "通过",
    }
    indices, reasons = WorkerThread._semantic_review_indices([row], [7])
    check("低置信意图进入语义Agent", indices == [0] and "意图低置信度" in reasons)

    bad = {"sender_email": "a@x.com", "subject": "测试", "客户": "1家公司",
           "项目": "德国WEEE", "需求": "注册", "代理": "甲"}
    good = dict(bad, 客户="深圳测试科技有限公司")
    kept, _ = WorkerThread._dedupe_placeholder_rows([bad, good], [1, 1])
    check("真实公司覆盖占位公司行", len(kept) == 1 and kept[0]["客户"] == "深圳测试科技有限公司",
          f"got {[r.get('客户') for r in kept]}")


def test_semantic_issue_template_fields():
    """语义复检输出要保留编号、证据、原值、建议值和短建议。"""
    print("== 语义复检问题模板字段 ==")
    from gui import WorkerThread

    class FakeLLM:
        model = "fake-model"

        def validate_extracted_fields_batch(self, items):
            return {
                "0": {
                    "status": "invalid", "confidence": "low",
                    "issues": [{
                        "code": "COMPANY_CODE_MIXED", "field": "company",
                        "reason": "公司字段混入编号", "evidence": "德国WEEE EG3164 RONG FANG",
                        "current_value": "德国WEEE EG3164 RONG FANG",
                        "suggested_value": "RONG FANG TECHNOLOGY LIMITED",
                        "suggestion": "拆分项目、编号和公司名称",
                    }],
                    "suggestions": {}, "reason": "字段需拆分",
                }
            }

    worker = type("W", (), {"_stop": False, "config": {"llm": {}}})()
    rows = [{"代理": "甲", "客户": "德国WEEE EG3164 RONG FANG", "项目": "德国WEEE",
             "需求": "注册", "subject": "测试", "body_text": "正文"}]
    out = WorkerThread._semantic_validate_rows(
        worker, rows, FakeLLM(), __import__("logging").getLogger("template-test"), mail_keys=[0]
    )
    row = out[0]
    check("问题编号写入", row.get("语义问题编号") == "COMPANY_CODE_MIXED")
    check("问题证据写入", "EG3164" in row.get("语义问题证据", ""))
    check("建议值写入", "RONG FANG" in row.get("语义建议值", ""))
    check("主建议保持简短", row.get("语义校验建议") == "拆分项目、编号和公司名称")


def test_workbench_add_delete_project():
    """工作台「新增项目 / 删除项目」：新增入库并自动推国家，删除只做标记、不动物料表，导出留痕"""
    print("== 工作台 新增/删除项目 ==")
    from pathlib import Path
    from openpyxl import Workbook, load_workbook
    import workbench_server as W

    tmp = Path(tempfile.mkdtemp())
    try:
        primary = tmp / "to_workorder_list.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "工单待查"
        ws.append(["发件人邮箱", "发件日期", "邮件主题", "邮件正文摘要(最多300字)", "附件名称",
                   "代理", "客户公司名称", "标准化项目名称", "需求", "置信度", "数据来源"])
        ws.append(["a@x.com", "2026-09-01", "测试主题", "正文", "a.xlsx",
                   "代理甲", "公司甲", "德国WEEE", "新增注册", "高", "工单待查"])
        wb.save(primary)

        table = tmp / "projects.xlsx"
        wb2 = Workbook()
        ws2 = wb2.active
        ws2.append(["项目编号", "项目名称", "国家", "业务类型"])
        ws2.append(["P001", "德国WEEE", "德国", "WEEE"])
        ws2.append(["P002", "瑞典包装法", "瑞典", "包装法"])
        wb2.save(table)

        old_state_path, old_root, old_table = W.REVIEW_STATE, W.APP_ROOT, W.project_table_path
        W.REVIEW_STATE = tmp / "state.json"
        W.APP_ROOT = tmp
        W.project_table_path = lambda: table
        try:
            # 测试会话必须隔离历史人工状态：昨天的“退回复核”不能影响本次页面。
            fresh_store = W.WorkbenchStore(
                primary_path=str(primary),
                review_path=str(tmp / "none.xlsx"),
                test_mode=True,
            )
            fresh_store.filtered_path = tmp / "no_filtered.xlsx"
            fresh_id = fresh_store.snapshot()["mails"][0]["details"][0]["id"]
            W._save_json(W.REVIEW_STATE, {"records": {fresh_id: {"status": "returned"}}})
            fresh_status = fresh_store.snapshot()["mails"][0]["details"][0]["status"]
            check(
                "工作台测试会话不读取历史退回状态",
                fresh_status == "ready" and fresh_store.state_path != W.REVIEW_STATE,
                {"status": fresh_status, "state_path": str(fresh_store.state_path)},
            )
            W._save_json(W.REVIEW_STATE, {})

            store = W.WorkbenchStore(primary_path=str(primary), review_path=str(tmp / "none.xlsx"))
            store.filtered_path = tmp / "no_filtered.xlsx"
            snap = store.snapshot()
            mail = snap["mails"][0]
            check("工作台: 基础快照读取正常", snap["counts"]["details"] == 1, snap["counts"])
            check("工作台: 项目名称表可选项目已带出", len(snap["projects"]) == 2, snap["projects"])

            added = store.action({
                "action": "add_project",
                "mail_id": mail["id"],
                "mail": {"sender": mail["sender"], "date": mail["date"], "subject": mail["subject"],
                         "body": mail["body"], "attachments": mail["attachments"]},
                "fields": {"company": "公司甲", "agent": "代理甲", "program": "瑞典包装法", "request": "新增注册"},
                "reason": "正文另要求瑞典包装法",
            })
            check("工作台: 新增项目返回新明细编号", bool(added.get("record_id")), added)

            snap2 = store.snapshot()
            mail2 = [m for m in snap2["mails"] if m["id"] == mail["id"]][0]
            new_detail = [d for d in mail2["details"] if d["id"] == added["record_id"]]
            check("工作台: 新增明细并入同一封邮件", len(mail2["details"]) == 2 and len(new_detail) == 1,
                  len(mail2["details"]))
            check("工作台: 新增明细来源标记为人工新增",
                  new_detail and new_detail[0]["source"] == "人工新增")
            check("工作台: 新增明细国家自动推导",
                  new_detail and new_detail[0]["fields"]["country"] == "瑞典",
                  new_detail[0]["fields"] if new_detail else None)
            check("工作台: 新增明细可导出字段齐全",
                  new_detail and new_detail[0]["initial_row"].get("标准化项目名称") == "瑞典包装法",
                  new_detail[0]["initial_row"] if new_detail else None)

            dup = False
            try:
                store.action({"action": "add_project", "mail_id": mail["id"], "mail": {},
                              "fields": {"company": "公司甲", "program": "瑞典包装法"}})
            except ValueError:
                dup = True
            check("工作台: 同一公司+项目重复新增被拦截", dup)

            warn = store.action({"action": "add_project", "mail_id": mail["id"], "mail": {},
                                 "fields": {"company": "公司乙", "program": "奥地利EPR"}})
            check("工作台: 非项目表项目给出提示", "不在项目名称表" in (warn.get("warning") or ""), warn)

            deleted = store.action({"action": "delete_project", "record_id": added["record_id"]})
            check("工作台: 删除返回成功", deleted.get("ok") is True, deleted)
            snap3 = store.snapshot()
            mail3 = [m for m in snap3["mails"] if m["id"] == mail["id"]][0]
            check("工作台: 删除后明细从队列消失",
                  all(d["id"] != added["record_id"] for d in mail3["details"]), len(mail3["details"]))
            check("工作台: 删除仅记录标记，原始产物不动",
                  added["record_id"] in W._load_json(W.REVIEW_STATE).get("deleted", {}))
            check("工作台: 被删项目不在导出主表",
                  all(d["id"] != added["record_id"] for m in snap3["mails"] for d in m["details"]))

            out = store.export()
            wb3 = load_workbook(out)
            check("工作台: 导出含已删除明细留痕表", "已删除明细" in wb3.sheetnames, wb3.sheetnames)
            rows = list(wb3["工单待查"].iter_rows(values_only=True))
            check("工作台: 导出主表行数=未删除明细数",
                  len(rows) - 1 == snap3["counts"]["details"], len(rows) - 1)
        finally:
            W.REVIEW_STATE, W.APP_ROOT, W.project_table_path = old_state_path, old_root, old_table
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_workbench_bulk_confirm():
    """批量确认只处理完整邮件，含人工复核明细的邮件必须跳过。"""
    print("== 工作台 批量确认邮件 ==")
    from pathlib import Path
    from openpyxl import Workbook
    import workbench_server as W

    tmp = Path(tempfile.mkdtemp())
    try:
        primary = tmp / "to_workorder_list.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "工单待查"
        ws.append(["发件人邮箱", "发件日期", "邮件主题", "邮件正文摘要(最多300字)", "附件名称",
                   "代理", "客户公司名称", "标准化项目名称", "需求", "置信度"])
        ws.append(["ready@example.com", "2026-09-16 09:00:00", "两家公司德国包装法",
                   "正文", "", "代理甲", "公司甲", "德国包装法", "注册", "high"])
        ws.append(["ready@example.com", "2026-09-16 09:00:00", "两家公司德国包装法",
                   "正文", "", "代理甲", "公司乙", "德国包装法", "注册", "high"])
        ws.append(["review@example.com", "2026-09-16 09:05:00", "待复核邮件",
                   "正文", "", "代理甲", "一家公司", "德国包装法", "注册", "high"])
        wb.save(primary)

        store = W.WorkbenchStore(
            primary_path=str(primary),
            review_path=str(tmp / "none.xlsx"),
            filtered_path=str(tmp / "none_filtered.xlsx"),
            state_path=str(tmp / "state.json"),
        )
        snap = store.snapshot()
        ready_mail = next(m for m in snap["mails"] if m["sender"] == "ready@example.com")
        review_mail = next(m for m in snap["mails"] if m["sender"] == "review@example.com")
        result = store.action({
            "action": "bulk_confirm",
            "mail_ids": [ready_mail["id"], review_mail["id"]],
            "reason": "批量字段确认",
        })
        check("批量确认: 完整邮件全部确认",
              result.get("confirmed_mails") == 1 and result.get("confirmed_details") == 2,
              result)
        check("批量确认: 含复核明细的邮件被跳过",
              result.get("skipped_count") == 1 and "未解决明细" in result["skipped"][0]["reason"],
              result)
        after = store.snapshot()
        ready_after = next(m for m in after["mails"] if m["id"] == ready_mail["id"])
        review_after = next(m for m in after["mails"] if m["id"] == review_mail["id"])
        check("批量确认: 完整邮件明细状态为已确认",
              all(d["status"] == "confirmed" for d in ready_after["details"]),
              [d["status"] for d in ready_after["details"]])
        check("批量确认: 复核邮件仍保持待复核",
              review_after["status"] == "review",
              review_after["status"])
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_workbench_attachment_record_audit():
    """工作台要把附件行数差异明确展示并强制进入复核。"""
    print("== 工作台 附件表格数量校验 ==")
    from pathlib import Path
    from openpyxl import Workbook
    import workbench_server as W

    tmp = Path(tempfile.mkdtemp())
    try:
        primary = tmp / "to_workorder_list.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "工单待查"
        ws.append([
            "发件人邮箱", "发件日期", "邮件主题", "邮件正文摘要(最多300字)", "附件名称",
            "附件明细来源", "附件表格记录数", "附件表格输出数", "附件表格数量校验",
            "代理", "客户公司名称", "标准化项目名称", "需求", "置信度",
        ])
        ws.append([
            "table@example.com", "2026-09-16 10:00:00", "附件 18 家客户", "正文只列出 15 家",
            "客户清单.xlsx", "附件表格：客户清单.xlsx / 明细 第19行", 18, 15, "需人工确认",
            "代理甲", "公司甲", "德国WEEE", "注册", "high",
        ])
        wb.save(primary)

        store = W.WorkbenchStore(
            primary_path=str(primary),
            review_path=str(tmp / "none.xlsx"),
            filtered_path=str(tmp / "none_filtered.xlsx"),
            state_path=str(tmp / "state.json"),
        )
        detail = store.snapshot()["mails"][0]["details"][0]
        check("工作台: 附件数量不一致自动转人工复核", detail["status"] == "review", detail["status"])
        check("工作台: 附件数量审计字段完整", detail["attachment_audit"] == {
            "expected": "18", "output": "15", "status": "需人工确认",
            "source": "附件表格：客户清单.xlsx / 明细 第19行",
        }, detail["attachment_audit"])
        count_evidence = next((item for item in detail["evidence"] if item["field"] == "attachment_record_count"), {})
        check("工作台: 附件数量作为字段证据显示", count_evidence.get("accepted") == "需人工确认", count_evidence)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_workbench_prefers_valid_history_over_stale_rows():
    """同一邮件的新输出全是脏占位行时，工作台回退到数据库合法记录。"""
    print("== 工作台 同邮件旧版脏记录回退 ==")
    from pathlib import Path
    from openpyxl import Workbook
    import workbench_server as W

    tmp = Path(tempfile.mkdtemp())
    original = (W.REVIEW_STATE, W.APP_ROOT)

    def write_rows(path, rows):
        wb = Workbook(); ws = wb.active; ws.title = "工单待查"
        ws.append(["发件人邮箱", "发件日期", "邮件主题", "邮件正文摘要(最多300字)",
                   "附件名称", "代理", "客户公司名称", "标准化项目名称", "需求", "置信度"])
        for row in rows:
            ws.append(row)
        wb.save(path)

    try:
        primary = tmp / "to_workorder_list.xlsx"
        mail_prefix = ["same@example.com", "2026-09-17 08:00:00", "同邮件版本切换", "正文", "", "代理甲"]
        write_rows(primary, [mail_prefix + ["龙岩市新罗区恒新致远贸易有限公司", "德国WEEE", "注册", "high"]])
        W.REVIEW_STATE = tmp / "state.json"
        W.APP_ROOT = tmp
        first = W.WorkbenchStore(primary_path=str(primary), review_path=str(tmp / "none.xlsx"),
                                 filtered_path=str(tmp / "none-filtered.xlsx"))
        first.snapshot()  # 先把合法历史写入 SQLite

        write_rows(primary, [
            mail_prefix + ["若不能提供，请客户确认好注册类别。", "德国WEEE", "注册", "high"],
            mail_prefix + ["注意事项", "意大利WEEE", "注册", "high"],
        ])
        second = W.WorkbenchStore(primary_path=str(primary), review_path=str(tmp / "none.xlsx"),
                                  filtered_path=str(tmp / "none-filtered.xlsx"))
        mail = second.snapshot()["mails"][0]
        check("同邮件脏新版回退合法历史主体",
              len(mail["details"]) == 1
              and mail["details"][0]["fields"]["company"] == "龙岩市新罗区恒新致远贸易有限公司"
              and mail["details"][0]["fields"]["program"] == "德国WEEE",
              mail["details"])
    finally:
        W.REVIEW_STATE, W.APP_ROOT = original
        shutil.rmtree(tmp, ignore_errors=True)


def test_workbench_persistent_history_totals():
    """正式工作台按邮件日期沉淀完成/未完成总表；临时测试状态不得污染它。"""
    print("== 工作台 持久历史总表 ==")
    from pathlib import Path
    from openpyxl import Workbook, load_workbook
    import workbench_server as W

    tmp = Path(tempfile.mkdtemp())
    original = (
        W.REVIEW_STATE, W.HISTORY_DIR, W.HISTORY_STATE, W.HISTORY_OUTPUT,
        W.COMPLETED_HISTORY_XLSX, W.UNFINISHED_HISTORY_XLSX,
    )
    try:
        primary = tmp / "to_workorder_list.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "工单待查"
        ws.append([
            "发件人邮箱", "发件日期", "邮件主题", "邮件正文摘要(最多300字)", "附件名称",
            "代理", "客户公司名称", "标准化项目名称", "需求", "置信度",
        ])
        ws.append(["done@example.com", "2026-09-16 08:10:00", "已完成邮件", "正文", "done.xlsx",
                   "代理甲", "公司甲", "德国WEEE", "注册", "high"])
        ws.append(["open@example.com", "2026-09-17 09:10:00", "未完成邮件", "正文", "open.xlsx",
                   "代理乙", "公司乙", "德国包装法", "注册", "high"])
        wb.save(primary)

        W.REVIEW_STATE = tmp / "storage" / "workbench_review.json"
        W.HISTORY_DIR = tmp / "storage" / "workbench_history"
        W.HISTORY_STATE = W.HISTORY_DIR / "mail_history.json"
        W.HISTORY_OUTPUT = tmp / "output" / "workbench_history"
        W.COMPLETED_HISTORY_XLSX = W.HISTORY_OUTPUT / "邮件处理完成总表.xlsx"
        W.UNFINISHED_HISTORY_XLSX = W.HISTORY_OUTPUT / "邮件未完成总表.xlsx"
        store = W.WorkbenchStore(
            primary_path=str(primary), review_path=str(tmp / "none.xlsx"),
            filtered_path=str(tmp / "none_filtered.xlsx"),
        )
        first = store.snapshot()
        done = next(mail for mail in first["mails"] if mail["sender"] == "done@example.com")
        store.action({
            "action": "confirm", "record_id": done["details"][0]["id"],
            "reason": "字段核对无误",
        })
        snap = store.snapshot()
        history = snap["history"]
        check("工作台: 正式会话启用持久历史", history["enabled"] is True, history)
        check("工作台: 完成与未完成各自归档", len(history["completed"]) == 1 and len(history["unfinished"]) == 1,
              {"completed": len(history["completed"]), "unfinished": len(history["unfinished"])})
        check("工作台: 历史保留人工确认留痕", history["completed"][0]["last_action"] == "人工确认整理完成",
              history["completed"][0])
        check("工作台: 两张历史总表已生成", W.COMPLETED_HISTORY_XLSX.exists() and W.UNFINISHED_HISTORY_XLSX.exists(),
              [str(W.COMPLETED_HISTORY_XLSX), str(W.UNFINISHED_HISTORY_XLSX)])
        complete_book = load_workbook(W.COMPLETED_HISTORY_XLSX, read_only=True, data_only=True)
        incomplete_book = load_workbook(W.UNFINISHED_HISTORY_XLSX, read_only=True, data_only=True)
        complete_rows = list(complete_book.active.iter_rows(values_only=True))
        incomplete_rows = list(incomplete_book.active.iter_rows(values_only=True))
        complete_book.close(); incomplete_book.close()
        check("工作台: 完成总表含邮件日期、收件人与操作留痕列",
              "邮件日期" in complete_rows[0] and "收件人" in complete_rows[0] and "最后操作" in complete_rows[0] and complete_rows[1][1] == "已完成邮件",
              complete_rows[:2])
        check("工作台: 未完成总表按邮件记录输出", incomplete_rows[1][1] == "未完成邮件", incomplete_rows[:2])
    finally:
        (
            W.REVIEW_STATE, W.HISTORY_DIR, W.HISTORY_STATE, W.HISTORY_OUTPUT,
            W.COMPLETED_HISTORY_XLSX, W.UNFINISHED_HISTORY_XLSX,
        ) = original
        shutil.rmtree(tmp, ignore_errors=True)


def test_workbench_filtered_route_round_trip():
    """询单与过滤邮件支持双向转移，原邮件和操作轨迹均不丢失。"""
    print("== 工作台 询单/过滤双向流转 ==")
    from pathlib import Path
    from openpyxl import Workbook
    import workbench_server as W

    tmp = Path(tempfile.mkdtemp())
    try:
        primary = tmp / "to_workorder_list.xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "工单待查"
        ws.append(["发件人邮箱", "发件日期", "邮件主题", "邮件正文摘要(最多300字)", "附件名称",
                   "代理", "客户公司名称", "标准化项目名称", "需求", "置信度"])
        ws.append(["active@example.com", "2026-09-17 10:00:00", "待处理询单", "正文", "a.xlsx",
                   "代理甲", "公司甲", "德国WEEE", "注册", "high"])
        wb.save(primary)
        filtered = tmp / "filtered_mail_record.xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "过滤日志"
        ws.append(["发件人邮箱", "发件日期", "主题", "正文摘要", "附件名称", "过滤原因", "处理时间戳"])
        ws.append(["filter@example.com", "2026-09-17 11:00:00", "原过滤邮件", "证书正文", "f.pdf",
                   "证书/下号通知", "2026-09-17 11:01:00"])
        wb.save(filtered)

        store = W.WorkbenchStore(
            primary_path=str(primary), review_path=str(tmp / "none.xlsx"),
            filtered_path=str(filtered), state_path=str(tmp / "state.json"), test_mode=True,
        )
        initial = store.snapshot()
        active = next(mail for mail in initial["mails"] if mail["sender"] == "active@example.com")
        source_filter = next(mail for mail in initial["filtered_mails"] if mail["sender"] == "filter@example.com")
        store.action({"action": "move_to_filtered", "mail_id": active["id"], "reason": "无需询单"})
        after_filter = store.snapshot()
        check("工作台: 询单可转入过滤页",
              all(mail["id"] != active["id"] for mail in after_filter["mails"]),
              [mail["id"] for mail in after_filter["mails"]])
        check("工作台: 转入过滤保留原邮件证据", any(mail["id"] == active["id"] for mail in after_filter["filtered_mails"]),
              after_filter["filtered_mails"])
        store.action({"action": "move_to_review", "mail_id": active["id"], "reason": "重新核查"})
        after_restore = store.snapshot()
        restored = next((mail for mail in after_restore["mails"] if mail["id"] == active["id"]), None)
        check("工作台: 人工过滤邮件可转回询单复核", restored is not None and restored["details"][0]["fields"]["company"] == "公司甲", restored)
        store.action({"action": "move_to_filtered", "mail_id": active["id"], "reason": "批量恢复回归准备"})
        store.snapshot()
        store.action({"action": "bulk_restore", "mail_ids": [active["id"]], "reason": "批量恢复回归"})
        bulk_restored = store.snapshot()
        check("工作台: 过滤列表支持批量恢复", any(mail["id"] == active["id"] for mail in bulk_restored["mails"]), bulk_restored)
        store.action({"action": "move_to_review", "mail_id": source_filter["id"], "reason": "证书邮件疑似误过滤"})
        filter_restored = store.snapshot()
        check("工作台: 已定论过滤邮件也可人工转入询单复核",
              any(mail["sender"] == "filter@example.com" for mail in filter_restored["mails"]),
              [mail["sender"] for mail in filter_restored["mails"]])
        check("工作台: 转回操作留痕已保存",
              len(W._load_json(store.state_path).get("mail_routes", {}).get(active["id"], {}).get("events", [])) == 4,
              W._load_json(store.state_path))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_workbench_history_records_workorder_outcomes():
    """工单“找到/漏单/查询失败”必须进入按邮件日期归档的完成或未完成总表。"""
    print("== 工作台 工单结论历史归档 ==")
    from pathlib import Path
    from openpyxl import Workbook, load_workbook
    import workbench_server as W

    tmp = Path(tempfile.mkdtemp())
    original = (
        W.REVIEW_STATE, W.HISTORY_DIR, W.HISTORY_STATE, W.HISTORY_OUTPUT,
        W.COMPLETED_HISTORY_XLSX, W.UNFINISHED_HISTORY_XLSX,
    )
    try:
        primary = tmp / "to_workorder_list.xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "工单待查"
        ws.append(["发件人邮箱", "发件日期", "邮件主题", "邮件正文摘要(最多300字)", "附件名称",
                   "代理", "客户公司名称", "标准化项目名称", "需求", "置信度"])
        ws.append(["found@example.com", "2026-09-17 11:00:00", "已找到邮件", "正文", "", "代理甲", "公司甲", "德国WEEE", "注册", "high"])
        ws.append(["missing@example.com", "2026-09-17 12:00:00", "漏单邮件", "正文", "", "代理乙", "公司乙", "德国包装法", "注册", "high"])
        ws.append(["retry@example.com", "2026-09-17 13:00:00", "查询失败邮件", "正文", "", "代理丙", "公司丙", "法国WEEE", "注册", "high"])
        wb.save(primary)
        result = tmp / "workorder_check_result.xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "工单核对"
        ws.append(["发件人邮箱", "发件日期", "邮件主题", "是否已录单", "匹配状态", "RPA查询状态", "查询时间戳"])
        ws.append(["found@example.com", "2026-09-17 11:00:00", "已找到邮件", "是", "精确匹配", "实时查询", "2026-09-17 15:00:00"])
        ws.append(["missing@example.com", "2026-09-17 12:00:00", "漏单邮件", "否", "漏单", "实时查询", "2026-09-17 15:01:00"])
        ws.append(["retry@example.com", "2026-09-17 13:00:00", "查询失败邮件", "否", "待复核", "查询失败", "2026-09-17 15:02:00"])
        wb.save(result)
        W.REVIEW_STATE = tmp / "storage" / "workbench_review.json"
        W.HISTORY_DIR = tmp / "storage" / "workbench_history"
        W.HISTORY_STATE = W.HISTORY_DIR / "mail_history.json"
        W.HISTORY_OUTPUT = tmp / "output" / "workbench_history"
        W.COMPLETED_HISTORY_XLSX = W.HISTORY_OUTPUT / "邮件处理完成总表.xlsx"
        W.UNFINISHED_HISTORY_XLSX = W.HISTORY_OUTPUT / "邮件未完成总表.xlsx"
        store = W.WorkbenchStore(
            primary_path=str(primary), review_path=str(tmp / "none.xlsx"),
            filtered_path=str(tmp / "none_filtered.xlsx"), workorder_result_path=str(result),
        )
        history = store.snapshot()["history"]
        completed = {item["subject"]: item for item in history["completed"]}
        unfinished = {item["subject"]: item for item in history["unfinished"]}
        check("工作台: 找到和漏单均归入完成总表", set(completed) == {"已找到邮件", "漏单邮件"}, completed)
        check("工作台: 查询失败仍在未完成总表", set(unfinished) == {"查询失败邮件"}, unfinished)
        check("工作台: 找到/漏单数量写入历史", completed["已找到邮件"]["workorder_found"] == 1 and completed["漏单邮件"]["workorder_missing"] == 1, completed)
        book = load_workbook(W.COMPLETED_HISTORY_XLSX, read_only=True, data_only=True)
        headers = list(next(book.active.iter_rows(values_only=True)))
        book.close()
        check("工作台: 完成总表含工单找到/未找到列", "已找到数" in headers and "未找到数" in headers, headers)
    finally:
        (
            W.REVIEW_STATE, W.HISTORY_DIR, W.HISTORY_STATE, W.HISTORY_OUTPUT,
            W.COMPLETED_HISTORY_XLSX, W.UNFINISHED_HISTORY_XLSX,
        ) = original
        shutil.rmtree(tmp, ignore_errors=True)


def test_workorder_result_table_parse():
    """结果表解析：表头必须取自 header-wrapper/thead，读不到就显式失败。"""
    print("== M5 结果表解析 (表头/表体分离) ==")
    from modules.workorder_checker import WorkOrderChecker

    class FakePage:
        def __init__(self, payload=None, error=None):
            self.payload = payload if payload is not None else []
            self.error = error

        async def evaluate(self, script):
            if self.error:
                raise RuntimeError(self.error)
            return self.payload

        async def screenshot(self, path, full_page):
            return None

    dump_dir = tempfile.mkdtemp(prefix="mail_audit_dump_")

    def build():
        return WorkOrderChecker({
            "url": "u", "username": "a", "password": "b",
            "query_cache_path": os.path.join(
                tempfile.mkdtemp(prefix="mail_audit_cache_"), "c.json"
            ),
            "show_cached_inputs": False,
            # 解析失败会抓现场快照，别写到真实 output/debug
            "debug_dump_dir": dump_dir,
        })

    async def extract(payload=None, error=None):
        checker = build()
        checker._page = FakePage(payload, error)
        return checker, await checker._extract_table_data()

    base_table = {
        "headers": ["工单编号", "客户", "所属代理", "服务项目", "下单日期"],
        "headerFound": True,
        "rows": [["WO-1", "深圳市甲科技有限公司", "乐天", "德国WEEE", "2026-09-03 10:00:00"]],
        "kind": "el-table",
        "empty": False,
    }

    checker, recs = asyncio.run(extract([base_table]))
    check("结果表按列索引映射成功",
          len(recs) == 1 and recs[0]["客户"] == "深圳市甲科技有限公司"
          and recs[0]["下单日期"] == "2026-09-03 10:00:00", f"got {recs}")
    check("解析成功时不标记失败", checker._last_table_parse_status == "ok")

    # 固定列会让 el-table 多渲染一组表（列更少、行数相同）→ 取列更全的那组
    _, recs = asyncio.run(extract([base_table, {
        "headers": ["工单编号", "客户"], "headerFound": True,
        "rows": [["WO-1", "深圳市甲科技有限公司"]], "kind": "el-table", "empty": False,
    }]))
    check("固定列重复表取列数更全的一组", "所属代理" in recs[0], f"got {list(recs[0])}")

    # 表头全空（旧实现把表体第一行当表头时的典型症状）→ 必须失败
    checker2, recs = asyncio.run(extract([{
        "headers": [], "headerFound": False,
        "rows": [["编辑", "审核", "详情"]], "kind": "el-table", "empty": False,
    }]))
    check("表头读不到时显式失败",
          recs == [] and checker2._last_table_parse_status == "failed",
          f"got {recs} / {checker2._last_table_parse_status}")

    checker3, _ = asyncio.run(extract([]))
    check("页面上无表格时显式失败", checker3._last_table_parse_status == "failed")

    checker4, _ = asyncio.run(extract(None, "boom"))
    check("结果表读取异常时显式失败", checker4._last_table_parse_status == "failed")

    checker5, recs = asyncio.run(extract([{
        "headers": ["工单编号", "客户"], "headerFound": True,
        "rows": [], "kind": "el-table", "empty": True,
    }]))
    check("页面显示暂无数据时按 0 条处理",
          recs == [] and checker5._last_table_parse_status == "ok")

    # 工单平台是 Ant Design（不是 Element UI）：表头在 .ant-table-header，
    # 表体在 .ant-table-body，空结果时表体里只有一个 ant-table-placeholder。
    ant_headers = ["", "状态", "下单日期", "注册类型", "主工单编号/客户号", "证书号",
                   "公司名称", "所属代理/直客", "国家", "服务项目", "备注",
                   "备忘（仅内部使用）", "服务商", "最近交付进度", "驳回原因", "操作", ""]
    _, recs = asyncio.run(extract([{
        "headers": ant_headers, "headerFound": True,
        "rows": [["", "待审核", "2026-09-05 10:00:00", "新注册", "WO-1", "CERT-1",
                  "保文商貿有限公司", "海维", "意大利", "包装法", "", "", "服务商A",
                  "已交付", "", "详情", ""]],
        "kind": "ant-table", "empty": False,
    }]))
    check("ant 结果表：合并列别名归并出 客户/所属代理/项目",
          len(recs) == 1 and recs[0].get("客户") == "保文商貿有限公司"
          and recs[0].get("所属代理") == "海维" and recs[0].get("项目") == "包装法",
          f"got {recs}")
    check("ant 结果表：无表头且整列为空的选择框/滚动条列不产出字段",
          "col_0" not in recs[0] and "col_16" not in recs[0],
          f"got {list(recs[0])}")

    checker_ant, recs = asyncio.run(extract([{
        "headers": ant_headers, "headerFound": True, "rows": [],
        "kind": "ant-table", "empty": True,
    }]))
    check("ant 空结果（暂无数据占位行）按 0 条处理",
          recs == [] and checker_ant._last_table_parse_status == "ok",
          f"got {recs} / {checker_ant._last_table_parse_status}")

    # 读到表头却 0 行、页面又没显示空状态：无法区分“真的没录单”和“表体没读到”，
    # 必须显式失败，绝不能默认成 0 条漏单。
    checker_amb, recs = asyncio.run(extract([{
        "headers": ["工单编号", "客户"], "headerFound": True,
        "rows": [], "kind": "plain", "empty": False,
    }]))
    check("读不到表体且无空状态标记时显式失败",
          recs == [] and checker_amb._last_table_parse_status == "failed",
          f"got {recs} / {checker_amb._last_table_parse_status}")
    check("歧义空结果的原因写进解析备注",
          "未显示空状态" in checker_amb._last_table_parse_note,
          f"got {checker_amb._last_table_parse_note}")

    # 表头必须优先于行数：否则“表体表有行、表头表无行”时行多者胜出，表头就丢了
    checker_hd, recs = asyncio.run(extract([
        {"headers": ["工单编号", "客户", "所属代理"], "headerFound": True, "rows": [],
         "kind": "ant-table", "empty": True},
        {"headers": [], "headerFound": False,
         "rows": [["暂无数据"]], "kind": "ant-table", "empty": False},
    ]))
    check("候选择优时表头优先于行数",
          checker_hd._last_table_parse_status == "ok" and recs == [],
          f"got {recs} / {checker_hd._last_table_parse_status}")

    checker6, recs = asyncio.run(extract([{
        "headers": ["A", "B"], "headerFound": True,
        "rows": [["1", "2"]], "kind": "plain", "empty": False,
    }]))
    check("列名与业务列完全不符时拒绝产出假漏单",
          recs == [] and checker6._last_table_parse_status == "failed",
          f"got {recs} / {checker6._last_table_parse_status}")

    dup = WorkOrderChecker._rows_to_records(["客户", "客户"], [["甲", "乙"]])
    check("重名表头加后缀互不覆盖",
          dup[0]["客户"] == "甲" and dup[0]["客户#2"] == "乙", f"got {dup}")
    flat = WorkOrderChecker._rows_to_records(["工单编号"], [["WO-1", "客户甲"]])
    check("超出表头的列退化为 col_N 且不塌缩成空 key",
          flat[0]["col_1"] == "客户甲" and "" not in flat[0], f"got {flat}")

    # 解析失败必须抓现场快照（否则只能靠日志猜表结构），且同一实例只抓一次
    solo_dir = tempfile.mkdtemp(prefix="mail_audit_dump_solo_")
    solo = WorkOrderChecker({
        "url": "u", "username": "a", "password": "b",
        "query_cache_path": os.path.join(
            tempfile.mkdtemp(prefix="mail_audit_c_"), "c.json"
        ),
        "show_cached_inputs": False,
        "debug_dump_dir": solo_dir,
    })
    solo._page = FakePage([{
        "headers": [], "headerFound": False,
        "rows": [["编辑", "审核"]], "kind": "el-table", "empty": False,
    }])
    asyncio.run(solo._extract_table_data())
    snap_first = sorted(os.listdir(solo_dir))
    asyncio.run(solo._extract_table_data())
    snap_second = sorted(os.listdir(solo_dir))
    check("解析失败时保存结果页现场快照",
          any(n.startswith("table_structure_") for n in snap_first), f"got {snap_first}")
    check("同一实例只抓一次现场快照",
          len(snap_first) == len(snap_second), f"{snap_first} -> {snap_second}")
    check("快照自身失败不影响失败判定（桩无 content 也不抛异常）",
          solo._last_table_parse_status == "failed" and solo._result_page_dumped is True,
          f"got {solo._last_table_parse_status}")

    async def parse_failure_in_search_one():
        checker = build()
        checker._logged_in = True
        checker._page = FakePage([])

        async def ok(**kwargs):
            return True

        async def yes():
            return True

        async def fingerprint():
            return "fp"

        async def wait_for_refresh(fp):
            return None

        checker._fill_query_form = ok
        checker._click_query_button = yes
        checker._table_fingerprint = fingerprint
        checker._wait_for_table_refresh = wait_for_refresh
        orders = await checker.search_one(
            company="深圳市甲科技有限公司", agent="乐天"
        )
        return checker, orders

    checker7, orders = asyncio.run(parse_failure_in_search_one())
    check("解析失败不产出 0 条结论",
          orders == [] and checker7._last_query_status == "结果表解析失败",
          f"got {orders} / {checker7._last_query_status}")
    check("解析失败不写查询缓存", not checker7._query_cache, f"got {list(checker7._query_cache)}")


def test_output_file_lock_fallback():
    """稳定名 xlsx 被 Excel/WPS 打开占用时，必须降级到时序副本，不能整轮崩掉。

    2026-09-14 真机踩过：用户开着 workorder_check_result.xlsx 再跑阶段二，
    openpyxl 写稳定名抛 PermissionError，因为旧实现“先写稳定名再复制副本”，
    时序副本也没生成，一整轮 RPA 查询结果全丢。
    """
    from modules.excel_writer import ExcelWriter

    lock_dir = tempfile.mkdtemp(prefix="mail_audit_lock_")
    stable = os.path.join(lock_dir, "workorder_check_result.xlsx")
    with open(stable, "wb") as f:
        f.write(b"placeholder")
    # 目标只读 → 打开写句柄直接 PermissionError，与 Excel 独占占用时的效果一致
    # （Excel 用独占共享模式打开，其他进程连 open(w) 都进不来，原文件不会被截断）
    os.chmod(stable, 0o444)
    try:
        primary, ts = ExcelWriter(output_dir=lock_dir).write_workorder_check_result(
            [{"客户公司名称": "测试公司", "是否已录单": "否"}], [], []
        )
        check("稳定名不可写时降级为时序副本且不抛异常",
              primary == ts and primary.endswith(".xlsx"), f"got {primary}")
        check("降级后时序副本已落盘",
              os.path.exists(ts) and os.path.getsize(ts) > 0, f"got {ts}")
        check("稳定名不可写时不破坏原文件",
              os.path.getsize(stable) == len(b"placeholder"))
    finally:
        os.chmod(stable, 0o666)

    free_dir = tempfile.mkdtemp(prefix="mail_audit_lock_free_")
    primary2, _ = ExcelWriter(output_dir=free_dir).write_workorder_check_result(
        [{"客户公司名称": "测试公司"}], [], []
    )
    check("未被占用时仍写稳定名",
          primary2.endswith("workorder_check_result.xlsx") and os.path.exists(primary2),
          f"got {primary2}")

    shutil.rmtree(lock_dir, ignore_errors=True)
    shutil.rmtree(free_dir, ignore_errors=True)


def test_query_cache_drops_collapsed_records():
    """旧版表头塌缩写进缓存的脏记录，必须在加载时丢弃，不能继续复用。"""
    print("== 阶段二缓存自愈 (丢弃塌缩脏记录) ==")
    import tempfile, os, json
    from modules.workorder_checker import WorkOrderChecker

    tmpdir = tempfile.mkdtemp(prefix="qa_cache_dirty_")
    try:
        cache_path = os.path.join(tmpdir, "cache.json")
        now = datetime.now().isoformat(timespec="seconds")
        with open(cache_path, "w", encoding="utf-8") as f:
            json.dump({
                "代理甲|公司X|德国_WEEE": {
                    "orders": [{"": "编辑\n审核\n详情\n更多"}], "timestamp": now,
                },
                "代理乙|公司Y|德国_WEEE": {
                    "orders": [{"客户": "公司Y", "服务项目": "德国WEEE",
                                "创建时间": "2026-09-03 10:00:00"}],
                    "timestamp": now,
                },
            }, f, ensure_ascii=False)

        checker = WorkOrderChecker({
            "url": "u", "username": "a", "password": "b",
            "query_cache_path": cache_path,
        })
        check("塌缩脏记录被丢弃", "代理甲|公司X|德国_WEEE" not in checker._query_cache,
              f"got {list(checker._query_cache)}")
        check("正常记录仍然保留", "代理乙|公司Y|德国_WEEE" in checker._query_cache)
        check("真实工单字段判定为非塌缩",
              not checker._is_collapsed_record({"客户": "甲", "服务项目": "德国WEEE"}))
        check("空记录判定为塌缩", checker._is_collapsed_record({"": "编辑"}))
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def test_workorder_table_extract_js():
    """结果表读取脚本的 DOM 逻辑（用 node + 最小 DOM 桩验证；无 node 则跳过）。"""
    print("== M5 结果表结构读取脚本 (DOM 桩测) ==")
    import subprocess

    root = os.path.dirname(os.path.abspath(__file__))
    script = os.path.join(root, "tests", "test_table_extract_js.js")
    node = os.environ.get("NODE_BIN") or shutil.which("node") or ""
    if not node:
        versions = os.path.expanduser(
            os.path.join("~", ".workbuddy", "binaries", "node", "versions")
        )
        if os.path.isdir(versions):
            for name in sorted(os.listdir(versions), reverse=True):
                candidate = os.path.join(versions, name, "node.exe")
                if os.path.exists(candidate):
                    node = candidate
                    break
    if not node or not os.path.exists(script):
        print("  [SKIP] 未找到 node 或桩测脚本，跳过（不影响其它断言）")
        return

    proc = subprocess.run(
        [node, script], capture_output=True, text=True,
        encoding="utf-8", errors="replace", timeout=120,
    )
    check(
        "结果表读取脚本通过 DOM 桩测",
        proc.returncode == 0,
        ((proc.stdout or "") + (proc.stderr or ""))[-500:],
    )


def test_workorder_batch_guards():
    """批次层：解析失败与无效客户名都不能被落成“漏单”。"""
    print("== M5 批次级解析失败 / 客户名守卫 ==")
    from modules.workorder_checker import WorkOrderChecker

    checker = WorkOrderChecker({
        "url": "u", "username": "a", "password": "b",
        "query_interval_seconds": 0, "max_consecutive_query_failures": 5,
    })
    checker._logged_in = True
    queried = []

    async def navigation_ok():
        return True

    async def fake_search_one(**kwargs):
        queried.append(kwargs["company"])
        checker._last_query_status = "结果表解析失败"
        checker._last_table_parse_note = "结果表结构无法识别"
        return []

    checker._navigate_to_order_list = navigation_ok
    checker.search_one = fake_search_one
    rows = [
        {"客户": "1家公司", "代理": "乐天", "项目": "德国WEEE"},
        {"客户": "深圳市甲科技有限公司", "代理": "乐天", "项目": "德国WEEE"},
    ]
    result = asyncio.run(checker.search_batch(rows))
    check("无效客户名不进入工单查询", queried == ["深圳市甲科技有限公司"], f"got {queried}")
    check("无效客户名标记为未比对并给出原因",
          result[0]["是否已录单"] == "未比对"
          and result[0]["匹配状态"] == "客户名称无效-跳过查询"
          and bool(result[0]["日期筛选说明"]),
          f"got {result[0].get('匹配状态')} / {result[0].get('日期筛选说明')}")
    check("结果表解析失败按未比对处理而非漏单",
          result[1]["是否已录单"] == "未比对"
          and result[1]["RPA查询状态"] == "结果表解析失败",
          f"got {result[1].get('是否已录单')} / {result[1].get('RPA查询状态')}")
    check("解析失败写入可追踪原因",
          "结果表结构无法识别" in result[1].get("_query_error", ""))


def test_workorder_customer_and_agent_guards():
    """客户字段守卫、空代理匹配、同名代理去重与全称别名。"""
    print("== M5 客户字段守卫 / 代理维度 ==")
    from modules.workorder_checker import WorkOrderChecker

    reason = WorkOrderChecker._customer_reject_reason
    check("数量描述不是公司名", bool(reason("1家公司")))
    check("指代句子不是公司名", bool(reason("这家公司需要修改公司")))
    check("请求动作片段不是公司名", bool(reason("补充法人签名和公司")))
    check("正常中文公司名可通过", reason("深圳市甲科技有限公司") == "", f"got {reason('深圳市甲科技有限公司')}")
    check("英文公司名可通过", reason("ABC TRADING CO., LTD") == "", f"got {reason('ABC TRADING CO., LTD')}")
    check("短名（如“华为”）不被误杀", reason("华为") == "", f"got {reason('华为')}")
    check("只有后缀的“公司”被拒绝", bool(reason("公司")))
    check("带后缀的短名不被误杀", reason("A公司") == "", f"got {reason('A公司')}")

    checker = WorkOrderChecker({"url": "u", "username": "a", "password": "b"})
    check("空代理不再直接判否",
          checker._match_agent("", "乐天") is True and checker._match_agent("乐天", "") is True)
    check("代理为空且其余字段命中→已录单并标注复核",
          (lambda r: r["是否已录单"] == "是" and r.get("代理校验", "").startswith("代理为空"))(
              checker.match_records(
                  [{"date": datetime(2026, 9, 1), "代理": "", "客户": "深圳市甲科技有限公司",
                    "项目": "德国WEEE"}],
                  [{"代理": "乐天", "客户": "深圳市甲科技有限公司", "服务项目": "德国WEEE",
                    "创建时间": "2026-09-03 10:00:00"}],
              )[0]
          ))
    check("同名多邮箱拼接值收敛成单名",
          WorkOrderChecker._dedupe_joined_names("向善 / 向善") == "向善")
    check("不同代理拼接保留",
          WorkOrderChecker._dedupe_joined_names("向善 / 乐天") == "向善 / 乐天")
    check("代理全称映射为下拉简称",
          checker._canonical_dropdown_value("所属代理", "深圳华之星检测有限公司") == "华之星")

    # 阶段一多匹配时不产出“向善 / 向善”
    from modules.field_extractor import FieldExtractor
    fx = FieldExtractor({
        "xiangshan@abctest.cn": {"代理": "向善", "代理简称": "向善"},
        "xiangshan1@abctest.com": {"代理": "向善", "代理简称": "向善"},
    }, project_names=[])
    agent = fx._extract_agent("xiangshan@abctest.com", "", "", [])
    check("多邮箱命中同一代理时输出单个名称",
          agent["agent"] == "向善", f"got {agent['agent']}")


def test_workorder_fuzzy_company_terms():
    """模糊重查的公司名口径：去类型说明括号 + 行政区划**逐级单独删**。

    业务口径（用户 2026-09-14 明确）：
      「xx市xx区xx有限公司」→ 第一刀删市级「xx区xx有限公司」
                              → 第二刀删区级「xx市xx有限公司」
                              → 最后一刀全删「xx有限公司」
      某一刀若把主体名剃光（只剩区名接后缀，如「xx区有限公司」），这刀作废——
      不能把区名当主体名去搜。
      「xx有限公司（个体工商户）」→ 去括号查「xx有限公司」
      括号**只删主体类型说明类**（个体工商户/自然人独资/分公司…）；括号里是
      地名或国别（实体标识）时整块保留——「深圳某某（上海）贸易有限公司」
      「某某（中国）有限公司」这类一律不动，删了会拼出平台上真实存在的
      别家公司名 → 假"已录单"。
      **公司后缀（有限公司/有限责任公司/集团…）必须保留，不得删除**，
      也不许做「取前 N 字」截短——那样会搜出一大批无关公司。
    同时必须挡住把品牌词当行政区划切掉的误伤（「城市之家」「某某夜市」）。
    """
    print("== M5 模糊公司名候选词 ==")
    from modules.workorder_checker import WorkOrderChecker

    checker = object.__new__(WorkOrderChecker)
    terms = checker._get_fuzzy_company_terms

    def got(name):
        return terms(name)

    def real(name):
        """调用方实际会去查的词（原始全名会被跳过）"""
        return [t for t in terms(name) if t != name]

    # 规则一：行政区划逐级单独删，公司后缀始终跟着主体名保留
    t = got("深圳市宝安区宝安某某有限公司")
    check("第一刀删市级（xx市→去掉，留 xx区xx）",
          t[0] == "宝安区宝安某某有限公司", f"got {t}")
    check("第二刀删区级（留 xx市xx）",
          t[1] == "深圳市宝安某某有限公司", f"got {t}")
    check("最后一刀全删区划（留主体名+后缀）",
          t[2] == "宝安某某有限公司", f"got {t}")
    check("删区级后市级保留",
          got("杭州市余杭区XX电子商务有限公司")[1] == "杭州市XX电子商务有限公司")
    check("三级区划逐级递进（省→市→区）",
          got("广东省广州市天河区某某网络科技有限公司")[0]
          == "广州市天河区某某网络科技有限公司")
    check("含“新区/开发区”的区划也能删",
          got("上海市浦东新区某某有限公司")[0] == "浦东新区某某有限公司")

    # 闸门：某一刀会把主体名剃光（只剩区名接后缀），这一刀必须作废
    check("不产出把区名当主体名的废词",
          "宝安区有限公司" not in got("深圳市宝安区有限公司")
          and real("深圳市宝安区有限公司") == [],
          f"got {got('深圳市宝安区有限公司')}")

    # 规则二：括号**只删「主体类型说明」类**（白名单，选严格侧）
    t = got("xx有限公司（个体工商户）")
    check("去全角类型说明括号", t[0] == "xx有限公司", f"got {t}")
    check("去半角括号里的类型说明",
          got("某某(自然人独资)有限公司")[0] == "某某有限公司")
    check("空括号照删（无信息量）", got("xx有限公司（）")[0] == "xx有限公司")

    # 危险侧闸门：括号里是地名/国别 → 整块保留，宁可零候选，也不拼出别家公司名
    check("地名括号保留（不拼出“某某贸易有限公司”）",
          real("深圳某某（上海）贸易有限公司") == [],
          f"got {real('深圳某某（上海）贸易有限公司')}")
    check("半角地名括号同样保留",
          real("某某(上海)贸易有限公司") == [],
          f"got {real('某某(上海)贸易有限公司')}")
    check("国名括号保留（不拼出真实存在的“某某有限公司”）",
          real("某某（中国）有限公司") == [],
          f"got {real('某某（中国）有限公司')}")
    check("境外地名括号保留",
          real("香港某某（国际）有限公司") == [],
          f"got {real('香港某某（国际）有限公司')}")

    # 去括号与删区划是两个独立维度：先出「去掉括号但区划还在」的名字，
    # 再在这条干净的基线上逐级删区划
    check("括号+区划叠加后的完整候选序列",
          real("杭州市余杭区XX电子商务有限公司（个体工商户）")
          == [
              "杭州市余杭区XX电子商务有限公司",
              "余杭区XX电子商务有限公司",
              "杭州市XX电子商务有限公司",
              "XX电子商务有限公司",
          ],
          f"got {real('杭州市余杭区XX电子商务有限公司（个体工商户）')}")

    # 后缀保留：不得出现去后缀的核心词，也不得出现前 N 字截短
    check("后缀保留→不产出“璟深电子商务”这类去后缀词",
          "璟深电子商务" not in got("深圳市璟深电子商务有限责任公司"),
          f"got {got('深圳市璟深电子商务有限责任公司')}")
    check("不作前 N 字截短→不产出“青岛积米”",
          "青岛积米" not in got("青岛积米厨房用品有限公司"),
          f"got {got('青岛积米厨房用品有限公司')}")
    check("去区划后仍带后缀（有限责任公司）",
          real("深圳市璟深电子商务有限责任公司") == ["璟深电子商务有限责任公司"],
          f"got {real('深圳市璟深电子商务有限责任公司')}")

    # 误伤守卫：品牌词里的“市/区”不能被当成行政区划
    check("“城市之家”不被误切成“之家有限公司”",
          "之家有限公司" not in got("城市之家有限公司"),
          f"got {got('城市之家有限公司')}")
    check("“某某夜市”不被误切成“有限公司/某某夜市”",
          "有限公司" not in got("某某夜市有限公司")
          and "某某夜市" not in got("某某夜市有限公司"),
          f"got {got('某某夜市有限公司')}")
    check("没有行政区划词时原样返回",
          WorkOrderChecker._strip_admin_prefix("上海某某科技有限公司")
          == "上海某某科技有限公司")
    check("剥完只剩公司后缀时放弃剥离",
          WorkOrderChecker._strip_admin_prefix("某某夜市有限公司") == "某某夜市有限公司")

    # 无任何结构可剥成分时，简体名不产候选；繁体名允许安全转换成简体再查
    check("无括号无区划→零候选", real("上海某某科技有限公司") == [],
          f"got {real('上海某某科技有限公司')}")
    check("繁体名无区划词→补充简体候选",
          real("保文商貿有限公司") == ["保文商贸有限公司"],
          f"got {real('保文商貿有限公司')}")

    # 格式变体只去空白，不删除后缀或实体标识
    check("英文公司名只补空白格式变体",
          "RONGFANGTECHNOLOGYLIMITED" in real("RONG FANG TECHNOLOGY LIMITED"),
          f"got {real('RONG FANG TECHNOLOGY LIMITED')}")
    check("格式变体仍保留法律后缀",
          all("LIMITED" in t.upper() for t in real("RONG FANG TECHNOLOGY LIMITED")),
          f"got {real('RONG FANG TECHNOLOGY LIMITED')}")

    check("原始全名排在候选末尾",
          got("深圳市甲科技有限公司")[-1] == "深圳市甲科技有限公司")
    check("空公司名不产候选", got("") == [])


def test_llm_json_lenient():
    """LLM 输出常带围栏/解释文字/尾逗号，不能因为格式细节整条丢弃。"""
    print("== LLM 输出宽松解析 ==")
    from modules.llm_intent import LLMIntentClient
    parse = LLMIntentClient._loads_lenient
    check("去除 markdown 围栏", parse("```json\n{\"is_registration\": true}\n```")["is_registration"] is True)
    check("剥离前后解释文字", parse('好的，结果是 {"is_registration": false} 请查收')["is_registration"] is False)
    check("容忍对象尾逗号", parse('{"a": [1, 2,], "b": 3,}')["b"] == 3)
    try:
        parse("完全不是 JSON 的一句话")
        bad = False
    except ValueError:
        bad = True
    check("确实无法解析时仍然报错", bad)


def test_agent_workflow_schema_gates():
    """三个 Agent 输出必须经严格类型校验；错误只修复一次并安全转人工。"""
    print("== Agent Workflow + Pydantic 输出闸门 ==")
    import json as _json
    from pydantic import ValidationError
    from modules.agent_schemas import ExtractionOutput, IntentOutput
    from modules.business_validator import apply_row_guard
    from modules.excel_writer import ExcelWriter
    from modules.llm_intent import LLMIntentClient
    from modules.mail_filter import MailFilter

    # 直接验证两个过去会被静默接受的危险类型。
    try:
        IntentOutput.model_validate({
            "is_target": "false", "intent": "", "email_type": "其他",
            "confidence": "high", "reason": "不是目标邮件",
        })
        string_false_rejected = False
    except ValidationError:
        string_false_rejected = True
    check("严格布尔: 字符串 false 被拒绝", string_false_rejected)

    try:
        ExtractionOutput.model_validate({
            "代理": "甲", "客户": "A公司", "项目": "德国WEEE",
            "记录": [], "需求": "注册", "confidence": "high",
        })
        project_string_rejected = False
    except ValidationError:
        project_string_rejected = True
    check("严格数组: 项目字符串不能冒充列表", project_string_rejected)

    english_alias = ExtractionOutput.model_validate({
        "代理": "甲", "客户": "A公司", "项目": ["德国WEEE"],
        "records": [], "需求": "注册", "confidence": "high",
    }).model_dump(by_alias=True)
    check("提取协议兼容 records 并统一输出中文记录键",
          "记录" in english_alias and "records" not in english_alias)

    class FakeCompletions:
        def __init__(self, payloads):
            self.payloads = list(payloads)
            self.calls = []

        def create(self, **kwargs):
            self.calls.append(kwargs)
            payload = self.payloads.pop(0)
            message = types.SimpleNamespace(content=_json.dumps(payload, ensure_ascii=False))
            return types.SimpleNamespace(
                choices=[types.SimpleNamespace(message=message)],
                usage=types.SimpleNamespace(prompt_tokens=0, completion_tokens=0, total_tokens=0),
                model="mock-model", _request_id="",
            )

    def build_client(payloads, audit_path):
        client = LLMIntentClient({
            "api_key": "mock", "base_url": "https://mock.invalid", "model": "mock-model",
            "audit_log_path": audit_path,
        })
        completions = FakeCompletions(payloads)
        client.client = types.SimpleNamespace(
            chat=types.SimpleNamespace(completions=completions)
        )
        return client, completions

    tmpdir = tempfile.mkdtemp(prefix="agent_workflow_")
    try:
        invalid_intent = {
            "is_target": "false", "intent": "", "email_type": "其他",
            "confidence": "high", "reason": "非目标",
        }
        valid_intent = {
            "is_target": False, "intent": "", "email_type": "其他",
            "confidence": "high", "reason": "非目标",
        }
        client, completions = build_client(
            [invalid_intent, valid_intent], os.path.join(tmpdir, "repair_ok.jsonl")
        )
        result = client.classify_email("普通通知", "无需注册", [])
        summary = client.get_usage_summary()
        check("结构错误后只修复一次并成功", result and result["is_target"] is False
              and len(completions.calls) == 2)
        check("修复统计可审计", summary["output_validation_failures"] == 1
              and summary["repair_attempts"] == 1 and summary["repair_successes"] == 1)
        repair_prompt = completions.calls[1]["messages"][1]["content"]
        check("修复请求保留原始任务上下文", "普通通知" in repair_prompt and "结构校验问题" in repair_prompt)

        failed_client, failed_calls = build_client(
            [invalid_intent, invalid_intent], os.path.join(tmpdir, "repair_failed.jsonl")
        )
        failed_result = failed_client.classify_email("普通通知", "正文", [])
        check("连续两次结构错误后停止，不发第三次", failed_result is None and len(failed_calls.calls) == 2)
        check("连续失败记录人工兜底", failed_client.get_usage_summary()["manual_review_fallbacks"] == 1)

        def review_item(rid):
            return {
                "id": rid, "status": "valid", "confidence": "high", "issues": [],
                "suggestions": {"company": "", "agent": "", "country": "", "program": "", "request": ""},
                "reason": "一致",
            }

        batch_client, batch_calls = build_client([
            {"results": [review_item("0")]},
            {"results": [review_item("0"), review_item("1")]},
        ], os.path.join(tmpdir, "batch_repair.jsonl"))
        batch = batch_client.validate_extracted_fields_batch([
            {"id": "0", "fields": {}}, {"id": "1", "fields": {}},
        ])
        check("批量复检缺失 id 会修复并一一对应", set(batch) == {"0", "1"}
              and len(batch_calls.calls) == 2)

        # 意图 Agent 本身失败时不是“确认过滤”，而是保留为人工复核候选。
        fake_llm = types.SimpleNamespace(enabled=True, classify_email=lambda *_a, **_k: None)
        mail = {"subject": "模糊邮件", "body_text": "", "attachments": [],
                "filter_reason": "外部无关邮件", "sender_email": "a@x.com"}
        recovered, filtered = MailFilter(set()).llm_second_pass([mail], fake_llm)
        check("意图失败转人工而非静默过滤", len(recovered) == 1 and not filtered
              and recovered[0]["filter_status"] == "uncertain")

        placeholder_row = {"客户": "一家公司", "项目": "德国WEEE", "需求": "注册"}
        apply_row_guard(placeholder_row)
        check("占位公司名被确定性规则拦截", placeholder_row["置信度"] == "需人工确认"
              and "占位词" in placeholder_row["业务规则问题"])

        form_sentence_row = {
            "客户": "在其他欧盟国家或第三国设立的公司",
            "项目": "比利时包装法", "需求": "注册",
        }
        apply_row_guard(form_sentence_row)
        check(
            "申请表说明句被确定性规则拦截",
            form_sentence_row["置信度"] == "需人工确认"
            and "说明句" in form_sentence_row["业务规则问题"],
            form_sentence_row,
        )

        good = {"代理": "甲", "客户": "A公司", "项目": "德国WEEE", "需求": "注册",
                "代理匹配方式": "邮箱精确匹配", "置信度": "high"}
        low = {**good, "客户": "B公司", "置信度": "low"}
        missing = {**good, "客户": "C公司", "项目": ""}
        to_work, to_review = ExcelWriter._partition_rows([good, low, missing])
        check("阶段二闸门只放行完整高置信行", len(to_work) == 1 and len(to_review) == 2)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def test_umbrella_epr_dedupe():
    """「国家+EPR」统称：同国已有具体项目时丢弃，只有统称时保留并标注。"""
    print("== 国家+EPR 统称去重 ==")
    from modules.project_normalizer import bare_epr_country, prune_umbrella_epr

    check("识别裸统称 奥地利EPR", bare_epr_country("奥地利EPR") == "奥地利")
    check("识别带空格 奥地利 EPR", bare_epr_country("奥地利 EPR") == "奥地利")
    check("具体项目不算统称", bare_epr_country("奥地利WEEE") == "")
    check("没有国名的 EPR 不算", bare_epr_country("EPR") == "")
    check("非项目表国家不算", bare_epr_country("测试EPR") == "")

    # 同封邮件同时给了统称和三项 → 统称冗余，丢弃；三项保留
    drop, flag = prune_umbrella_epr(["奥地利EPR", "奥地利WEEE", "奥地利电池法", "奥地利包装法"])
    check("同国具体项已存在时丢弃统称", drop == {0} and flag == set(), f"drop={drop} flag={flag}")

    # 整封只写了统称 → 保留但要人工确认
    drop, flag = prune_umbrella_epr(["奥地利EPR"])
    check("只有统称时保留并标注", drop == set() and flag == {0}, f"drop={drop} flag={flag}")

    # 多国混合：奥地利有具体项，德国没有
    drop, flag = prune_umbrella_epr(["波兰包装法", "奥地利EPR", "奥地利WEEE", "德国EPR"])
    check("按国家分别判定", drop == {1} and flag == {3}, f"drop={drop} flag={flag}")

    # 重复统称一起去重
    drop, flag = prune_umbrella_epr(["奥地利EPR", "奥地利EPR", "奥地利包装法"])
    check("重复统称全部丢弃", drop == {0, 1} and flag == set(), f"drop={drop} flag={flag}")

    # 无统称时不误伤
    drop, flag = prune_umbrella_epr(["德国WEEE", "法国电池法"])
    check("没有统称时不动任何行", drop == set() and flag == set(), f"drop={drop} flag={flag}")


def test_semantic_validation_same_mail_programs():
    """语义校验要把同封邮件的全部项目一并给 LLM，避免把拆分结果误判成漏拆。"""
    print("== 语义校验带上同封项目 ==")
    from gui import WorkerThread
    from modules.llm_intent import LLMIntentClient, SYSTEM_PROMPT_VALIDATE_BATCH

    check("提示词说明同封邮件全部项目", "同封邮件全部项目" in SYSTEM_PROMPT_VALIDATE_BATCH)
    check("提示词禁止按拆分行判漏拆", "漏拆" in SYSTEM_PROMPT_VALIDATE_BATCH)

    captured = {}

    class FakeLLM:
        enabled = True
        model = "fake-model"

        def validate_extracted_fields_batch(self, items):
            captured["items"] = items
            return {x["id"]: {"status": "valid", "confidence": "high", "issues": [],
                              "suggestions": {}, "reason": "ok"} for x in items}

    worker = types.SimpleNamespace(config={"llm": {"semantic_batch_size": 6}}, _stop=False)
    validate = WorkerThread._semantic_validate_rows
    rows = [
        {"项目": "奥地利WEEE", "客户": "C公司", "sender_email": "a@b.com",
         "subject": "奥地利三项", "body_text": "正文", "代理": "华之星", "需求": "注册"},
        {"项目": "奥地利电池法", "客户": "C公司", "sender_email": "a@b.com",
         "subject": "奥地利三项", "body_text": "正文", "代理": "华之星", "需求": "注册"},
        {"项目": "德国WEEE", "客户": "D公司", "sender_email": "c@d.com",
         "subject": "德国一封", "body_text": "正文", "代理": "华之星", "需求": "注册"},
    ]
    out = validate(worker, rows, FakeLLM(), __import__("logging").getLogger("t"),
                   mail_keys=[0, 0, 1])
    items = captured.get("items") or []
    check("三行都送进校验", len(items) == 3, f"got {len(items)}")
    check("同封两行共享完整项目清单",
          items[0]["same_mail_programs"] == ["奥地利WEEE", "奥地利电池法"]
          and items[1]["same_mail_programs"] == ["奥地利WEEE", "奥地利电池法"],
          f"got {items[0].get('same_mail_programs')}")
    check("另一封邮件不串项目", items[2]["same_mail_programs"] == ["德国WEEE"],
          f"got {items[2].get('same_mail_programs')}")
    check("校验结果正常落回", out[0]["语义校验状态"] == "valid" and len(out) == 3)

    # 多吐 id 的修复响应不应污染未校验的行
    class ExtraIdLLM(FakeLLM):
        def validate_extracted_fields_batch(self, items):
            res = super().validate_extracted_fields_batch(items)
            res["999"] = {"status": "invalid", "confidence": "low", "issues": [],
                          "suggestions": {}, "reason": "不属于本批"}
            return res

    rows2 = [dict(r) for r in rows]
    out2 = validate(worker, rows2, ExtraIdLLM(), __import__("logging").getLogger("t"),
                    mail_keys=[0, 0, 1])
    check("本批之外的 id 被丢弃", all(r["语义校验状态"] == "valid" for r in out2),
          [r["语义校验状态"] for r in out2])

    # 客户端确实把同封清单放进请求体
    client = LLMIntentClient.__new__(LLMIntentClient)
    sent = {}

    def fake_call(system, user, purpose=None, **kw):
        sent["user"] = user
        return {"results": [{
            "id": "0", "status": "valid", "confidence": "high",
            "issues": [],
            "suggestions": {"company": "", "agent": "", "country": "", "program": "", "request": ""},
            "reason": "字段一致",
        }]}

    client._call_api = fake_call
    client._loads_lenient = staticmethod(lambda raw: {})
    client.validate_extracted_fields_batch([
        {"id": "0", "subject": "s", "body": "b", "fields": {"program": "奥地利WEEE"},
         "same_mail_programs": ["奥地利WEEE", "奥地利电池法"]}
    ])
    check("请求体带上同封邮件全部项目", "同封邮件全部项目" in sent.get("user", ""), sent.get("user", "")[:120])


def test_workorder_query_timing():
    """查询提速：短沉降 + 指纹轮询，不再固定空等 2 秒 / 死等 5 秒。

    2026-09-14 实测（logs/run_20260914_161414.log）：旧实现里
    _click_query_button 写死 sleep(2)，_wait_for_table_refresh 又复用
    dropdown_timeout_seconds(5 秒) 当上限。命中新结果时每次查询 2 秒，
    而"本次结果与上次逐字相同"（两次都命中暂无数据）时会死等满 5 秒 ——
    模糊重查一条要连查 3~4 轮，十几秒全花在等上。
    """
    print("== M5 查询提速 ==")
    import asyncio
    import time

    from modules.workorder_checker import WorkOrderChecker

    base = {
        "url": "https://example.test/workbench/main",
        "username": "a",
        "password": "b",
    }

    # 1) 参数默认值与解耦：刷新上限不再跟着下拉超时走
    c = WorkOrderChecker(base)
    check("点击后沉降默认 0.3 秒（原为写死的 2 秒）",
          c.query_click_settle_seconds == 0.3, f"got {c.query_click_settle_seconds}")
    check("刷新等待上限默认 3 秒（原复用 dropdown 的 5 秒）",
          c.table_refresh_timeout_seconds == 3.0,
          f"got {c.table_refresh_timeout_seconds}")
    check("刷新上限与下拉超时是两个独立配置",
          c.table_refresh_timeout_seconds != c.dropdown_timeout_seconds)
    check("刷新上限可配置",
          WorkOrderChecker({**base, "table_refresh_timeout_seconds": 1.5})
          .table_refresh_timeout_seconds == 1.5)
    check("沉降可配成 0（完全不等）",
          WorkOrderChecker({**base, "query_click_settle_seconds": 0})
          .query_click_settle_seconds == 0.0)

    def seq_fingerprint(values):
        """按序吐出指纹的假指纹函数；序列用完后一直重复最后一个值。"""
        state = {"i": 0}

        async def fake():
            i = state["i"]
            state["i"] = i + 1
            return values[i] if i < len(values) else values[-1]

        return fake

    # 2) 结果变了且连续两次一致 → 立刻返回，不陪着睡满上限
    c2 = WorkOrderChecker({**base, "table_refresh_timeout_seconds": 3.0})
    c2._table_fingerprint = seq_fingerprint(["old", "new", "new", "new"])
    t0 = time.monotonic()
    asyncio.run(c2._wait_for_table_refresh("old"))
    dt2 = time.monotonic() - t0
    check("结果刷新后立即返回（约 0.2 秒，不等满 3 秒）", dt2 < 0.6, f"耗时 {dt2:.2f}s")

    # 3) 表格"变了又变"（清空 → 渲染）时必须等到稳定，否则会把空表读成 0 条
    c3 = WorkOrderChecker({**base, "table_refresh_timeout_seconds": 3.0})
    c3._table_fingerprint = seq_fingerprint(["old", "half", "full", "full", "full"])
    t0 = time.monotonic()
    asyncio.run(c3._wait_for_table_refresh("old"))
    dt3 = time.monotonic() - t0
    check("半渲染（变了又变）时继续等到指纹稳定", 0.2 <= dt3 < 1.2, f"耗时 {dt3:.2f}s")

    # 4) 指纹始终不变 → 按新上限返回，不再死等到 dropdown 的 5 秒
    c4 = WorkOrderChecker({**base, "table_refresh_timeout_seconds": 0.8,
                           "dropdown_timeout_seconds": 5})
    c4._table_fingerprint = seq_fingerprint(["same"])
    t0 = time.monotonic()
    asyncio.run(c4._wait_for_table_refresh("same"))
    dt4 = time.monotonic() - t0
    check("结果与上次相同时按新上限返回（不再死等 5 秒）", 0.6 <= dt4 < 1.6,
          f"耗时 {dt4:.2f}s")

    # 5) 无头下跳过过程截图（没人看，逐条累计很可观）；有头仍保留
    class ShotPage:
        def __init__(self):
            self.shots = []

        async def screenshot(self, path, full_page=False):
            self.shots.append(path)

    hp = WorkOrderChecker({**base, "headless": True})
    hp._page = ShotPage()
    asyncio.run(hp._debug_screenshot("output/x.png"))
    check("无头下不截过程图", hp._page.shots == [], f"got {hp._page.shots}")

    pp = WorkOrderChecker({**base, "headless": False})
    pp._page = ShotPage()
    asyncio.run(pp._debug_screenshot("output/x.png"))
    check("有头下保留过程图", pp._page.shots == ["output/x.png"], f"got {pp._page.shots}")


def test_workorder_headless_mode():
    """无头运行：配置解析 + 登录态复用 + 有头降级判定。

    登录页有图形验证码，无头下不可能人工输入，所以无头完全依赖
    storage/login_state.json 的登录态复用；失效时必须能切回可视化窗口。
    """
    print("== M5 无头运行 ==")
    import json

    from modules.workorder_checker import WorkOrderChecker

    base = {"url": "https://example.test/workbench/main", "username": "a", "password": "b"}

    # 1) 无头下"给人看的停顿"必须归零（没人看，纯空转）
    headed = WorkOrderChecker({**base, "headless": False,
                               "input_step_delay_seconds": 0.8,
                               "cached_input_hold_seconds": 1.2})
    check("非无头保留可视化停留", headed.input_step_delay == 0.8 and headed.cached_input_hold == 1.2,
          f"got {headed.input_step_delay}/{headed.cached_input_hold}")
    check("默认不是无头（老配置不受影响）", headed.headless is False)

    hless = WorkOrderChecker({**base, "headless": True,
                              "input_step_delay_seconds": 0.8,
                              "cached_input_hold_seconds": 1.2})
    check("无头下字段停留归零", hless.input_step_delay == 0.0, f"got {hless.input_step_delay}")
    check("无头下缓存重放停留归零", hless.cached_input_hold == 0.0, f"got {hless.cached_input_hold}")
    check("无头下跳过缓存输入重放（没人看，纯省时间）",
          hless.show_cached_inputs is False, f"got {hless.show_cached_inputs}")
    check("非无头仍保留缓存输入重放",
          WorkOrderChecker({**base, "headless": False,
                            "show_cached_inputs": True}).show_cached_inputs is True)
    check("登录态路径默认落在 storage/",
          hless.login_state_path.replace("\\", "/") == "storage/login_state.json",
          hless.login_state_path)
    check("登录态失效默认允许弹窗降级", hless.headless_fallback is True)
    check("登录等待时长可配置", WorkOrderChecker(
        {**base, "headless": True, "login_wait_seconds": 60}).login_wait_seconds == 60)

    # 2) 登录页判定以"密码框可见"为准；只认可见元素，隐藏的不算
    class FakeEl:
        def __init__(self, visible):
            self._visible = visible

        async def is_visible(self):
            return self._visible

    class FakePage:
        def __init__(self, visible):
            self._visible = visible

        async def query_selector(self, _sel):
            return FakeEl(self._visible) if self._visible is not None else None

    c = object.__new__(WorkOrderChecker)
    from modules.workorder_checker import DEFAULT_SELECTORS
    c.selectors = DEFAULT_SELECTORS
    c._page = FakePage(True)
    check("登录表单页被识别", asyncio.run(c._is_login_form_page()) is True)
    c._page = FakePage(None)
    check("已登录页不被误判为登录表单", asyncio.run(c._is_login_form_page()) is False)

    # 3) 登录态落盘：把上下文里的 cookie 写到指定路径
    with tempfile.TemporaryDirectory() as tmp:
        state_path = os.path.join(tmp, "sub", "login_state.json")

        class FakeContext:
            def __init__(self):
                self.saved_to = None

            async def storage_state(self, path=None):
                self.saved_to = path
                with open(path, "w", encoding="utf-8") as fh:
                    json.dump({"cookies": [{"name": "token", "value": "x"}],
                               "origins": [{"origin": "https://example.test",
                                            "localStorage": [{"name": "token", "value": "y"}]}]}, fh)
                return {"cookies": [], "origins": []}

        c2 = object.__new__(WorkOrderChecker)
        c2.login_state_path = state_path
        c2._context = FakeContext()
        c2.logger = None
        c2._log = lambda *a, **k: None
        ok = asyncio.run(c2._save_login_state())
        check("登录态写入成功", ok is True)
        check("自动创建 storage 目录", os.path.exists(state_path))
        saved = json.load(open(state_path, encoding="utf-8"))
        check("登录态含 cookie 与 localStorage",
              bool(saved.get("cookies")) and bool(saved.get("origins")),
              json.dumps(saved, ensure_ascii=False)[:120])

    # 4) storage_state 取 cookie 快照时不落盘（path=None）也不能报错
    c3 = object.__new__(WorkOrderChecker)
    c3.login_state_path = os.path.join(tempfile.gettempdir(), "no_such_dir_x", "s.json")
    c3._log = lambda *a, **k: None

    class BadContext:
        async def storage_state(self, path=None):
            raise RuntimeError("boom")

    c3._context = BadContext()
    check("登录态保存失败不抛异常（只是少一份缓存）",
          asyncio.run(c3._save_login_state()) is False)

    # 5) login() 的三条关键分支：直接复用 / 无头下拒绝弹窗 / 无头下降级弹窗
    class FakePage:
        url = "https://example.test/workbench/main"

        async def goto(self, url, **kw):
            type(self).url = url

        async def wait_for_load_state(self, *a, **k):
            pass

        async def screenshot(self, **kw):
            pass

        async def evaluate(self, _js):
            return ""

        async def wait_for_selector(self, *a, **k):
            raise RuntimeError("no element")

    def build(headless, fallback, login_form_results, restarts, restores):
        ck = object.__new__(WorkOrderChecker)
        ck._logged_in = False
        ck.headless = headless
        ck.headless_fallback = fallback
        ck._headless_restore_failed = False
        ck._headed_fallback_used = False
        ck.selectors = DEFAULT_SELECTORS
        ck.login_wait_seconds = 1
        ck.login_state_path = os.path.join(tempfile.gettempdir(), "no_login_state_here.json")
        ck.url = "https://example.test/workbench/main"
        ck._context = None
        ck._page = None
        ck.logger = None
        ck._log = lambda *a, **k: None
        state = {"saved": 0}
        results = list(login_form_results)

        async def fake_ensure(headless=None):
            ck._page = FakePage()

        async def fake_is_login_form():
            return results.pop(0) if results else False

        async def fake_save():
            state["saved"] += 1
            return True

        async def fake_restart(headless):
            restarts.append(headless)
            ck._page = FakePage()

        async def fake_restore():
            restores.append(True)
            return True

        async def fake_after():
            pass

        ck._ensure_browser = fake_ensure
        ck._is_login_form_page = fake_is_login_form
        ck._save_login_state = fake_save
        ck._restart_browser = fake_restart
        ck._restore_headless = fake_restore
        ck._after_login_ready = fake_after
        ck._state = state
        return ck

    # 5a) 登录态有效：直接复用，不重启浏览器、不弹窗
    restarts, restores = [], []
    ck = build(True, True, [False], restarts, restores)
    check("无头+登录态有效 → 直接复用免验证码", asyncio.run(ck.login()) is True)
    check("复用成功不重开浏览器", restarts == [], restarts)
    check("复用成功回写登录态（刷新有效期）", ck._state["saved"] == 1, ck._state)

    # 5b) 登录态失效 + 关掉降级：直接放弃，绝不偷偷弹窗
    restarts, restores = [], []
    ck = build(True, False, [True], restarts, restores)
    check("无头+失效+禁降级 → 返回失败", asyncio.run(ck.login()) is False)
    check("禁降级时不弹浏览器", restarts == [], restarts)

    # 5c) 登录态失效 + 允许降级：切有头，登录后切回无头
    restarts, restores = [], []
    ck = build(True, True, [True, False], restarts, restores)
    check("无头+失效 → 切有头并免验证码成功", asyncio.run(ck.login()) is True)
    check("先切有头", restarts == [False], restarts)
    check("登录后切回无头", restores == [True], restores)
    check("降级登录也回写登录态", ck._state["saved"] == 1, ck._state)


def test_workbench_database():
    """独立工作台数据库应支持增量幂等写入和按操作日期读取。"""
    print("== 工作台 SQLite 增量数据库 ==")
    from pathlib import Path
    from workbench_database import WorkbenchDatabase

    tmp = Path(tempfile.mkdtemp())
    try:
        db = WorkbenchDatabase(tmp / "workbench.db")
        row = {
            "sender_email": "agent@example.com",
            "date": "2026-09-17 09:00:00",
            "subject": "德国WEEE注册",
            "body_text": "公司甲",
            "attachments": [{"filename": "申请表.xlsx"}],
            "agent": "代理甲",
            "company": "公司甲",
            "program": "德国WEEE",
            "request": "注册",
            "_id": "legacy-1",
        }
        first = db.ingest_rows([row], "active", "API测试")
        second = db.ingest_rows([row], "active", "API测试")
        rows = db.read_rows("active")
        check("数据库首次写入新增一条", first.get("inserted") == 1 and first.get("updated") == 0, first)
        check("数据库重复写入幂等更新", second.get("inserted") == 0 and second.get("updated") == 1, second)
        check("数据库保存内部字段并规范附件名", len(rows) == 1 and rows[0].get("客户公司名称") == "公司甲" and rows[0].get("附件名称") == "申请表.xlsx", rows)
        db.record_operation("confirm", record_id=rows[0]["_id"], operation_date="2026-09-17", reason="测试确认")
        operations = db.read_operations("2026-09-17")
        check("数据库按操作日期读取留痕", len(operations) == 1 and operations[0]["action"] == "confirm", operations)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_workbench_persistent_mail_and_detail_numbers():
    """邮件/明细编号跨重复读取保持不变，部分处理后可从 SQLite/JSON 恢复。"""
    print("== 工作台 邮件与明细持久编号 ==")
    from pathlib import Path
    from openpyxl import Workbook
    import workbench_server as W

    tmp = Path(tempfile.mkdtemp())
    old_review_state = W.REVIEW_STATE
    old_history_dir = W.HISTORY_DIR
    old_history_state = W.HISTORY_STATE
    old_history_output = W.HISTORY_OUTPUT
    try:
        primary = tmp / "to_workorder_list.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "工单待查"
        ws.append(["发件人邮箱", "发件日期", "邮件主题", "邮件正文摘要(最多300字)", "附件名称",
                   "代理", "客户公司名称", "标准化项目名称", "需求", "置信度"])
        ws.append(["idtest@example.com", "2026-09-21 09:00:00", "编号持久化测试", "正文", "申请表.xlsx",
                   "代理甲", "测试公司甲", "德国包装法", "注册", "high"])
        ws.append(["idtest@example.com", "2026-09-21 09:00:00", "编号持久化测试", "正文", "申请表.xlsx",
                   "代理甲", "测试公司乙", "德国包装法", "注册", "high"])
        wb.save(primary)

        W.REVIEW_STATE = tmp / "storage" / "workbench_review.json"
        W.HISTORY_DIR = tmp / "storage" / "workbench_history"
        W.HISTORY_STATE = W.HISTORY_DIR / "mail_history.json"
        W.HISTORY_OUTPUT = tmp / "output" / "workbench_history"
        store = W.WorkbenchStore(
            primary_path=str(primary), review_path=str(tmp / "none.xlsx"),
            filtered_path=str(tmp / "none_filtered.xlsx"),
        )
        first = store.snapshot()
        mail = first["mails"][0]
        details = mail["details"]
        assert mail.get("mail_number", "").startswith("MAIL-")
        assert len(details) == 2 and all(d.get("detail_number", "").startswith("DETAIL-") for d in details)
        mail_number = mail["mail_number"]
        detail_numbers = {d["id"]: d["detail_number"] for d in details}

        result = store.action({
            "action": "confirm_detail",
            "mail_id": mail["id"],
            "mail_number": mail_number,
            "record_id": details[0]["id"],
            "detail_number": details[0]["detail_number"],
            "reason": "只处理第一条",
        })
        assert result["ok"] and result["audit"]["mail_number"] == mail_number
        assert result["audit"]["detail_number"] == details[0]["detail_number"]

        restarted = W.WorkbenchStore(
            primary_path=str(primary), review_path=str(tmp / "none.xlsx"),
            filtered_path=str(tmp / "none_filtered.xlsx"),
        )
        second = restarted.snapshot()
        mail2 = second["mails"][0]
        detail2 = {d["id"]: d for d in mail2["details"]}
        assert mail2["mail_number"] == mail_number
        assert {d["id"]: d["detail_number"] for d in mail2["details"]} == detail_numbers
        assert detail2[details[0]["id"]]["status"] == "confirmed"
        assert detail2[details[1]["id"]]["status"] != "confirmed"
        operations = restarted.database.read_operations()
        assert operations and operations[0]["result"]["audit"]["detail_number"] == details[0]["detail_number"]
    finally:
        W.REVIEW_STATE = old_review_state
        W.HISTORY_DIR = old_history_dir
        W.HISTORY_STATE = old_history_state
        W.HISTORY_OUTPUT = old_history_output
        shutil.rmtree(tmp, ignore_errors=True)


def test_workbench_change_exports_and_idempotent_actions():
    """人工修改/新增明细应实时留痕，重试同一请求不能重复写入。"""
    print("== 工作台 修改/新增明细导出与幂等操作 ==")
    from pathlib import Path
    from openpyxl import Workbook, load_workbook
    import workbench_server as W

    tmp = Path(tempfile.mkdtemp())
    old_review_state, old_history_dir = W.REVIEW_STATE, W.HISTORY_DIR
    old_history_state, old_history_output = W.HISTORY_STATE, W.HISTORY_OUTPUT
    try:
        primary = tmp / "to_workorder_list.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "工单待查"
        ws.append(["发件人邮箱", "发件日期", "邮件主题", "代理", "客户公司名称", "标准化项目名称", "需求", "置信度"])
        ws.append(["audit@example.com", "2026-09-21 09:00:00", "变更导出测试", "代理甲", "原公司", "德国包装法", "注册", "high"])
        wb.save(primary)
        W.REVIEW_STATE = tmp / "storage" / "workbench_review.json"
        W.HISTORY_DIR = tmp / "storage" / "workbench_history"
        W.HISTORY_STATE = W.HISTORY_DIR / "mail_history.json"
        W.HISTORY_OUTPUT = tmp / "output" / "workbench_history"
        store = W.WorkbenchStore(primary_path=str(primary), review_path=str(tmp / "none.xlsx"), filtered_path=str(tmp / "none_filtered.xlsx"))
        snap = store.snapshot()
        mail, detail = snap["mails"][0], snap["mails"][0]["details"][0]
        payload = {
            "action": "edit", "request_id": "test-edit-once", "mail_id": mail["id"],
            "mail_number": mail["mail_number"], "detail_number": detail["detail_number"],
            "record_id": detail["id"], "subject": mail["subject"],
            "fields": {"company": "修改后公司"}, "reason": "测试修改",
        }
        first = store.action(payload)
        second = store.action(payload)
        assert first.get("change_type") == "modified_detail" and second == first
        state = json.loads(W.REVIEW_STATE.read_text(encoding="utf-8"))
        assert len([x for x in state.get("change_log", []) if x.get("kind") == "modified_detail"]) == 1
        added = store.action({
            "action": "add_project", "request_id": "test-add-once", "mail_id": mail["id"],
            "mail_number": mail["mail_number"], "mail": {"sender": mail["sender"], "date": mail["date"], "subject": mail["subject"]},
            "fields": {"agent": "代理甲", "company": "新增公司", "country": "德国", "program": "德国WEEE", "request": "注册"},
            "reason": "测试新增",
        })
        assert added.get("change_type") == "added_detail"
        out = store.export()
        exported = load_workbook(out, read_only=True, data_only=True)
        assert "修改明细" in exported.sheetnames and "新增明细" in exported.sheetnames
        assert exported["修改明细"].max_row >= 2 and exported["新增明细"].max_row >= 2
        exported.close()
    finally:
        W.REVIEW_STATE, W.HISTORY_DIR = old_review_state, old_history_dir
        W.HISTORY_STATE, W.HISTORY_OUTPUT = old_history_state, old_history_output
        shutil.rmtree(tmp, ignore_errors=True)


def test_workbench_missing_mail_round_trip():
    """漏单可人工同步，也可进入自动重查队列，且结果覆盖旧漏单。"""
    print("== 工作台 漏单回流与工单同步 ==")
    from pathlib import Path
    from openpyxl import Workbook
    import workbench_server as W

    tmp = Path(tempfile.mkdtemp())
    original = (
        W.REVIEW_STATE, W.HISTORY_DIR, W.HISTORY_STATE, W.HISTORY_OUTPUT,
        W.COMPLETED_HISTORY_XLSX, W.UNFINISHED_HISTORY_XLSX, W.WORKORDER_RETRY_QUEUE,
    )
    try:
        primary = tmp / "to_workorder_list.xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "工单待查"
        ws.append(["发件人邮箱", "发件日期", "邮件主题", "邮件正文摘要(最多300字)", "附件名称",
                   "代理", "客户公司名称", "标准化项目名称", "需求", "置信度"])
        ws.append(["missing@example.com", "2026-09-17 12:00:00", "漏单回流测试", "请查公司甲",
                   "申请表.xlsx", "代理甲", "公司甲", "德国WEEE", "注册", "high"])
        wb.save(primary)
        result = tmp / "workorder_check_result.xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "工单核对"
        ws.append(["发件人邮箱", "发件日期", "邮件主题", "客户公司名称", "标准化项目名称", "需求",
                   "代理", "是否已录单", "匹配状态", "RPA查询状态", "查询时间戳"])
        ws.append(["missing@example.com", "2026-09-17 12:00:00", "漏单回流测试", "公司甲", "德国WEEE", "注册",
                   "代理甲", "否", "漏单", "实时查询", "2026-09-17 15:00:00"])
        wb.save(result)

        W.REVIEW_STATE = tmp / "storage" / "workbench_review.json"
        W.HISTORY_DIR = tmp / "storage" / "workbench_history"
        W.HISTORY_STATE = W.HISTORY_DIR / "mail_history.json"
        W.HISTORY_OUTPUT = tmp / "output" / "workbench_history"
        W.COMPLETED_HISTORY_XLSX = W.HISTORY_OUTPUT / "邮件处理完成总表.xlsx"
        W.UNFINISHED_HISTORY_XLSX = W.HISTORY_OUTPUT / "邮件未处理完成总表.xlsx"
        W.WORKORDER_RETRY_QUEUE = tmp / "storage" / "workorder_retry_queue.json"
        store = W.WorkbenchStore(
            primary_path=str(primary), review_path=str(tmp / "none.xlsx"),
            filtered_path=str(tmp / "none_filtered.xlsx"), workorder_result_path=str(result),
        )
        initial = store.snapshot()
        mail = initial["mails"][0]
        detail = mail["details"][0]
        check("漏单进入独立视图", len(initial["missing_mails"]) == 1, initial["missing_mails"])

        manual = store.action({
            "action": "manual_workorder_result", "mail_id": mail["id"],
            "record_id": detail["id"], "mail": mail, "fields": detail["fields"],
            "found": "是", "match_status": "人工确认-已找到", "workorder_number": "WO-1001",
            "workorder_date": "2026-09-17", "note": "人工在工单系统查到一条",
        })
        after_manual = store.snapshot()
        summary = after_manual["mails"][0]["workorder_summary"]
        check("人工同步覆盖旧漏单", manual["ok"] and summary.get("found") == 1 and summary.get("missing") == 0, summary)
        check("人工工单同步写入明细留痕", any(e.get("action") == "manual_workorder_result" for e in after_manual["mails"][0]["details"][0]["events"]), after_manual["mails"][0]["details"][0]["events"])
        check("人工工单同步写入操作日志", store.database.summary().get("operations", 0) == 1, store.database.summary())

        queued = store.action({
            "action": "queue_workorder_retry", "mail_id": mail["id"],
            "record_id": detail["id"], "mail": mail, "fields": detail["fields"],
            "reason": "再次自动核查",
        })
        queue = W._load_json(W.WORKORDER_RETRY_QUEUE)
        check("漏单可加入自动重查队列", queued["ok"] and len(queue.get("items", [])) == 1, queue)
        check("自动重查动作也写入明细留痕", any(e.get("action") == "queue_workorder_retry" for e in after_manual["mails"][0]["details"][0]["events"]) is False, "动作发生在下一次快照")
        after_queue = store.snapshot()
        check("自动重查动作刷新后可追溯", any(e.get("action") == "queue_workorder_retry" for e in after_queue["mails"][0]["details"][0]["events"]), after_queue["mails"][0]["details"][0]["events"])
    finally:
        (
            W.REVIEW_STATE, W.HISTORY_DIR, W.HISTORY_STATE, W.HISTORY_OUTPUT,
            W.COMPLETED_HISTORY_XLSX, W.UNFINISHED_HISTORY_XLSX, W.WORKORDER_RETRY_QUEUE,
        ) = original
        shutil.rmtree(tmp, ignore_errors=True)


def main():
    tests = [
        test_mail_reader_uses_real_uids,
        test_mail_reader_invalid_date_is_tolerated,
        test_mail_filter,
        test_field_extractor,
        test_project_normalizer,
        test_country_of_project,
        test_non_registration_filter,
        test_workorder_match,
        test_workorder_navigation_contract,
        test_workorder_cancellation,
        test_workorder_row_failure_continues,
        test_workorder_agent_alias,
        test_workorder_result_table_parse,
        test_output_file_lock_fallback,
        test_workorder_table_extract_js,
        test_workorder_batch_guards,
        test_workorder_customer_and_agent_guards,
        test_workorder_fuzzy_company_terms,
        test_workorder_query_timing,
        test_workorder_headless_mode,
        test_llm_json_lenient,
        test_agent_workflow_schema_gates,
        test_umbrella_epr_dedupe,
        test_semantic_validation_same_mail_programs,
        test_excel_writer,
        test_ocr_compat,
        test_archive_safety,
        test_ocr_fallback,
        test_stage2_preprocess,
        test_stage2_preprocess_xlsx_headers,
        test_query_cache_roundtrip,
        test_query_cache_drops_collapsed_records,
        test_cached_query_replays_visible_inputs,
        test_stage2_outputs,
        test_llm_intent_mock,
        test_llm_audit_metadata,
        test_llm_fallback_in_field_extractor,
        test_epr_form_parser,
        test_epr_checkbox_priority,
        test_attachment_xlsx_record_authority,
        test_multi_project_company_count_and_generic_phrase,
        test_company_name_cross_source_consensus_and_code_prefix,
        test_epr_form_labels_are_not_customer_records,
        test_stage1_semantic_routing_and_placeholder_dedupe,
        test_semantic_issue_template_fields,
        test_workbench_add_delete_project,
        test_workbench_bulk_confirm,
        test_workbench_attachment_record_audit,
        test_workbench_prefers_valid_history_over_stale_rows,
        test_workbench_persistent_history_totals,
        test_workbench_filtered_route_round_trip,
        test_workbench_history_records_workorder_outcomes,
        test_workbench_database,
        test_workbench_persistent_mail_and_detail_numbers,
        test_workbench_change_exports_and_idempotent_actions,
        test_workbench_missing_mail_round_trip,
    ]
    for t in tests:
        t()
    print("=" * 40)
    print(f"结果: {PASS} 通过, {FAIL} 失败")
    sys.exit(1 if FAIL else 0)


if __name__ == "__main__":
    main()
