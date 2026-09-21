"""本机邮件询单人工复核工作台。

工作台是阶段一输出的可视化复核层，不替代现有邮件解析和阶段二 RPA：
原始 Excel 只读，人工修改保存在 storage/workbench_review.json，确认后另存
为 output/manual_review/workbench_reviewed_*.xlsx。
"""

from __future__ import annotations

import hashlib
import io
import json
import mimetypes
import os
import re
import socket
import threading
import uuid
import webbrowser
from datetime import datetime, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
from urllib.parse import parse_qs, quote, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from modules.project_normalizer import country_of_project
from workbench_database import WorkbenchDatabase
from utils.runtime_paths import APP_ROOT, resolve_runtime_path


OUTPUT_ROOT = APP_ROOT / "output"
STAGE1_OUTPUT = OUTPUT_ROOT / "stage1_email"
STAGE2_OUTPUT = OUTPUT_ROOT / "stage2_workorder"
MANUAL_OUTPUT = OUTPUT_ROOT / "manual_review"
DEFAULT_PRIMARY = STAGE1_OUTPUT / "to_workorder_list.xlsx"
DEFAULT_REVIEW = STAGE1_OUTPUT / "to_review_list.xlsx"
DEFAULT_FILTERED = STAGE1_OUTPUT / "filtered_mail_record.xlsx"
DEFAULT_WORKORDER_RESULT = STAGE2_OUTPUT / "workorder_check_result.xlsx"
DEFAULT_MISSING_WORKORDER = STAGE2_OUTPUT / "漏单.xlsx"
LEGACY_PRIMARY = OUTPUT_ROOT / "to_workorder_list.xlsx"
LEGACY_REVIEW = OUTPUT_ROOT / "to_review_list.xlsx"
LEGACY_FILTERED = OUTPUT_ROOT / "filtered_mail_record.xlsx"
LEGACY_WORKORDER_RESULT = OUTPUT_ROOT / "workorder_check_result.xlsx"
LEGACY_MISSING_WORKORDER = OUTPUT_ROOT / "漏单.xlsx"
REVIEW_STATE = APP_ROOT / "storage" / "workbench_review.json"
HISTORY_DIR = APP_ROOT / "storage" / "workbench_history"
HISTORY_STATE = HISTORY_DIR / "mail_history.json"
HISTORY_OUTPUT = OUTPUT_ROOT / "workbench_history"
DATABASE_PATH = APP_ROOT / "storage" / "workbench.db"
WORKORDER_RETRY_QUEUE = APP_ROOT / "storage" / "workorder_retry_queue.json"
COMPLETED_HISTORY_XLSX = HISTORY_OUTPUT / "邮件处理完成总表.xlsx"
UNFINISHED_HISTORY_XLSX = HISTORY_OUTPUT / "邮件未完成总表.xlsx"
HTML_PATH = APP_ROOT / "workbench.html"
SESSION_STATE = APP_ROOT / "session_state.json"
CONFIG_PATH = APP_ROOT / "config.yaml"

# 过滤日志里已经"定论"的过滤原因：证书/下号通知是规则确定的交付通知，
# 非注册业务(收款/报价/续费/信息变更/发票)是词表确定的运营事务，
# LLM确认=... 是二次识别已判定为非注册邮件。它们只参与右上角计数，
# 不再进入复核队列；其余过滤原因（可能误过滤）仍保留给人工判断。
SETTLED_FILTER_HINTS = ("证书/下号通知", "非注册业务", "LLM确认")


def _is_settled_filter(reason: str) -> bool:
    text = _text(reason)
    return any(hint in text for hint in SETTLED_FILTER_HINTS)


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _json_list(value: Any) -> List[dict]:
    """读取阶段一写入的附件证据 JSON；旧文件/手工 Excel 为空时兼容。"""
    if isinstance(value, list):
        return [item for item in value if isinstance(item, dict)]
    text = _text(value)
    if not text:
        return []
    try:
        parsed = json.loads(text)
    except (TypeError, ValueError, json.JSONDecodeError):
        # 旧版为了适配 Excel 单元格上限把 JSON 从中间截断。不能把整列
        # 直接判成“无附件”：从截断前仍完整的 filename/row 对象恢复一份
        # 可预览的结构化快照，原始文件不存在时也能生成可打开的预览表。
        return _recover_truncated_attachment_evidence(text)
    return [item for item in parsed if isinstance(item, dict)] if isinstance(parsed, list) else []


def _json_object(value: Any) -> Dict[str, Any]:
    """读取阶段一/阶段二写入的专项核对 JSON，失败时返回空对象。"""
    if isinstance(value, dict):
        return value
    text = _text(value)
    if not text:
        return {}
    try:
        parsed = json.loads(text)
    except (TypeError, ValueError, json.JSONDecodeError):
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _recover_truncated_attachment_evidence(text: str) -> List[dict]:
    """尽可能恢复被 Excel 单元格长度截断的附件证据。

    恢复目标是证据展示，不把不完整 JSON 当成完整数据：只接受
    ``json.JSONDecoder.raw_decode`` 能完整解析的附件名、工作表行和文本片段。
    这样旧数据不会再静默变成空附件，新数据仍由正常 json.loads 处理。
    """
    decoder = json.JSONDecoder()
    filename_matches = list(re.finditer(r'"filename"\s*:', text))
    if not filename_matches:
        return []
    recovered: List[dict] = []
    seen_names = set()
    for index, match in enumerate(filename_matches):
        start = match.end()
        while start < len(text) and text[start].isspace():
            start += 1
        try:
            filename, _ = decoder.raw_decode(text, start)
        except (TypeError, ValueError, json.JSONDecodeError):
            continue
        filename = _text(filename)
        if not filename:
            continue
        end = filename_matches[index + 1].start() if index + 1 < len(filename_matches) else len(text)
        segment = text[start:end]
        rows: List[dict] = []
        seen_rows = set()
        for row_match in re.finditer(r'\{"row_number"\s*:', segment):
            try:
                row, _ = decoder.raw_decode(segment, row_match.start())
            except (TypeError, ValueError, json.JSONDecodeError):
                continue
            if not isinstance(row, dict):
                continue
            key = (_text(row.get("row_number")), json.dumps(row.get("cells") or [], ensure_ascii=False))
            if key in seen_rows:
                continue
            seen_rows.add(key)
            rows.append({
                "row_number": row.get("row_number") or "?",
                "cells": [_text(cell)[:160] for cell in (row.get("cells") or [])[:24]],
            })
        sheet_names = []
        for sheet_match in re.finditer(r'"sheet_name"\s*:', segment):
            pos = sheet_match.end()
            while pos < len(segment) and segment[pos].isspace():
                pos += 1
            try:
                sheet_name, _ = decoder.raw_decode(segment, pos)
            except (TypeError, ValueError, json.JSONDecodeError):
                continue
            if _text(sheet_name) and _text(sheet_name) not in sheet_names:
                sheet_names.append(_text(sheet_name))
        text_preview = ""
        text_match = re.search(r'"text"\s*:', segment)
        if text_match:
            pos = text_match.end()
            while pos < len(segment) and segment[pos].isspace():
                pos += 1
            try:
                text_value, _ = decoder.raw_decode(segment, pos)
                text_preview = _text(text_value)[:1200]
            except (TypeError, ValueError, json.JSONDecodeError):
                pass
        item: Dict[str, Any] = {"filename": filename, "records": [], "sheets": [], "text": text_preview}
        if rows:
            item["sheets"] = [{"sheet_name": sheet_names[0] if sheet_names else "工作表", "rows": rows}]
        if filename not in seen_names:
            seen_names.add(filename)
            recovered.append(item)
    return recovered


def _looks_like_company_name(value: Any) -> bool:
    """仅接受带法定后缀的附件公司候选，拒绝城市/提示语。"""
    text = _text(value)
    if not text or _is_non_company_customer_value(text):
        return False
    if re.search(r"(?:有限责任公司|有限公司|股份有限公司|集团公司|合伙企业|公司|（个体工商户）|\(个体工商户\)|经营部|门市部|服务部|商店|商行|工厂|工作室|店)$", text):
        return len(text) >= 4
    return bool(re.search(
        r"(?i)(?:^|\s)(?:limited|ltd\.?|llc|gmbh|s\.?\s*p\.?\s*z\.?\s*o\.?\s*o\.?|"
        r"sarl|sas|b\.?v\.?|a\.?b\.?|s\.?l\.?|oy|a\.?/\.?s\.?|plc|inc\.?|corp\.?|company)"
        r"(?:$|\s|[,.)])",
        text,
    ))


def _subject_loose_company(subject: Any) -> str:
    """从旧阶段一记录的主题中恢复无公司后缀的英文商号。

    只处理“案件编号 + 英文商号 + 国家/项目”这一种明确格式，避免把
    中文正文中的说明句重新猜成公司。该兜底用于读取旧 Excel，新的阶段一
    结果仍以 field_extractor 的完整证据链为准。
    """
    text = _text(subject)
    code = re.search(
        r"(?<![A-Za-z0-9])(?:[A-Za-z]{1,3}\s*[-_]\s*)?"
        r"[A-Za-z]{1,8}\s*[-_]?\s*\d{2,8}(?![A-Za-z0-9])",
        text,
        re.I,
    )
    if not code:
        return ""
    tail = text[code.end():].strip(" \t-—_:：+＋")
    boundary = re.search(
        r"(?:德国|比利时|比利時|法国|法國|意大利|義大利|西班牙|荷兰|荷蘭|"
        r"波兰|波蘭|瑞典|爱尔兰|愛爾蘭|葡萄牙|奥地利|奧地利|芬兰|芬蘭|"
        r"WEEE|EEE|EPR|电池法|電池法|包装法|包裝法|注册|註冊|新注册|新註冊|"
        r"申报|申報|注销|註銷|修改|变更|變更)",
        tail,
        re.I,
    )
    candidate = (tail[:boundary.start()] if boundary else tail).strip(" \t-—_:：+＋,，;；")
    english = re.fullmatch(r"[A-Za-z][A-Za-z0-9 .,&'’()\-]{3,80}", candidate)
    chinese = re.fullmatch(
        r"[\u4e00-\u9fffA-Za-z0-9·（）()]{4,60}"
        r"(?:经营部|门市部|服务部|商店|商行|工厂|工作室|店|有限公司|公司|（个体工商户）|\(个体工商户\))",
        candidate,
    )
    if not english and not chinese:
        return ""
    if english and len(candidate.split()) < 2:
        return ""
    if _is_non_company_customer_value(candidate):
        return ""
    return candidate


def _repair_attachment_fields(row: Dict[str, Any]) -> Dict[str, Any]:
    """从附件证据行修复旧版“城市冒充公司”的阶段一结果。

    早期导出曾把注册表的 City 列写进客户公司字段。附件证据保留了
    ``代理 | 编号 | 公司中文名`` 的原始顺序，因此无需重新下载邮件即可
    修复已有工作台数据；人工状态仍由后续 state 覆盖。
    """
    if not isinstance(row, dict):
        return row
    current_company = _text(row.get("客户公司名称") or row.get("客户"))
    # 即使旧行没有附件证据，也能安全拆分“编号 + 法定英文名称”。
    code_prefix = re.match(
        r"^\s*(?:[A-Za-z]{1,3}\s*[-_]\s*)?[A-Za-z]{1,8}"
        r"\s*[-_]?\s*\d{2,8}\s+(.+?)\s*$",
        current_company,
    )
    if code_prefix and _looks_like_company_name(code_prefix.group(1)):
        cleaned_company = _text(code_prefix.group(1)).strip(" .,-")
        row["客户公司名称"] = cleaned_company
        row["客户"] = cleaned_company
        row["客户提取来源"] = "公司编号与法定名称拆分"
        current_company = cleaned_company
    # 旧记录中“不能与其它公司”“国公司”等确定性说明值可能遮住了主题中
    # 明确的英文商号。能按案件编号和国家/项目边界恢复时直接修复；否则保留
    # 空值/人工复核，不再把说明文字显示成公司。
    if _is_non_company_customer_value(current_company):
        fallback = _subject_loose_company(
            row.get("邮件主题") or row.get("主题") or row.get("subject")
        )
        if fallback:
            row["客户公司名称"] = fallback
            row["客户"] = fallback
            row["客户提取来源"] = "主题案件编号边界恢复"
            current_company = fallback
    evidence = _json_list(row.get("附件证据"))
    if not evidence:
        return row
    source = _text(row.get("附件明细来源"))
    code = _text(row.get("客户编号"))
    company_candidates: List[str] = []
    agent_candidate = ""
    for attachment in evidence:
        for record in attachment.get("records") or []:
            if not isinstance(record, dict):
                continue
            filename = _text(record.get("attachment_name") or attachment.get("filename"))
            sheet = _text(record.get("sheet_name") or "工作表")
            number = _text(record.get("row_number") or "?")
            locator = f"附件表格：{filename} / {sheet} 第{number}行"
            if source and locator not in source:
                continue
            raw = _text(record.get("raw_text"))
            parts = [_text(part) for part in raw.split("|") if _text(part)]
            if code and parts:
                code_index = next(
                    (index for index, part in enumerate(parts) if part == code), -1
                )
                if code_index >= 0:
                    if code_index > 0:
                        before = parts[code_index - 1]
                        if before and not re.search(r"\d", before) and len(before) <= 40:
                            agent_candidate = before
                    if code_index + 1 < len(parts):
                        after = parts[code_index + 1]
                        if _looks_like_company_name(after):
                            company_candidates.append(after)
            for part in parts:
                if _looks_like_company_name(part):
                    company_candidates.append(part)
    company_candidates = list(dict.fromkeys(company_candidates))
    replacement = next(
        (candidate for candidate in company_candidates if candidate != current_company),
        "",
    )
    if replacement and not _looks_like_company_name(current_company):
        row["客户公司名称"] = replacement
        row["客户"] = replacement
        row["客户提取来源"] = "附件表格中文公司名"
    if agent_candidate and not _text(row.get("代理")):
        row["代理"] = agent_candidate
        row["代理(可能空)"] = agent_candidate
        row["代理匹配方式"] = "附件表格代理列"
    return row


def _is_non_company_customer_value(value: Any) -> bool:
    """判断历史导入的客户值是否是 EPR 表单标签而非主体名称。

    旧版本曾把 EPR 申请表中的“POA/注册资本/签字时间”等字段标签写成
    客户。这里仅在工作台读取时隐藏这些确定性脏记录，不删除数据库历史；
    重新运行阶段一后，正确的客户记录会正常进入队列。
    """
    text = _text(value)
    if not text:
        return False
    compact = re.sub(r"\s+", "", text).lower()
    hints = (
        "poa", "legalrepresentative", "legalperson", "legalpositions",
        "nameoflegalperson", "placeofsignature", "signingtime",
        "registrationcapital", "companyname", "companyaddress", "companybusiness",
        "plz", "postcode", "amazonlink", "shoplink", "e-mail", "email", "tel", "phone",
        "legrepresentativename", "companyregistrationnumber", "registrationnumber", "uscc",
        "营业执照", "公司名称", "公司地址", "公司注册", "注册资本", "法人",
        "公司成立日期", "成立日期", "签字", "签署", "职位", "联系信息", "联系人", "联系电话", "邮箱",
        "邮政编码", "邮编", "地址", "姓名", "身份证", "护照", "性别", "店铺链接",
        "平台信息", "服务内容", "服务的国家", "销售量", "预计销售", "说明", "备注", "请提供", "请选择", "填写",
        "注意事项", "不能提供", "提供", "请客户", "确认好", "注册类别", "翻译公司", "盖章",
        "回收公司", "适用国家", "所有国家", "产品图片或说明书", "资料列表", "不用提供",
        "保证有就可以", "否则不接单", "要求北爱公司", "非中国公司", "中国公司",
        "外国公司", "国公司", "不能与其它公司", "不能与其他公司",
    )
    if any(hint in compact for hint in hints):
        return True
    # 旧版阶段一会把正文引导语中的公司后缀截出来，例如
    # “以下为广东省方信企业管理集团有限公司……提交……名单”。
    # 这不是客户主体；重新解析前先从工作台显示层隐藏这类历史脏行，
    # 但不删除数据库中的原始记录，便于审计追溯。
    if re.match(r"^(?:以下|下面|下列|现将|本次)(?:为|是)?", text, re.I) and re.search(
        r"(?:提交|报送|发送|提供|列出|名单|新注册|申请)", text, re.I
    ):
        return True
    # 旧行里通常只剩公司后缀前的截断值（如“以下为某某有限公司”），
    # 后面的“提交名单”已经不在客户字段中，因此单独按前缀+公司后缀识别。
    if re.match(r"^(?:以下|下面|下列)(?:为|是)", text, re.I) and re.search(
        r"(?:有限责任公司|股份有限公司|集团有限公司|有限公司|责任公司|公司|企业)$",
        text,
        re.I,
    ):
        return True
    if "@" in text or re.search(r"https?://|www\.", text, re.I):
        return True
    if re.fullmatch(r"[+()\-\s\d]{6,}", text):
        return True
    if re.fullmatch(r"[A-Za-z]{1,8}[-_]?\d{4,}", text):
        return True
    if re.fullmatch(r"[0-9一二三四五六七八九十多几]*\s*家\s*(?:公司|主体|企业)", text, re.I):
        return True
    if re.fullmatch(
        r"[0-9一二三四五六七八九十多几]+\s*家\s*(?:公司|主体|企业)"
        r"(?:\s*[-—:：,，/／+＋].*)?",
        text,
        re.I,
    ):
        return True
    # 旧版表单解析会把法人姓名写进客户列（如 Huiming Wu）。没有公司后缀、
    # 仅由英文名组成的值没有足够主体证据，工作台不应继续展示为客户公司。
    if re.fullmatch(r"[A-Z][a-z]{1,24}(?:\s+[A-Z][a-z]{1,24}){1,3}", text):
        return True
    return len(text) > 120


def _mail_identity_tuple(row: Dict[str, Any]) -> Tuple[str, str, str]:
    """返回用于替换同一封邮件旧版本记录的稳定身份。"""
    sender = _text(row.get("发件人邮箱") or row.get("sender_email") or row.get("sender")).lower()
    date = _text(row.get("发件日期") or row.get("date")).replace("T", " ")
    subject = _text(row.get("邮件主题") or row.get("主题") or row.get("subject"))
    return sender, date, subject


def _workorder_detail_tuple(row: Dict[str, Any]) -> Tuple[str, str, str, str, str, str]:
    """阶段二结果与人工回写结果的稳定业务明细身份。"""
    mail = _mail_identity_tuple(row)
    company = _text(row.get("客户公司名称") or row.get("客户") or row.get("company"))
    project = _text(row.get("标准化项目名称") or row.get("项目") or row.get("program"))
    request = _text(row.get("需求") or row.get("request"))
    return (*mail, company, project, request)


def _obvious_business_issues(row: Dict[str, Any]) -> List[dict]:
    """Return only high-confidence company-name anomalies for legacy rows.

    Stage-one workbooks have used both internal keys (客户/项目/需求) and
    exported headers (客户公司名称/标准化项目名称/需求).  Passing an old
    exported row directly to ``inspect_row`` makes valid data look empty and
    incorrectly changes its status to 待人工复核.  We still re-check known bad
    company patterns, but do not infer missing values from header differences.
    """
    try:
        from modules.business_validator import inspect_row

        normalized = {
            "客户": _text(row.get("客户") or row.get("客户公司名称")),
            "项目": _text(row.get("项目") or row.get("标准化项目名称")),
            "需求": _text(row.get("需求")),
        }
        if not normalized["客户"]:
            return []
        issues = inspect_row(normalized)
        return [
            issue for issue in issues
            if _text(issue.get("code")).startswith("COMPANY_")
        ]
    except Exception:
        return []


def _safe_path(path: Optional[str], fallback: Path) -> Path:
    # 兼容旧版 session/config 中残留的开发机绝对路径。显式传入但尚未
    # 生成的临时路径必须原样保留（测试/用户刚选择的空结果文件不能被
    # fallback 偷换成 output 里的历史文件）；只有路径为空时才使用默认值。
    if not path:
        return resolve_runtime_path(None, fallback)
    candidate = Path(str(path)).expanduser()
    if not candidate.is_absolute():
        candidate = APP_ROOT / candidate
    try:
        if candidate.is_file() or candidate.is_dir():
            return candidate.resolve()
    except OSError:
        pass
    if candidate.is_absolute() and candidate.name:
        for sibling in (
            APP_ROOT / "data" / candidate.name,
            APP_ROOT / "storage" / "imported" / candidate.name,
            APP_ROOT / "output" / candidate.name,
            APP_ROOT / candidate.name,
        ):
            try:
                if sibling.exists():
                    return sibling.resolve()
            except OSError:
                continue
    return candidate.resolve()


def _stable_id(source: str, row_number: int, row: Dict[str, Any]) -> str:
    raw = "|".join(
        [
            source,
            str(row_number),
            _text(row.get("发件人邮箱")),
            _text(row.get("发件日期")),
            _text(row.get("邮件主题")),
            _text(row.get("客户公司名称")),
            _text(row.get("标准化项目名称")),
        ]
    )
    return hashlib.sha1(raw.encode("utf-8", errors="ignore")).hexdigest()[:16]


def _read_sheet(path: Path, sheet_name: str, source: str) -> List[Dict[str, Any]]:
    if not path.exists():
        return []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception:
        return []
    if not rows:
        return []
    headers = [_text(v) for v in rows[0]]
    result: List[Dict[str, Any]] = []
    for row_number, values in enumerate(rows[1:], start=2):
        if not any(v not in (None, "") for v in values):
            continue
        item = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers)) if headers[i]}
        item["_source"] = source
        item["_row_number"] = row_number
        item["_id"] = _stable_id(source, row_number, item)
        result.append(item)
    return result


def _load_json(path: Path) -> Dict[str, Any]:
    try:
        with path.open("r", encoding="utf-8") as fh:
            value = json.load(fh)
        return value if isinstance(value, dict) else {}
    except (FileNotFoundError, json.JSONDecodeError, OSError):
        return {}


def _save_json(path: Path, value: Dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with tmp.open("w", encoding="utf-8") as fh:
        json.dump(value, fh, ensure_ascii=False, indent=2)
    os.replace(tmp, path)


def _mail_date_sort_value(value: Any) -> datetime:
    """把邮件日期转为排序键；无法解析的旧数据排在最后。"""
    text = _text(value).replace("T", " ")
    if not text:
        return datetime.min
    for candidate in (text, text[:19], text[:10]):
        try:
            parsed = datetime.fromisoformat(candidate)
            # 邮件头有时带 +00:00，有时是本地时间；排序键统一为无时区
            # 的 UTC 值，避免历史/新数据混排时触发 naive/aware 比较异常。
            if parsed.tzinfo is not None:
                parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
            return parsed
        except ValueError:
            continue
    return datetime.min


def _history_result(status: str) -> str:
    return {
        "confirmed": "处理完成",
        "filtered": "已过滤归档",
        "workorder_completed": "工单核对完成",
        "review": "待人工复核",
        "returned": "已退回复核",
        "needs_info": "待补资料",
        "ready": "待确认",
    }.get(_text(status), "待处理")


def _history_is_complete(status: str) -> bool:
    """人工完成、过滤归档及阶段二已出明确结论的邮件进入完成总表。"""
    return _text(status) in {"confirmed", "filtered", "workorder_completed"}


def _event_signature(detail_id: str, event: Dict[str, Any]) -> str:
    return "|".join([
        _text(detail_id), _text(event.get("at")), _text(event.get("action")),
        _text(event.get("reason")),
    ])


# ---- 项目名称表：给「新增项目」提供可选项目，避免人工手输造成阶段二匹配不上 ----

_PROJECT_CACHE: Dict[str, Tuple[float, List[Dict[str, str]]]] = {}
_PROJECT_LOCK = threading.Lock()


def _config_project_names() -> Optional[str]:
    """从 config.yaml 取项目名称表路径（不引入 YAML 依赖，只扫这一行）。"""
    try:
        text = CONFIG_PATH.read_text(encoding="utf-8", errors="ignore")
    except OSError:
        return None
    match = re.search(r"^\s*project_names:\s*(\S+)\s*$", text, re.MULTILINE)
    return match.group(1).strip().strip("'\"") if match else None


def project_table_path() -> Optional[Path]:
    """项目名称表位置：会话状态（GUI 手选） > config.yaml > data/project_names.xlsx。"""
    session = _load_json(SESSION_STATE)
    for candidate in (session.get("project_table_path"), _config_project_names(), "data/project_names.xlsx"):
        if not candidate:
            continue
        path = _safe_path(str(candidate), Path("data") / "project_names.xlsx")
        try:
            if path.is_file():
                return path
        except OSError:
            continue
    return None


def load_project_names() -> List[Dict[str, str]]:
    """读取项目名称表；按文件 mtime 缓存，避免每次刷新都读盘。"""
    path = project_table_path()
    if not path:
        return []
    try:
        mtime = path.stat().st_mtime
    except OSError:
        return []
    key = str(path)
    with _PROJECT_LOCK:
        cached = _PROJECT_CACHE.get(key)
        if cached and cached[0] == mtime:
            return cached[1]
    result: List[Dict[str, str]] = []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = [r for r in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
        wb.close()
    except Exception:
        return []
    if not rows:
        return []
    headers = [_text(v) for v in rows[0]]
    if "项目名称" in headers:
        i_name: Optional[int] = headers.index("项目名称")
    else:  # 没有表头时退化为第二列
        i_name = 1
    i_number = headers.index("项目编号") if "项目编号" in headers else None
    i_country = headers.index("国家") if "国家" in headers else None
    i_business = headers.index("业务类型") if "业务类型" in headers else None

    def cell(row: Iterable[Any], index: Optional[int]) -> str:
        row = list(row)
        return _text(row[index]) if index is not None and index < len(row) else ""

    for row in rows[1:]:
        name = cell(row, i_name)
        if not name:
            continue
        result.append({
            "name": name,
            "number": cell(row, i_number),
            "country": cell(row, i_country) or country_of_project(name),
            "business": cell(row, i_business),
        })
    with _PROJECT_LOCK:
        _PROJECT_CACHE[key] = (mtime, result)
    return result


class WorkbenchStore:
    """读取阶段一产物并叠加人工复核状态。"""

    def __init__(
        self,
        primary_path: Optional[str] = None,
        review_path: Optional[str] = None,
        filtered_path: Optional[str] = None,
        workorder_result_path: Optional[str] = None,
        state_path: Optional[str] = None,
        database_path: Optional[str] = None,
        test_mode: bool = False,
    ):
        primary_fallback = DEFAULT_PRIMARY if DEFAULT_PRIMARY.exists() else LEGACY_PRIMARY
        review_fallback = DEFAULT_REVIEW if DEFAULT_REVIEW.exists() else LEGACY_REVIEW
        filtered_fallback = DEFAULT_FILTERED if DEFAULT_FILTERED.exists() else LEGACY_FILTERED
        workorder_fallback = (
            DEFAULT_WORKORDER_RESULT
            if DEFAULT_WORKORDER_RESULT.exists()
            else LEGACY_WORKORDER_RESULT
        )
        self.primary_path = _safe_path(primary_path, primary_fallback)
        self.review_path = _safe_path(review_path, review_fallback)
        self.filtered_path = _safe_path(filtered_path, filtered_fallback)
        self.workorder_result_path = _safe_path(workorder_result_path, workorder_fallback)
        self.test_mode = bool(test_mode)
        if state_path:
            self.state_path = _safe_path(state_path, REVIEW_STATE)
        elif self.test_mode:
            # 测试会话不加载旧的“退回/确认/修改”等人工状态，也不覆盖正式状态。
            # 同一服务存活期间仍可刷新页面，继续当前测试。
            session_dir = APP_ROOT / "storage" / "workbench_test_sessions"
            token = f"{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:8]}"
            self.state_path = session_dir / f"workbench_test_{token}.json"
        else:
            self.state_path = REVIEW_STATE
        # 自动历史只属于正式工作台。单元测试/临时状态文件不能污染真实邮件台账。
        self.history_enabled = (
            not self.test_mode
            and self.state_path.resolve() == REVIEW_STATE.resolve()
        )
        # 正式工作台才启用共享数据库；测试/临时会话继续完全隔离，避免把测试数据
        # 写进正式台账。数据库默认放在当前 APP_ROOT 下，也支持独立启动时传入路径。
        self.database: Optional[WorkbenchDatabase] = None
        if self.history_enabled:
            db_fallback = self.state_path.parent / "workbench.db"
            self.database = WorkbenchDatabase(_safe_path(database_path, db_fallback))
        self._lock = threading.RLock()

    def _raw_records(self) -> List[Dict[str, Any]]:
        records = _read_sheet(self.primary_path, "工单待查", "待查名单")
        # 当前选择的阶段一输出是主数据；复查表是另一条待补全队列。
        records.extend(_read_sheet(self.review_path, "漏单复查", "人工补全"))
        # 兼容旧版阶段一：附件证据中的“代理 | 编号 | 公司中文名”
        # 可修复城市冒充公司、代理为空等确定性错误，再写入持久数据库。
        records = [_repair_attachment_fields(dict(row)) for row in records]
        if self.database is not None:
            # 先把本次阶段一输出增量写入数据库。数据库保存完整历史，
            # 但当前队列不能把同一封邮件在历次解析中产生的旧版本一起展示，
            # 否则一次修复会变成“7 个主体/32 条明细”这种重复污染。
            self.database.ingest_rows(
                records,
                dataset="active",
                source_path=f"{self.primary_path};{self.review_path}",
            )
            stored = [
                _repair_attachment_fields(dict(row))
                for row in self.database.read_rows("active")
            ]
            # 数据库中可能留有旧版本误抽取的表单字段。保留它们用于历史
            # 追溯，但不再把确定性脏值显示为业务明细或带入工单核对。
            stored = [
                row for row in stored
                if not _is_non_company_customer_value(
                    row.get("客户公司名称") or row.get("客户") or row.get("company")
                )
            ]
            if records:
                fresh = [
                    row for row in records
                    if not _is_non_company_customer_value(
                        row.get("客户公司名称") or row.get("客户") or row.get("company")
                    )
                ]
                fresh_by_mail: Dict[Tuple[str, str, str], List[Dict[str, Any]]] = {}
                stored_by_mail: Dict[Tuple[str, str, str], List[Dict[str, Any]]] = {}
                for row in fresh:
                    fresh_by_mail.setdefault(_mail_identity_tuple(row), []).append(row)
                for row in stored:
                    stored_by_mail.setdefault(_mail_identity_tuple(row), []).append(row)

                # 同一封邮件优先使用本次输出；如果本次输出全是旧版无法
                # 识别的空主体占位行，则回退到数据库中最近保存的合法主体。
                # 这让用户无需先清库就能看到修复后的真实两条项目。
                selected: List[Dict[str, Any]] = []
                fresh_mail_keys = set(fresh_by_mail)
                stable_by_legacy: Dict[str, str] = {}
                for row in stored:
                    legacy = _text(row.get("_legacy_id"))
                    stable = _text(row.get("_db_record_key"))
                    if legacy and stable:
                        stable_by_legacy.setdefault(legacy, stable)
                for mail_key in dict.fromkeys([*fresh_by_mail.keys(), *stored_by_mail.keys()]):
                    fresh_rows = fresh_by_mail.get(mail_key, [])
                    stored_rows = stored_by_mail.get(mail_key, [])
                    fresh_with_company = [
                        row for row in fresh_rows
                        if _text(row.get("客户公司名称") or row.get("客户") or row.get("company"))
                    ]
                    stored_with_company = [
                        row for row in stored_rows
                        if _text(row.get("客户公司名称") or row.get("客户") or row.get("company"))
                    ]
                    if fresh_with_company:
                        chosen = fresh_with_company
                    elif stored_with_company:
                        chosen = stored_with_company
                    else:
                        chosen = fresh_rows or stored_rows

                    # 将同一行映射回数据库稳定 ID，保证之前已保存的人工
                    # 修改仍能叠加到本次新输出；新行则保留本次原始 ID。
                    for row in chosen:
                        if row in fresh_rows:
                            legacy = _text(row.get("_id"))
                            stable = stable_by_legacy.get(legacy)
                            if stable:
                                row["_legacy_id"] = legacy
                                row["_id"] = stable
                    selected.extend(chosen)

                # 本次文件未覆盖的历史邮件仍保留在工作台，可按邮件日期
                # 继续处理；同一邮件只保留一个版本，不把当前邮件的旧版本
                # 再追加一次。
                if stored:
                    selected_keys = {
                        _text(row.get("_db_record_key")) or _text(row.get("_id"))
                        for row in selected
                    }
                    for row in stored:
                        if _mail_identity_tuple(row) in fresh_mail_keys:
                            continue
                        key = _text(row.get("_db_record_key")) or _text(row.get("_id"))
                        if key and key not in selected_keys:
                            selected.append(row)
                return selected
            if stored:
                return stored
        return records

    def _state(self) -> Dict[str, Any]:
        return _load_json(self.state_path)

    def _status(self, row: Dict[str, Any], saved: Dict[str, Any]) -> str:
        if saved.get("status") in {"review", "ready", "confirmed", "returned", "needs_info"}:
            return saved["status"]
        # 对旧的阶段一输出也重新执行一次确定性业务闸门，避免工作台继续把
        # “在其他欧盟国家或第三国设立的公司”等申请表说明句显示为可确认。
        if _obvious_business_issues(row):
            return "review"
        source = row.get("_source")
        confidence = _text(row.get("置信度")).lower()
        needs = _text(row.get("人工复核提示")) or _text(row.get("待人工补全"))
        semantic = _text(row.get("语义校验状态")).lower()
        attachment_count_check = _text(row.get("附件表格数量校验"))
        if source == "人工补全" or needs or attachment_count_check == "需人工确认" or semantic in {"uncertain", "invalid", "llm调用失败/未完成"} or confidence in {"low", "medium"} or not _text(row.get("代理")):
            return "review"
        return "ready"

    @staticmethod
    def _base_fields(row: Dict[str, Any]) -> Dict[str, str]:
        project = _text(row.get("标准化项目名称"))
        # 国家必须从项目名里解析出「国家」本身，不能按空格硬切
        # （"奥地利WEEE" 无空格，旧写法会让国家列显示成 "奥地利WEEE"）。
        country = country_of_project(project)
        return {
            "agent": _text(row.get("代理")) or _text(row.get("代理(可能空)")),
            "customer_code": _text(row.get("客户编号")),
            "company": _text(row.get("客户公司名称")),
            "country": country,
            "program": project,
            "request": _text(row.get("需求")),
        }

    def _record(self, row: Dict[str, Any], state: Dict[str, Any]) -> Dict[str, Any]:
        rid = row["_id"]
        saved_records = state.get("records", {}) if isinstance(state.get("records"), dict) else {}
        saved = saved_records.get(rid, {})
        # 数据库首次接管已有 Excel 时，使用旧行号生成的 ID 迁移历史人工状态，
        # 后续新操作则统一写入数据库稳定 ID。
        if not isinstance(saved, dict) or not saved:
            legacy_id = _text(row.get("_legacy_id"))
            saved = saved_records.get(legacy_id, {}) if legacy_id else {}
        if not isinstance(saved, dict):
            saved = {}
        fields = self._base_fields(row)
        fields.update(saved.get("fields", {}) if isinstance(saved.get("fields"), dict) else {})
        status = self._status(row, saved)
        business_issues = _obvious_business_issues(row)
        business_hint = "；".join(
            f"{item.get('field', '')}: {item.get('reason', '')}"
            for item in business_issues
            if item.get("reason")
        )
        evidence = [
            {"field": "customer_code", "label": "客户编号", "initial": self._base_fields(row)["customer_code"], "accepted": fields["customer_code"], "source": "主题/正文中的编号实体"},
            {"field": "company", "label": "客户公司", "initial": self._base_fields(row)["company"], "accepted": fields["company"], "source": _text(row.get("客户提取来源")) or "邮件正文"},
            {"field": "agent", "label": "代理", "initial": self._base_fields(row)["agent"], "accepted": fields["agent"], "source": _text(row.get("代理匹配方式")) or "待确认"},
            {"field": "program", "label": "服务项目", "initial": self._base_fields(row)["program"], "accepted": fields["program"], "source": _text(row.get("项目原始值")) or "正文/附件"},
            {"field": "request", "label": "需求", "initial": self._base_fields(row)["request"], "accepted": fields["request"], "source": "正文/主题"},
        ]
        attachment_expected = _text(row.get("附件表格记录数"))
        attachment_output = _text(row.get("附件表格输出数"))
        attachment_check = _text(row.get("附件表格数量校验"))
        attachment_source = _text(row.get("附件明细来源"))
        attachment_evidence = _json_list(row.get("附件证据"))
        attachment_files = _json_list(row.get("附件文件索引"))
        weee_items = _json_list(row.get("德国WEEE品类明细"))
        weee_check = _json_object(row.get("德国WEEE品类核对"))
        weee_enabled = _text(row.get("德国WEEE专项")) == "是"
        if attachment_expected or attachment_output or attachment_check or attachment_source:
            evidence.append({
                "field": "attachment_record_count",
                "label": "附件表格数量",
                "initial": " / ".join(value for value in (attachment_expected, attachment_output) if value) or "（未记录）",
                "accepted": attachment_check or "（未校验）",
                "source": attachment_source or "附件结构化解析",
            })
        if weee_enabled:
            evidence.append({
                "field": "weee_category",
                "label": "德国 WEEE 品类",
                "initial": "；".join(
                    f"{_text(item.get('brand')) + ' / ' if _text(item.get('brand')) else ''}{_text(item.get('category'))}"
                    for item in weee_items if _text(item.get("category")) or _text(item.get("brand"))
                ) or "（未提取到明确品类）",
                "accepted": _text(row.get("德国WEEE品类状态")) or "待工单核对",
                "source": "产品分类表 + 邮件/附件证据",
            })
        return {
            "id": rid,
            "source": row.get("_source", ""),
            "row_number": row.get("_row_number"),
            "status": status,
            "sender": _text(row.get("发件人邮箱")),
            "recipient": _text(row.get("收件人")) or _text(row.get("收件人邮箱")),
            "date": _text(row.get("发件日期")),
            "operation_date": _text(row.get("_db_operation_date")) or _text(row.get("发件日期"))[:10],
            "subject": _text(row.get("邮件主题")),
            # 新版阶段一保留“邮件正文原文”；旧工作簿没有该列时兼容摘要。
            "body": _text(row.get("邮件正文原文")) or _text(row.get("邮件正文摘要(最多300字)")),
            "attachments": _text(row.get("附件名称")),
            "attachment_files": attachment_files,
            "fields": fields,
            "evidence": evidence,
            "confidence": _text(row.get("置信度")) or "unknown",
            "agent_match": _text(row.get("代理匹配方式")) or "待确认",
            "extraction": _text(row.get("智能提取方式")) or "规则提取",
            "review_hint": (
                _text(row.get("人工复核提示"))
                or _text(row.get("待人工补全"))
                or (f"确定性业务校验: {business_hint}" if business_hint else "")
            ),
            "attachment_audit": {
                "expected": attachment_expected,
                "output": attachment_output,
                "status": attachment_check,
                "source": attachment_source,
            },
            "attachment_evidence": attachment_evidence,
            "weee": {
                "enabled": weee_enabled,
                "items": weee_items,
                "status": _text(row.get("德国WEEE品类状态")) or "待工单核对",
                "comparison": weee_check,
                "reason": _text(row.get("德国WEEE专项说明")),
            },
            "semantic": {
                "status": _text(row.get("语义校验状态")),
                "issue_codes": _text(row.get("语义问题编号")),
                "issues": _text(row.get("语义问题字段")),
                "evidence": _text(row.get("语义问题证据")),
                "current_values": _text(row.get("语义当前值")),
                "suggested_values": _text(row.get("语义建议值")),
                "suggestion": _text(row.get("语义校验建议")),
                "reason": _text(row.get("语义校验原因")),
                "model": _text(row.get("语义校验模型")),
            },
            "data_source": _text(row.get("数据来源")),
            "initial_row": {k: _text(v) for k, v in row.items() if not k.startswith("_")},
            "events": saved.get("events", []) if isinstance(saved.get("events"), list) else [],
        }

    @staticmethod
    def _deleted_map(state: Dict[str, Any]) -> Dict[str, Any]:
        deleted = state.get("deleted")
        return deleted if isinstance(deleted, dict) else {}

    @staticmethod
    def _added_map(state: Dict[str, Any]) -> Dict[str, Any]:
        added = state.get("added")
        return added if isinstance(added, dict) else {}

    def _added_record(self, item: Dict[str, Any], mail: Dict[str, Any], state: Dict[str, Any]) -> Dict[str, Any]:
        """把「人工新增的项目」包装成与阶段一明细同构的记录。

        人工新增的明细没有原始 Excel 行，因此 initial_row 由录入字段拼出，
        导出时能直接写回"工单待查"表头，阶段二仍可识别。
        """
        rid = _text(item.get("id"))
        saved = state.get("records", {}).get(rid, {}) if isinstance(state.get("records"), dict) else {}
        raw = item.get("fields") if isinstance(item.get("fields"), dict) else {}
        fields = {k: _text(raw.get(k)) for k in ("agent", "customer_code", "company", "country", "program", "request")}
        saved_fields = saved.get("fields") if isinstance(saved.get("fields"), dict) else {}
        fields.update({k: _text(v) for k, v in saved_fields.items() if k in fields})
        if not fields["country"]:
            fields["country"] = country_of_project(fields["program"])
        status = _text(saved.get("status")) or "ready"
        note = _text(item.get("note"))
        created = _text(item.get("created_at"))
        evidence = [
            {"field": "company", "label": "客户公司", "initial": fields["company"], "accepted": fields["company"], "source": "人工新增（原邮件人工判读）"},
            {"field": "agent", "label": "代理", "initial": fields["agent"], "accepted": fields["agent"], "source": "人工新增（原邮件人工判读）"},
            {"field": "program", "label": "服务项目", "initial": fields["program"], "accepted": fields["program"], "source": "人工新增（项目名称表）"},
            {"field": "request", "label": "需求", "initial": fields["request"], "accepted": fields["request"], "source": "人工新增（原邮件人工判读）"},
        ]
        initial_row = {
            "发件人邮箱": _text(mail.get("sender")),
            "发件日期": _text(mail.get("date")),
            "邮件主题": _text(mail.get("subject")),
            "邮件正文摘要(最多300字)": _text(mail.get("body")),
            "附件名称": _text(mail.get("attachments")),
            "代理": fields["agent"],
            "代理匹配方式": "人工指定",
            "客户编号": fields["customer_code"],
            "客户公司名称": fields["company"],
            "客户提取来源": "人工新增",
            "标准化项目名称": fields["program"],
            "项目原始值": fields["program"],
            "需求": fields["request"],
            "置信度": "人工录入",
            "智能提取方式": "人工新增",
            "数据来源": "人工新增",
            "人工复核提示": note,
        }
        events = [{"at": created, "action": "add_project", "label": "人工新增项目", "reason": note}]
        events.extend(saved.get("events") if isinstance(saved.get("events"), list) else [])
        return {
            "id": rid,
            "source": "人工新增",
            "row_number": None,
            "status": status,
            "sender": _text(mail.get("sender")),
            "date": _text(mail.get("date")),
            "operation_date": created[:10] or _text(mail.get("date"))[:10],
            "subject": _text(mail.get("subject")),
            "body": _text(mail.get("body")),
            "attachments": _text(mail.get("attachments")),
            "fields": fields,
            "evidence": evidence,
            "confidence": "人工录入",
            "agent_match": "人工指定",
            "extraction": "人工新增",
            "review_hint": note,
            "semantic": {"status": "", "issues": "", "suggestion": "", "reason": "", "model": ""},
            "data_source": "人工新增",
            "initial_row": initial_row,
            "events": events,
            "created_at": created,
        }

    @staticmethod
    def _route_map(state: Dict[str, Any]) -> Dict[str, Any]:
        routes = state.get("mail_routes")
        return routes if isinstance(routes, dict) else {}

    @staticmethod
    def _history_mail_key(mail: Dict[str, Any]) -> str:
        """历史按邮件身份归档，不依赖阶段一 Excel 的行号。"""
        sender, date, subject = WorkbenchStore._mail_identity(
            mail.get("sender"), mail.get("date"), mail.get("subject")
        )
        raw = "|".join([sender, date, subject])
        return hashlib.sha1(raw.encode("utf-8", errors="ignore")).hexdigest()[:20]

    @staticmethod
    def _mail_identity(sender: Any, date: Any, subject: Any) -> Tuple[str, str, str]:
        """阶段一和阶段二按同一组邮件元数据关联，兼容 Excel 读出的日期对象。"""
        parsed = _mail_date_sort_value(date)
        date_key = parsed.strftime("%Y-%m-%d %H:%M:%S") if parsed != datetime.min else _text(date)
        return (_text(sender).lower(), date_key, _text(subject))

    def _merged_workorder_rows(self) -> List[Dict[str, Any]]:
        """读取阶段二结果并叠加工作台人工同步结果。"""
        rows = _read_sheet(self.workorder_result_path, "工单核对", "工单核对")
        if self.database is not None:
            self.database.ingest_rows(
                rows,
                dataset="workorder",
                source_path=str(self.workorder_result_path),
            )
            stored = self.database.read_rows("workorder")
            if stored:
                rows = stored
            # 工作台人工核查结果是阶段二结果的覆盖层。它不能直接写回原始
            # Excel（原始文件仍由阶段二输出），否则下一次刷新会把人工结论
            # 覆盖掉；按邮件+公司+项目+需求身份替换同一条结果。
            manual_rows = self.database.read_rows("workorder_manual")
            if manual_rows:
                manual_by_detail = {
                    _workorder_detail_tuple(row): row
                    for row in manual_rows
                    if isinstance(row, dict)
                }
                replaced = set()
                merged: List[Dict[str, Any]] = []
                for row in rows:
                    key = _workorder_detail_tuple(row)
                    override = manual_by_detail.get(key)
                    if override is not None:
                        # 自动阶段二若在人工同步之后再次完成查询，以较新的
                        # 查询时间为准；否则继续显示人工核查结果。
                        base_at = _mail_date_sort_value(row.get("查询时间戳"))
                        manual_at = _mail_date_sort_value(override.get("查询时间戳"))
                        merged.append(override if manual_at >= base_at else row)
                        replaced.add(key)
                    else:
                        merged.append(row)
                merged.extend(
                    row for key, row in manual_by_detail.items()
                    if key not in replaced
                )
                rows = merged
        return rows

    def _workorder_summaries(self) -> Dict[Tuple[str, str, str], Dict[str, Any]]:
        """读取阶段二结论，按邮件聚合“找到/未找到/待核对”。

        只有页面/RPA 已返回明确“是”或“漏单/否”时才算已经处理完成；查询失败、
        未查询和待确认不会被误归档为完成。
        """
        rows = self._merged_workorder_rows()
        grouped: Dict[Tuple[str, str, str], Dict[str, Any]] = {}
        for row in rows:
            key = self._mail_identity(
                row.get("发件人邮箱"), row.get("发件日期"),
                row.get("邮件主题") or row.get("主题"),
            )
            summary = grouped.setdefault(key, {
                "found": 0, "missing": 0, "pending": 0, "total": 0,
                "checked_at": "", "labels": [],
            })
            found = _text(row.get("是否已录单"))
            match_status = _text(row.get("匹配状态"))
            rpa_status = _text(row.get("RPA查询状态"))
            weee_status = _text(row.get("德国WEEE品类状态"))
            checked_at = _text(row.get("查询时间戳"))
            summary["total"] += 1
            if _text(row.get("德国WEEE专项")) == "是" and weee_status in {"未找到", "缺失"}:
                summary["missing"] += 1
                summary["labels"].append("WEEE品类未找到")
            elif _text(row.get("德国WEEE专项")) == "是" and weee_status in {"待人工核对", "待工单品类核对", "未比对"}:
                summary["pending"] += 1
                summary["labels"].append("WEEE品类待核对")
            elif found == "是":
                summary["found"] += 1
                summary["labels"].append("已找到")
            elif (
                "漏单" in match_status
                or (found == "否" and not any(token in rpa_status for token in ("失败", "未查询", "待")))
            ):
                summary["missing"] += 1
                summary["labels"].append("未找到（漏单）")
            else:
                summary["pending"] += 1
                summary["labels"].append("待核对")
            if _mail_date_sort_value(checked_at) >= _mail_date_sort_value(summary["checked_at"]):
                summary["checked_at"] = checked_at
        return grouped

    @staticmethod
    def _history_projection(mail: Dict[str, Any]) -> Dict[str, Any]:
        details = mail.get("details") if isinstance(mail.get("details"), list) else []
        companies = list(dict.fromkeys(
            _text((item.get("fields") or {}).get("company"))
            for item in details if isinstance(item, dict)
            and _text((item.get("fields") or {}).get("company"))
        ))
        projects = list(dict.fromkeys(
            _text((item.get("fields") or {}).get("program"))
            for item in details if isinstance(item, dict)
            and _text((item.get("fields") or {}).get("program"))
        ))
        status = _text(mail.get("status")) or "review"
        workorder = mail.get("workorder_summary") if isinstance(mail.get("workorder_summary"), dict) else {}
        if int(workorder.get("total") or 0) and not int(workorder.get("pending") or 0):
            status = "workorder_completed"
        workorder_result = ""
        if workorder:
            workorder_result = (
                f"已找到 {int(workorder.get('found') or 0)} 条；"
                f"未找到 {int(workorder.get('missing') or 0)} 条；"
                f"待核对 {int(workorder.get('pending') or 0)} 条"
            )
        return {
            "mail_date": _text(mail.get("date")),
            "sender": _text(mail.get("sender")),
            "recipient": _text(mail.get("recipient")),
            "subject": _text(mail.get("subject")),
            "attachments": _text(mail.get("attachments")),
            "source": _text(mail.get("source")),
            "status": status,
            "result": _history_result(status),
            "detail_count": len(details),
            "companies": "；".join(companies),
            "projects": "；".join(projects),
            "workorder_result": workorder_result,
            "workorder_found": int(workorder.get("found") or 0),
            "workorder_missing": int(workorder.get("missing") or 0),
            "workorder_pending": int(workorder.get("pending") or 0),
            "workorder_checked_at": _text(workorder.get("checked_at")),
        }

    @staticmethod
    def _history_events(mail: Dict[str, Any]) -> List[Dict[str, str]]:
        events: List[Dict[str, str]] = []
        for detail in mail.get("details") or []:
            if not isinstance(detail, dict):
                continue
            detail_id = _text(detail.get("id"))
            for raw in detail.get("events") or []:
                if not isinstance(raw, dict):
                    continue
                events.append({
                    "signature": _event_signature(detail_id, raw),
                    "detail_id": detail_id,
                    "at": _text(raw.get("at")),
                    "action": _text(raw.get("action")),
                    "label": _text(raw.get("label")),
                    "reason": _text(raw.get("reason")),
                })
        for raw in mail.get("route_events") or []:
            if not isinstance(raw, dict):
                continue
            detail_id = f"mail:{_text(mail.get('id'))}"
            events.append({
                "signature": _event_signature(detail_id, raw),
                "detail_id": detail_id,
                "at": _text(raw.get("at")),
                "action": _text(raw.get("action")),
                "label": _text(raw.get("label")),
                "reason": _text(raw.get("reason")),
            })
        return sorted(events, key=lambda item: (item["at"], item["signature"]))

    def _write_history_total(self, path: Path, title: str, records: List[Dict[str, Any]]) -> None:
        """将持久历史按邮件日期导出为可直接打开的总表。"""
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = title
        headers = [
            "邮件日期", "邮件标题", "发件人", "收件人", "业务明细数", "客户公司",
            "服务项目", "当前状态", "处理结果", "工单核对结果", "已找到数", "未找到数", "待核对数", "工单查询时间", "最后操作", "最后操作时间",
            "最近原因", "首次归档时间", "最后更新时间", "来源", "附件名称",
        ]
        ws.append(headers)
        header_fill = PatternFill("solid", fgColor="0F766E")
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
        for item in records:
            mail_date = _mail_date_sort_value(item.get("mail_date"))
            mail_value: Any = mail_date if mail_date != datetime.min else _text(item.get("mail_date"))
            ws.append([
                mail_value,
                _text(item.get("subject")),
                _text(item.get("sender")),
                _text(item.get("recipient")),
                int(item.get("detail_count") or 0),
                _text(item.get("companies")),
                _text(item.get("projects")),
                _text(item.get("status")),
                _text(item.get("result")),
                _text(item.get("workorder_result")),
                int(item.get("workorder_found") or 0),
                int(item.get("workorder_missing") or 0),
                int(item.get("workorder_pending") or 0),
                _text(item.get("workorder_checked_at")),
                _text(item.get("last_action")),
                _text(item.get("last_action_at")),
                _text(item.get("last_reason")),
                _text(item.get("first_seen_at")),
                _text(item.get("updated_at")),
                _text(item.get("source")),
                _text(item.get("attachments")),
            ])
        for row in ws.iter_rows(min_row=2, max_col=1):
            row[0].number_format = "yyyy-mm-dd hh:mm:ss"
        widths = [20, 46, 28, 28, 12, 34, 28, 18, 18, 34, 12, 12, 12, 20, 20, 20, 34, 20, 20, 16, 38]
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(path)

    def _sync_persistent_history(
        self, mails: List[Dict[str, Any]], filtered_mails: List[Dict[str, Any]]
    ) -> Dict[str, Any]:
        """同步邮件历史及完成/未完成总表；测试会话绝不写入正式留痕。"""
        if not self.history_enabled:
            return {
                "enabled": False,
                "message": "测试或临时会话：不读取或写入正式历史留痕",
                "completed": [],
                "unfinished": [],
                "completed_count": 0,
                "unfinished_count": 0,
                "completed_path": "",
                "unfinished_path": "",
            }

        state = _load_json(HISTORY_STATE)
        entries = state.setdefault("mails", {})
        if not isinstance(entries, dict):
            entries = {}
            state["mails"] = entries
        now = _now()

        normal_keys = {
            (_text(mail.get("sender")), _text(mail.get("date")), _text(mail.get("subject")))
            for mail in mails
        }
        history_mails = list(mails)
        # 已在询单复核队列的过滤候选由其队列状态记录，避免同一封邮件双份入账。
        for filtered in filtered_mails:
            key = (_text(filtered.get("sender")), _text(filtered.get("date")), _text(filtered.get("subject")))
            if key in normal_keys:
                continue
            history_mails.append({
                "id": _text(filtered.get("id")),
                "sender": _text(filtered.get("sender")),
                "date": _text(filtered.get("date")),
                "subject": _text(filtered.get("subject")),
                "attachments": _text(filtered.get("attachments")),
                "source": "过滤日志",
                "status": "filtered",
                "details": [],
            })

        for mail in history_mails:
            key = self._history_mail_key(mail)
            entry = entries.get(key) if isinstance(entries.get(key), dict) else {}
            projection = self._history_projection(mail)
            known_signatures = set(entry.get("event_signatures") or [])
            stored_events = entry.get("events") if isinstance(entry.get("events"), list) else []
            for event in self._history_events(mail):
                if event["signature"] not in known_signatures:
                    stored_events.append(event)
                    known_signatures.add(event["signature"])
            last_event = stored_events[-1] if stored_events else {}
            entry.update(projection)
            entry["id"] = key
            entry["first_seen_at"] = _text(entry.get("first_seen_at")) or now
            entry["updated_at"] = now
            entry["events"] = stored_events
            entry["event_signatures"] = list(known_signatures)
            entry["last_action"] = _text(last_event.get("label"))
            entry["last_action_at"] = _text(last_event.get("at"))
            entry["last_reason"] = _text(last_event.get("reason"))
            entries[key] = entry

        records = sorted(
            (item for item in entries.values() if isinstance(item, dict)),
            key=lambda item: (_mail_date_sort_value(item.get("mail_date")), _text(item.get("updated_at"))),
            reverse=True,
        )
        completed = [item for item in records if _history_is_complete(item.get("status", ""))]
        unfinished = [item for item in records if not _history_is_complete(item.get("status", ""))]
        state["updated_at"] = now
        _save_json(HISTORY_STATE, state)
        self._write_history_total(COMPLETED_HISTORY_XLSX, "处理完成", completed)
        self._write_history_total(UNFINISHED_HISTORY_XLSX, "未处理完成", unfinished)
        return {
            "enabled": True,
            "message": "正式历史已按邮件日期归档",
            "completed": completed,
            "unfinished": unfinished,
            "completed_count": len(completed),
            "unfinished_count": len(unfinished),
            "completed_path": str(COMPLETED_HISTORY_XLSX),
            "unfinished_path": str(UNFINISHED_HISTORY_XLSX),
        }

    def snapshot(self) -> Dict[str, Any]:
        with self._lock:
            state = self._state()
            deleted = self._deleted_map(state)
            added = self._added_map(state)
            routes = self._route_map(state)
            details = [self._record(row, state) for row in self._raw_records() if row["_id"] not in deleted]
            # 把阶段二刚完成的德国 WEEE 品类核对结果回填到对应业务明细，
            # 工作台刷新后即可看到“邮件品类 / 工单品类 / 逐项状态”。
            workorder_rows = self._merged_workorder_rows()
            workorder_by_detail = {
                _workorder_detail_tuple(row): row
                for row in workorder_rows
                if isinstance(row, dict)
            }
            for detail in details:
                weee = detail.get("weee") if isinstance(detail.get("weee"), dict) else {}
                if not weee.get("enabled"):
                    continue
                identity = _workorder_detail_tuple(detail.get("initial_row") or {})
                workorder_row = workorder_by_detail.get(identity)
                if workorder_row is None:
                    continue
                comparison = _json_object(workorder_row.get("德国WEEE品类核对"))
                if comparison:
                    weee["comparison"] = comparison
                weee["status"] = _text(workorder_row.get("德国WEEE品类状态")) or weee.get("status", "待工单核对")
                weee["reason"] = _text(workorder_row.get("德国WEEE专项说明")) or weee.get("reason", "")
                detail["weee"] = weee
            filtered = _read_sheet(self.filtered_path, "过滤日志", "过滤日志")
            if self.database is not None:
                self.database.ingest_rows(
                    filtered,
                    dataset="filtered",
                    source_path=str(self.filtered_path),
                )
                stored_filtered = self.database.read_rows("filtered")
                if stored_filtered:
                    filtered = stored_filtered
            # 过滤邮件单独提供给工作台的“已过滤邮件”视图。这里不沿用主队列
            # 的 settled 过滤规则：即使是已经确认的证书通知/非注册业务，仍要
            # 保留原始证据，方便人工追溯“为什么没有进入询单队列”。
            filtered_mails: List[Dict[str, Any]] = []
            for row in filtered:
                if row["_id"] in deleted:
                    continue
                filtered_id = f"filtered:{row['_id']}"
                route = routes.get(filtered_id) if isinstance(routes.get(filtered_id), dict) else {}
                # 已由人工转入询单复核的过滤邮件不能同时留在过滤页。
                if _text(route.get("route")) == "review":
                    continue
                reason = _text(row.get("过滤原因"))
                body = (
                    _text(row.get("正文原文"))
                    or _text(row.get("邮件正文原文"))
                    or _text(row.get("正文摘要(最多300字)"))
                    or _text(row.get("正文摘要"))
                    or _text(row.get("正文"))
                )
                filtered_mails.append({
                    "id": filtered_id,
                    "sender": _text(row.get("发件人邮箱")),
                    "recipient": _text(row.get("收件人")) or _text(row.get("收件人邮箱")),
                    "date": _text(row.get("发件日期")),
                    "operation_date": _text(row.get("_db_operation_date")) or _text(row.get("处理时间戳"))[:10] or _text(row.get("发件日期"))[:10],
                    "subject": _text(row.get("主题")) or _text(row.get("邮件主题")),
                    "body": body,
                    "attachments": _text(row.get("附件名称")),
                    "reason": reason,
                    "intent_status": _text(row.get("意图LLM状态")),
                    "processed_at": _text(row.get("处理时间戳")),
                    "settled": _is_settled_filter(reason),
                    "source": "过滤日志",
                    "route_events": route.get("events") if isinstance(route.get("events"), list) else [],
                })
            # 过滤日志不再只作为右上角数量展示，而是进入同一复核队列；
            # 保留过滤原因，人工可以判断是否误过滤并补录字段。
            # 例外（只计数、不进队列）：证书/下号通知等已定论的过滤邮件，
            # 以及被人工明确删除的明细。
            for row in filtered:
                filtered_id = f"filtered:{row['_id']}"
                route = routes.get(filtered_id) if isinstance(routes.get(filtered_id), dict) else {}
                should_restore = _text(route.get("route")) == "review"
                if row["_id"] in deleted or (not should_restore and _is_settled_filter(row.get("过滤原因"))):
                    continue
                row = dict(row)
                row["发件人邮箱"] = row.get("发件人邮箱", "")
                row["收件人"] = row.get("收件人", row.get("收件人邮箱", ""))
                row["发件日期"] = row.get("发件日期", "")
                row["邮件主题"] = row.get("主题", "")
                row["邮件正文摘要(最多300字)"] = row.get("过滤原因", "")
                row["附件名称"] = row.get("附件名称", "")
                row["客户公司名称"] = ""
                row["标准化项目名称"] = ""
                row["需求"] = ""
                row["人工复核提示"] = row.get("过滤原因", "")
                row["数据来源"] = "过滤日志"
                details.append(self._record(row, state))
            groups: Dict[Tuple[str, str, str], Dict[str, Any]] = {}
            for item in details:
                key = (item["sender"], item["date"], item["subject"])
                if key not in groups:
                    groups[key] = {
                        "id": hashlib.sha1("|".join(key).encode("utf-8", errors="ignore")).hexdigest()[:16],
                        "sender": item["sender"],
                        "recipient": item.get("recipient", ""),
                        "date": item["date"],
                        "operation_date": item.get("operation_date", "") or item["date"][:10],
                        "subject": item["subject"],
                        "body": item["body"],
                        "attachments": item["attachments"],
                        "details": [],
                        "source": item["source"],
                    }
                groups[key]["details"].append(item)
            # 人工新增的项目：可能挂到已有邮件下，也可能是阶段一完全没有产出的邮件
            known_ids = {g["id"] for g in groups.values()}
            for mail_id, bucket in added.items():
                if mail_id in known_ids or not isinstance(bucket, dict):
                    continue
                meta_raw = bucket.get("mail") if isinstance(bucket.get("mail"), dict) else {}
                meta = {k: _text(meta_raw.get(k)) for k in ("sender", "date", "subject", "body", "attachments")}
                key = (meta["sender"], meta["date"], meta["subject"])
                if key in groups:  # 同键邮件已在队列里，不要覆盖它原有明细
                    continue
                groups[key] = {
                    "id": mail_id,
                        "sender": meta["sender"],
                        "recipient": _text(meta_raw.get("recipient")),
                        "date": meta["date"],
                        "operation_date": _text(meta_raw.get("operation_date")) or meta["date"][:10],
                        "subject": meta["subject"],
                    "body": meta["body"],
                    "attachments": meta["attachments"],
                    "details": [],
                    "source": "人工新增",
                }
            for mail_id, bucket in added.items():
                if not isinstance(bucket, dict):
                    continue
                target = next((g for g in groups.values() if g["id"] == mail_id), None)
                if target is None:
                    meta_raw = bucket.get("mail") if isinstance(bucket.get("mail"), dict) else {}
                    target = groups.get((_text(meta_raw.get("sender")), _text(meta_raw.get("date")), _text(meta_raw.get("subject"))))
                if target is None:
                    continue
                for raw_item in bucket.get("items") or []:
                    if not isinstance(raw_item, dict) or _text(raw_item.get("id")) in deleted:
                        continue
                    target["details"].append(self._added_record(raw_item, target, state))
            mails = [g for g in groups.values() if g["details"]]
            # 统计口径以最终队列为准（含人工新增的明细，不含被删除的明细）
            all_details = [d for m in mails for d in m["details"]]
            # 同一封邮件的状态由最需要人工介入的一条明细决定。
            for mail in mails:
                states = {d["status"] for d in mail["details"]}
                if "review" in states:
                    mail["status"] = "review"
                elif "returned" in states:
                    mail["status"] = "returned"
                elif states and states <= {"confirmed"}:
                    mail["status"] = "confirmed"
                elif "needs_info" in states:
                    mail["status"] = "needs_info"
                else:
                    mail["status"] = "ready"
                route = routes.get(mail["id"]) if isinstance(routes.get(mail["id"]), dict) else {}
                mail["route_events"] = route.get("events") if isinstance(route.get("events"), list) else []
                # 处理日期筛选必须绑定邮件实际发生日期，不能随着人工操作记录更新而漂移。
                # 人工操作时间仍保存在 route_events/details.events 中，供审计使用。
                mail["operation_date"] = _text(mail.get("date"))[:10]
                # 邮件级统计与“附件表格行数”分开，避免把主题里的数字、项目数
                # 或旧表单行误当成公司数量。工作台据此明确显示主体/项目/明细口径。
                mail["company_count"] = len({
                    _text((detail.get("fields") or {}).get("company"))
                    for detail in mail["details"]
                    if _text((detail.get("fields") or {}).get("company"))
                })
                mail["project_count"] = len({
                    _text((detail.get("fields") or {}).get("program"))
                    for detail in mail["details"]
                    if _text((detail.get("fields") or {}).get("program"))
                })
                mail["detail_count"] = len(mail["details"])

            # 人工将一封询单转为过滤后，原始阶段一行不删除，转而出现在过滤页；
            # 转回询单复核时只改路由状态，保留原始明细和完整操作记录。
            active_mails: List[Dict[str, Any]] = []
            manual_filtered: List[Dict[str, Any]] = []
            for mail in mails:
                route = routes.get(mail["id"]) if isinstance(routes.get(mail["id"]), dict) else {}
                if _text(route.get("route")) != "filtered":
                    active_mails.append(mail)
                    continue
                manual_filtered.append({
                    "id": mail["id"],
                    "sender": mail["sender"],
                    "recipient": mail.get("recipient", ""),
                    "date": mail["date"],
                    "subject": mail["subject"],
                    "body": mail["body"],
                    "attachments": mail["attachments"],
                    "reason": _text(route.get("reason")) or "人工转为已过滤邮件",
                    "intent_status": "人工路由",
                    "processed_at": _text(route.get("at")),
                    "settled": False,
                    "source": "人工过滤",
                    "route_events": route.get("events") if isinstance(route.get("events"), list) else [],
                })
            mails = active_mails
            filtered_mails.extend(manual_filtered)
            workorder_summaries = self._workorder_summaries()
            for mail in mails:
                mail["workorder_summary"] = workorder_summaries.get(
                    self._mail_identity(mail["sender"], mail["date"], mail["subject"]), {}
                )
            # 漏单单独形成可处理视图：它保留原邮件、附件和业务明细，
            # 但不从询单队列删除；处理完成后仍可在历史总表追溯。
            missing_mails = sorted(
                [
                    mail for mail in mails
                    if int((mail.get("workorder_summary") or {}).get("missing") or 0) > 0
                ],
                key=lambda mail: _mail_date_sort_value(mail.get("date")),
                reverse=True,
            )
            all_details = [detail for mail in mails for detail in mail["details"]]
            counts = {
                "mails": len(mails),
                "details": len(all_details),
                "review": sum(1 for d in all_details if d["status"] == "review"),
                "ready": sum(1 for d in all_details if d["status"] == "ready"),
                "confirmed": sum(1 for d in all_details if d["status"] == "confirmed"),
                "returned": sum(1 for d in all_details if d["status"] == "returned"),
                "needs_info": sum(1 for d in all_details if d["status"] == "needs_info"),
                # 当前过滤页包含阶段一过滤日志及人工转入的邮件。
                "filtered": len(filtered_mails),
            }
            history = self._sync_persistent_history(mails, filtered_mails)
            table = project_table_path()
            return {
                "ok": True,
                "generated_at": _now(),
                "paths": {
                    "primary": str(self.primary_path),
                    "review": str(self.review_path),
                    "filtered": str(self.filtered_path),
                    "workorder_result": str(self.workorder_result_path),
                    "missing_workorder": str(
                        DEFAULT_MISSING_WORKORDER
                        if DEFAULT_MISSING_WORKORDER.exists()
                        else (LEGACY_MISSING_WORKORDER if LEGACY_MISSING_WORKORDER.exists() else DEFAULT_MISSING_WORKORDER)
                    ),
                    "database": str(self.database.path) if self.database is not None else "",
                    "project_table": str(table) if table else "",
                },
                "session": {
                    "mode": "test" if self.test_mode else "persistent",
                    "history_loaded": self.history_enabled,
                },
                "counts": counts,
                "mails": mails,
                "missing_mails": missing_mails,
                "projects": load_project_names(),
                "filtered": [{k: _text(v) for k, v in row.items() if not k.startswith("_")} for row in filtered],
                "filtered_mails": filtered_mails,
                "history": history,
                "database": self.database_info(),
            }

    def _locate(self, state: Dict[str, Any], rid: str) -> Dict[str, str]:
        """按明细编号找出该条记录，用于删除时写审计摘要。"""
        for bucket in self._added_map(state).values():
            if not isinstance(bucket, dict):
                continue
            for item in bucket.get("items") or []:
                if isinstance(item, dict) and _text(item.get("id")) == rid:
                    fields = item.get("fields") if isinstance(item.get("fields"), dict) else {}
                    meta = bucket.get("mail") if isinstance(bucket.get("mail"), dict) else {}
                    return {
                        "company": _text(fields.get("company")),
                        "program": _text(fields.get("program")),
                        "subject": _text(meta.get("subject")),
                        "source": "人工新增",
                    }
        for row in self._raw_records():
            if row["_id"] == rid:
                rec = self._record(row, state)
                return {
                    "company": rec["fields"].get("company", ""),
                    "program": rec["fields"].get("program", ""),
                    "subject": rec["subject"],
                    "source": rec["source"],
                }
        return {}

    def _add_project(self, state: Dict[str, Any], payload: Dict[str, Any]) -> Dict[str, Any]:
        mail_id = _text(payload.get("mail_id"))
        if not mail_id:
            raise ValueError("缺少所属邮件，无法新增项目")
        raw = payload.get("fields") if isinstance(payload.get("fields"), dict) else {}
        fields = {k: _text(raw.get(k)) for k in ("agent", "company", "country", "program", "request")}
        if not fields["company"] and not fields["program"]:
            raise ValueError("请至少填写客户公司或服务项目")
        if not fields["country"]:
            fields["country"] = country_of_project(fields["program"])
        meta_raw = payload.get("mail") if isinstance(payload.get("mail"), dict) else {}
        mail = {
            k: _text(meta_raw.get(k))
            for k in ("sender", "recipient", "date", "subject", "body", "attachments")
        }
        added = state.setdefault("added", {})
        bucket = added.setdefault(mail_id, {"mail": {}, "items": []})
        # 邮件元信息用于「整封邮件只剩人工新增明细」时仍能还原队列分组
        if any(mail.values()):
            bucket["mail"] = mail
        items = bucket.setdefault("items", [])
        for existing in items:
            old = existing.get("fields") if isinstance(existing.get("fields"), dict) else {}
            if (
                _text(old.get("company")) == fields["company"]
                and _text(old.get("program")) == fields["program"]
                and _text(old.get("country")) == fields["country"]
            ):
                raise ValueError("该邮件下已存在相同的公司+项目明细，未重复新增")
        item_id = "add_" + hashlib.sha1(f"{mail_id}|{_now()}|{len(items)}".encode("utf-8")).hexdigest()[:12]
        items.append({"id": item_id, "fields": fields, "created_at": _now(), "note": _text(payload.get("reason"))})
        _save_json(self.state_path, state)
        warning = ""
        known = {p["name"] for p in load_project_names()}
        if known and fields["program"] and fields["program"] not in known:
            warning = f"服务项目「{fields['program']}」不在项目名称表中，阶段二可能匹配不到"
        return {
            "ok": True,
            "record_id": item_id,
            "message": f"已新增项目：{fields['company'] or fields['program']}",
            "warning": warning,
        }

    def _delete_project(self, state: Dict[str, Any], rid: str) -> Dict[str, Any]:
        info = self._locate(state, rid)
        deleted = state.setdefault("deleted", {})
        deleted[rid] = {
            "at": _now(),
            "subject": info.get("subject", ""),
            "company": info.get("company", ""),
            "program": info.get("program", ""),
            "source": info.get("source", ""),
        }
        _save_json(self.state_path, state)
        label = info.get("company") or info.get("program") or info.get("subject") or rid
        return {"ok": True, "message": f"已删除项目：{label}"}

    def _bulk_confirm(self, state: Dict[str, Any], payload: Dict[str, Any]) -> Dict[str, Any]:
        """Confirm selected mail groups without bypassing unresolved review states.

        The UI selects whole mail groups, while the persisted state is per detail
        row.  A mail is eligible only when every detail is already ``ready`` or
        ``confirmed``.  Review/returned/needs-info rows are reported back to the
        operator instead of being silently forced into ``confirmed``.
        """
        raw_ids = payload.get("mail_ids")
        if not isinstance(raw_ids, list):
            raise ValueError("批量确认需要邮件ID列表")
        mail_ids = list(dict.fromkeys(_text(value) for value in raw_ids if _text(value)))
        if not mail_ids:
            raise ValueError("请至少选择一封邮件")
        if len(mail_ids) > 200:
            raise ValueError("一次最多批量确认200封邮件")

        reason = _text(payload.get("reason")) or "批量人工确认：邮件字段已核对"
        current = {mail["id"]: mail for mail in self.snapshot().get("mails", [])}
        records = state.setdefault("records", {})
        confirmed_mails: List[str] = []
        confirmed_details = 0
        skipped: List[Dict[str, str]] = []
        changed = False
        for mail_id in mail_ids:
            mail = current.get(mail_id)
            if not mail:
                skipped.append({"mail_id": mail_id, "reason": "邮件不存在或已从当前队列移除"})
                continue
            details = mail.get("details") or []
            blocked = [detail for detail in details if detail.get("status") not in {"ready", "confirmed"}]
            if blocked:
                states = ", ".join(sorted({str(detail.get("status") or "未判定") for detail in blocked}))
                skipped.append({"mail_id": mail_id, "reason": f"含未解决明细：{states}"})
                continue
            pending = [detail for detail in details if detail.get("status") != "confirmed"]
            if not pending:
                skipped.append({"mail_id": mail_id, "reason": "该邮件已确认"})
                continue
            for detail in pending:
                rid = _text(detail.get("id"))
                if not rid:
                    continue
                entry = records.setdefault(rid, {"fields": {}, "events": []})
                if not isinstance(entry.get("events"), list):
                    entry["events"] = []
                entry["status"] = "confirmed"
                entry["events"].append({
                    "at": _now(),
                    "action": "bulk_confirm",
                    "label": "批量人工确认整理完成",
                    "reason": reason,
                })
                entry["updated_at"] = _now()
                confirmed_details += 1
                changed = True
            confirmed_mails.append(mail_id)

        if changed:
            _save_json(self.state_path, state)
        return {
            "ok": True,
            "message": f"已批量确认{len(confirmed_mails)}封邮件、{confirmed_details}条明细",
            "confirmed_mails": len(confirmed_mails),
            "confirmed_details": confirmed_details,
            "skipped": skipped,
            "skipped_count": len(skipped),
        }

    @staticmethod
    def _write_route_event(
        state: Dict[str, Any], mail_id: str, route: str, label: str, reason: str
    ) -> None:
        routes = state.setdefault("mail_routes", {})
        if not isinstance(routes, dict):
            routes = {}
            state["mail_routes"] = routes
        entry = routes.setdefault(mail_id, {"events": []})
        if not isinstance(entry, dict):
            entry = {"events": []}
            routes[mail_id] = entry
        events = entry.setdefault("events", [])
        if not isinstance(events, list):
            events = []
            entry["events"] = events
        event = {
            "at": _now(),
            "action": "move_to_filtered" if route == "filtered" else "move_to_review",
            "label": label,
            "reason": reason,
        }
        events.append(event)
        entry.update({"route": route, "at": event["at"], "reason": reason})

    def _move_mail_route(self, state: Dict[str, Any], payload: Dict[str, Any], route: str) -> Dict[str, Any]:
        mail_id = _text(payload.get("mail_id"))
        if not mail_id:
            raise ValueError("缺少邮件，无法变更队列")
        reason = _text(payload.get("reason"))
        if route == "filtered":
            label = "转为已过滤邮件"
            reason = reason or "人工确认：该邮件不进入询单复核"
        else:
            label = "转入询单复核"
            reason = reason or "人工恢复：需要进入询单复核"
        self._write_route_event(state, mail_id, route, label, reason)
        _save_json(self.state_path, state)
        return {"ok": True, "message": label}

    def _bulk_filter(self, state: Dict[str, Any], payload: Dict[str, Any]) -> Dict[str, Any]:
        raw_ids = payload.get("mail_ids")
        if not isinstance(raw_ids, list):
            raise ValueError("批量过滤需要邮件ID列表")
        mail_ids = list(dict.fromkeys(_text(value) for value in raw_ids if _text(value)))
        if not mail_ids:
            raise ValueError("请至少选择一封邮件")
        if len(mail_ids) > 200:
            raise ValueError("一次最多批量处理200封邮件")
        reason = _text(payload.get("reason")) or "批量人工确认：不进入询单复核"
        for mail_id in mail_ids:
            self._write_route_event(state, mail_id, "filtered", "批量转为已过滤邮件", reason)
        _save_json(self.state_path, state)
        return {"ok": True, "message": f"已将{len(mail_ids)}封邮件转入已过滤邮件", "count": len(mail_ids)}

    def _bulk_restore(self, state: Dict[str, Any], payload: Dict[str, Any]) -> Dict[str, Any]:
        """批量把已过滤邮件恢复到询单复核，并保留每封邮件的流转理由。"""
        raw_ids = payload.get("mail_ids")
        if not isinstance(raw_ids, list):
            raise ValueError("批量恢复需要邮件ID列表")
        mail_ids = list(dict.fromkeys(_text(value) for value in raw_ids if _text(value)))
        if not mail_ids:
            raise ValueError("请至少选择一封邮件")
        if len(mail_ids) > 200:
            raise ValueError("一次最多批量处理200封邮件")
        reason = _text(payload.get("reason")) or "批量人工确认：恢复询单复核"
        for mail_id in mail_ids:
            self._write_route_event(state, mail_id, "review", "批量转入询单复核", reason)
        _save_json(self.state_path, state)
        return {"ok": True, "message": f"已将{len(mail_ids)}封邮件转入询单复核", "count": len(mail_ids)}

    def _manual_workorder_result(self, state: Dict[str, Any], payload: Dict[str, Any]) -> Dict[str, Any]:
        """保存操作人员在工单系统核查后的单条结果。

        结果写入独立的 ``workorder_manual`` 覆盖层，不直接改写阶段二原始
        Excel；刷新工作台时按邮件+业务明细身份覆盖旧的“漏单/待核对”结论。
        """
        if self.database is None:
            raise ValueError("当前测试会话未启用持久数据库，不能同步工单结果")
        mail = payload.get("mail") if isinstance(payload.get("mail"), dict) else {}
        fields = payload.get("fields") if isinstance(payload.get("fields"), dict) else {}
        found = _text(payload.get("found"))
        if found not in {"是", "否", "待复核"}:
            raise ValueError("工单结果必须选择：是、否或待复核")
        sender = _text(mail.get("sender"))
        mail_date = _text(mail.get("date"))
        subject = _text(mail.get("subject"))
        company = _text(fields.get("company"))
        project = _text(fields.get("program"))
        if not sender or not mail_date or not subject or not company or not project:
            raise ValueError("缺少邮件、公司或项目身份，无法同步工单结果")
        now = _now()
        match_status = _text(payload.get("match_status")) or {
            "是": "人工确认-已找到",
            "否": "人工确认-漏单",
            "待复核": "人工确认-待复核",
        }[found]
        query_note = _text(payload.get("note"))
        result_row = {
            "发件人邮箱": sender,
            "收件人": _text(mail.get("recipient")),
            "发件日期": mail_date,
            "邮件主题": subject,
            "邮件正文摘要(最多300字)": _text(mail.get("body")),
            "附件名称": _text(mail.get("attachments")),
            "代理": _text(fields.get("agent")),
            "客户公司名称": company,
            "标准化项目名称": project,
            "需求": _text(fields.get("request")),
            "是否已录单": found,
            "工单编号": _text(payload.get("workorder_number")),
            "工单日期": _text(payload.get("workorder_date")),
            "下单日期": _text(payload.get("placed_date")),
            "匹配状态": match_status,
            "RPA查询状态": "人工查询同步",
            "查询方式": "人工同步",
            "模糊查询词": "",
            "查询时间戳": now,
            "人工核查备注": query_note,
        }
        self.database.ingest_rows(
            [result_row],
            dataset="workorder_manual",
            source_path="工作台人工工单同步",
        )
        record_id = _text(payload.get("record_id"))
        records = state.setdefault("records", {})
        if record_id:
            entry = records.setdefault(record_id, {"fields": {}, "events": []})
            if not isinstance(entry, dict):
                entry = {"fields": {}, "events": []}
                records[record_id] = entry
            events = entry.setdefault("events", [])
            if not isinstance(events, list):
                events = []
                entry["events"] = events
            events.append({
                "at": now,
                "action": "manual_workorder_result",
                "label": f"人工同步工单结果：{found}",
                "reason": query_note or match_status,
            })
            entry["workorder_result"] = result_row
            entry["updated_at"] = now
        _save_json(self.state_path, state)
        return {
            "ok": True,
            "message": f"已同步工单结果：{found}",
            "record_id": record_id,
            "mail_id": _text(payload.get("mail_id")),
            # 操作日志只保存结果摘要，不把正文/附件内容重复写入审计表。
            "result": {
                "found": found,
                "match_status": match_status,
                "workorder_number": _text(payload.get("workorder_number")),
                "query_time": now,
            },
        }

    def _queue_workorder_retry(self, state: Dict[str, Any], payload: Dict[str, Any]) -> Dict[str, Any]:
        """把漏单明细加入自动重查队列；下一次阶段二运行会强制跳过旧缓存。"""
        if not self.history_enabled:
            raise ValueError("当前测试会话不会写入正式自动重查队列")
        mail = payload.get("mail") if isinstance(payload.get("mail"), dict) else {}
        fields = payload.get("fields") if isinstance(payload.get("fields"), dict) else {}
        row = {
            "发件人邮箱": _text(mail.get("sender")),
            "发件日期": _text(mail.get("date")),
            "邮件主题": _text(mail.get("subject")),
            "代理": _text(fields.get("agent")),
            "客户公司名称": _text(fields.get("company")),
            "标准化项目名称": _text(fields.get("program")),
            "需求": _text(fields.get("request")),
        }
        key = "\x1f".join(_workorder_detail_tuple(row))
        if not all(_workorder_detail_tuple(row)):
            raise ValueError("缺少邮件、公司或项目身份，无法加入自动重查")
        now = _now()
        queue = _load_json(WORKORDER_RETRY_QUEUE)
        items = queue.get("items") if isinstance(queue.get("items"), list) else []
        items = [item for item in items if isinstance(item, dict) and _text(item.get("key")) != key]
        items.append({
            "key": key,
            "mail_id": _text(payload.get("mail_id")),
            "record_id": _text(payload.get("record_id")),
            "requested_at": now,
            "row": row,
            "status": "pending",
        })
        _save_json(WORKORDER_RETRY_QUEUE, {"updated_at": now, "items": items})
        record_id = _text(payload.get("record_id"))
        records = state.setdefault("records", {})
        if record_id:
            entry = records.setdefault(record_id, {"fields": {}, "events": []})
            if not isinstance(entry, dict):
                entry = {"fields": {}, "events": []}
                records[record_id] = entry
            events = entry.setdefault("events", [])
            if not isinstance(events, list):
                events = []
                entry["events"] = events
            events.append({
                "at": now,
                "action": "queue_workorder_retry",
                "label": "已加入工单自动重查队列",
                "reason": _text(payload.get("reason")) or "漏单重新进入工单核查",
            })
            entry["updated_at"] = now
        _save_json(self.state_path, state)
        return {
            "ok": True,
            "message": "已加入自动重查队列；下次运行阶段二时将强制重新查询",
            "queue_path": str(WORKORDER_RETRY_QUEUE),
            "requested_at": now,
        }

    def _finish_action(self, payload: Dict[str, Any], result: Dict[str, Any]) -> Dict[str, Any]:
        """兼容旧 JSON 状态的同时，把每次操作复制到 SQLite 留痕表。"""
        if self.database is not None and result.get("ok"):
            try:
                self.database.record_operation(
                    _text(payload.get("action")),
                    record_id=_text(payload.get("record_id")) or _text(result.get("record_id")),
                    mail_id=_text(payload.get("mail_id")),
                    reason=_text(payload.get("reason")),
                    result=result,
                    operation_date=_text(payload.get("operation_date")),
                )
            except Exception:
                # 数据库留痕不能阻断原有人工操作；错误仍可从服务日志/JSON 状态追溯。
                pass
        return result

    def ingest(self, payload: Dict[str, Any]) -> Dict[str, Any]:
        """接收外部程序增量写入的数据，不要求启动 PyQt GUI。

        推荐传入 ``{"dataset":"active", "rows":[...]}``，rows 使用阶段一 Excel
        的表头。为便于其它脚本接入，也接受单条 ``row``。
        """
        if self.database is None:
            return {"ok": False, "error": "当前是测试/临时会话，未启用正式数据库"}
        dataset = _text(payload.get("dataset")) or "active"
        rows = payload.get("rows")
        if isinstance(rows, dict):
            rows = [rows]
        if not isinstance(rows, list):
            row = payload.get("row")
            rows = [row] if isinstance(row, dict) else []
        if not rows:
            raise ValueError("数据库导入需要 rows 或 row")
        result = self.database.ingest_rows(
            rows,
            dataset=dataset,
            source_path=_text(payload.get("source_path")) or "API导入",
        )
        return {"ok": True, **result, "message": f"已增量写入数据库：{result.get('rows', 0)} 条"}

    def operations(self, operation_date: str = "") -> Dict[str, Any]:
        if self.database is None:
            return {"ok": True, "enabled": False, "operations": []}
        return {
            "ok": True,
            "enabled": True,
            "operation_date": _text(operation_date),
            "operations": self.database.read_operations(operation_date),
        }

    def database_info(self) -> Dict[str, Any]:
        if self.database is None:
            return {"ok": True, "enabled": False, "path": "", "records": {}, "operations": 0}
        return {"ok": True, "enabled": True, **self.database.summary()}

    def action(self, payload: Dict[str, Any]) -> Dict[str, Any]:
        rid = _text(payload.get("record_id"))
        action = _text(payload.get("action"))
        allowed = {
            "edit", "confirm", "confirm_detail", "return", "needs_info", "agent_confirm", "reset",
            "add_project", "delete_project", "bulk_confirm", "move_to_filtered",
            "move_to_review", "bulk_filter", "bulk_restore",
            "manual_workorder_result", "queue_workorder_retry",
        }
        mail_actions = {
            "add_project", "bulk_confirm", "bulk_filter", "bulk_restore",
            "move_to_filtered", "move_to_review", "manual_workorder_result",
            "queue_workorder_retry",
        }
        if action not in allowed or (action not in mail_actions and not rid):
            raise ValueError("无效的工作台操作")
        with self._lock:
            state = self._state()
            if action == "add_project":
                return self._finish_action(payload, self._add_project(state, payload))
            if action == "delete_project":
                return self._finish_action(payload, self._delete_project(state, rid))
            if action == "bulk_confirm":
                return self._finish_action(payload, self._bulk_confirm(state, payload))
            if action == "bulk_filter":
                return self._finish_action(payload, self._bulk_filter(state, payload))
            if action == "bulk_restore":
                return self._finish_action(payload, self._bulk_restore(state, payload))
            if action == "manual_workorder_result":
                return self._finish_action(payload, self._manual_workorder_result(state, payload))
            if action == "queue_workorder_retry":
                return self._finish_action(payload, self._queue_workorder_retry(state, payload))
            if action == "move_to_filtered":
                return self._finish_action(payload, self._move_mail_route(state, payload, "filtered"))
            if action == "move_to_review":
                return self._finish_action(payload, self._move_mail_route(state, payload, "review"))
            records = state.setdefault("records", {})
            entry = records.setdefault(rid, {"fields": {}, "events": []})
            if not isinstance(entry.get("events"), list):
                entry["events"] = []
            reason = _text(payload.get("reason"))
            if action == "edit":
                fields = payload.get("fields") or {}
                allowed = {"agent", "company", "country", "program", "request"}
                entry.setdefault("fields", {}).update({k: _text(v) for k, v in fields.items() if k in allowed})
                entry["status"] = "review"
                label = "人工修正字段"
            elif action in {"confirm", "confirm_detail"}:
                entry["status"] = "confirmed"
                label = "人工确认当前业务" if action == "confirm_detail" else "人工确认整理完成"
            elif action == "return":
                entry["status"] = "returned"
                label = "退回复核"
            elif action == "needs_info":
                entry["status"] = "needs_info"
                label = "标记待补资料"
            elif action == "agent_confirm":
                fields = payload.get("fields") or {}
                if _text(fields.get("agent")):
                    entry.setdefault("fields", {})["agent"] = _text(fields["agent"])
                entry["status"] = "ready" if _text(fields.get("agent")) else "review"
                label = "确认代理归属"
            else:
                records.pop(rid, None)
                _save_json(self.state_path, state)
                return self._finish_action(payload, {"ok": True, "message": "已撤销当前人工复核状态"})
            entry["events"].append({"at": _now(), "action": action, "label": label, "reason": reason})
            entry["updated_at"] = _now()
            _save_json(self.state_path, state)
        return self._finish_action(payload, {"ok": True, "message": label})

    def attachment_bytes(self, mail_id: str, name: str = "", token: str = "") -> Tuple[bytes, str, str, str]:
        """返回原始附件；原始文件缺失时返回由结构化证据重建的 xlsx 预览。

        返回值为 ``(内容, 文件名, MIME, 来源)``。token 只允许是缓存目录中的
        basename，不能让浏览器借此读取任意本机路径。
        """
        requested = _text(name)
        safe_token = os.path.basename(_text(token))
        cache_dir = (APP_ROOT / "cache" / "attachments").resolve()
        candidates: List[Path] = []
        if safe_token and safe_token not in {".", ".."}:
            candidate = (cache_dir / safe_token).resolve()
            try:
                candidate.relative_to(cache_dir)
            except ValueError:
                candidate = None
            if candidate is not None and candidate.is_file():
                candidates.append(candidate)

        snap = self.snapshot()
        mail = next((item for item in snap.get("mails", []) if _text(item.get("id")) == _text(mail_id)), None)
        if mail:
            wanted_key = Path(requested).name.casefold()
            for detail in mail.get("details") or []:
                for item in detail.get("attachment_files") or []:
                    item_name = _text(item.get("filename"))
                    item_token = os.path.basename(_text(item.get("token")))
                    if wanted_key and Path(item_name).name.casefold() != wanted_key:
                        continue
                    if item_token:
                        candidate = (cache_dir / item_token).resolve()
                        try:
                            candidate.relative_to(cache_dir)
                        except ValueError:
                            continue
                        if candidate.is_file():
                            candidates.append(candidate)
            unique = []
            seen = set()
            for candidate in candidates:
                if str(candidate) not in seen:
                    seen.add(str(candidate))
                    unique.append(candidate)
            candidates = unique
        if candidates:
            path = candidates[0]
            filename = requested or path.name.split("_", 1)[-1]
            mime = mimetypes.guess_type(filename)[0] or "application/octet-stream"
            return path.read_bytes(), filename, mime, "原始附件"

        # 旧工作簿没有保留附件本体，但可能保留了可恢复的 xlsx 行快照。
        evidence: List[dict] = []
        if mail:
            seen = set()
            for detail in mail.get("details") or []:
                for item in detail.get("attachment_evidence") or []:
                    item_name = _text(item.get("filename"))
                    if not item_name or (requested and Path(item_name).name.casefold() != Path(requested).name.casefold()):
                        continue
                    key = (item_name, json.dumps(item, ensure_ascii=False, sort_keys=True))
                    if key not in seen:
                        seen.add(key)
                        evidence.append(item)
        if not evidence:
            raise FileNotFoundError("原始附件未保留，当前数据也没有可重建的结构化预览")
        out_name = requested or _text(evidence[0].get("filename")) or "附件预览.xlsx"
        if not Path(out_name).suffix:
            out_name += ".xlsx"
        wb = Workbook()
        default = wb.active
        wb.remove(default)
        used_titles = set()
        for attachment in evidence:
            sheets = attachment.get("sheets") or []
            if not sheets and attachment.get("text"):
                sheets = [{"sheet_name": "文本预览", "rows": [{"row_number": 1, "cells": [attachment.get("text", "")] }]}]
            for sheet in sheets:
                title = _text(sheet.get("sheet_name")) or "工作表"
                title = re.sub(r"[\\/*?:\[\]]", "_", title)[:31] or "工作表"
                base = title
                suffix = 2
                while title in used_titles:
                    title = f"{base[:28]}_{suffix}"
                    suffix += 1
                used_titles.add(title)
                ws = wb.create_sheet(title)
                for row in sheet.get("rows") or []:
                    values = [row.get("row_number", "")]
                    values.extend(row.get("cells") or [])
                    ws.append(values)
                ws.freeze_panes = "A2"
                for column in ws.columns:
                    width = min(42, max(12, max((len(_text(cell.value)) for cell in column), default=0) + 2))
                    ws.column_dimensions[get_column_letter(column[0].column)].width = width
        if not wb.sheetnames:
            ws = wb.create_sheet("预览")
            ws.append(["说明"])
            ws.append(["当前仅能从邮件快照重建预览，原始附件文件未保存。"])
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue(), out_name, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "结构化预览（非原始附件）"

    def export(self) -> Path:
        with self._lock:
            snap = self.snapshot()
            output_dir = MANUAL_OUTPUT
            output_dir.mkdir(parents=True, exist_ok=True)
            out = output_dir / f"workbench_reviewed_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
            wb = Workbook()
            ws = wb.active
            # 主表沿用阶段一“工单待查”表头，因此可直接被阶段二候选列表识别。
            ws.title = "工单待查"
            source_rows = self._raw_records()
            original_headers: List[str] = []
            for row in source_rows:
                for key in row:
                    if not key.startswith("_") and key not in original_headers:
                        original_headers.append(key)
            required = ["代理", "客户公司名称", "标准化项目名称", "需求"]
            headers = required + [h for h in original_headers if h not in required]
            headers += ["工作台状态", "人工修改原因", "工作台最后操作时间"]
            ws.append(headers)
            for mail in snap["mails"]:
                for detail in mail["details"]:
                    events = detail.get("events") or []
                    reason = events[-1].get("reason", "") if events else ""
                    values = dict(detail.get("initial_row") or {})
                    values.update({
                        "代理": detail["fields"].get("agent", ""),
                        "客户公司名称": detail["fields"].get("company", ""),
                        "标准化项目名称": detail["fields"].get("program", ""),
                        "需求": detail["fields"].get("request", ""),
                        "工作台状态": detail["status"],
                        "人工修改原因": reason,
                        "工作台最后操作时间": events[-1].get("at", "") if events else "",
                    })
                    ws.append([values.get(header, "") for header in headers])
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for idx, header in enumerate(headers, start=1):
                ws.column_dimensions[chr(64 + idx) if idx <= 26 else "A"].width = min(42, max(14, len(header) + 4))
            log_ws = wb.create_sheet("人工操作记录")
            log_ws.append(["明细编号", "邮件主题", "操作时间", "操作", "原因"])
            for mail in snap["mails"]:
                for detail in mail["details"]:
                    for event in detail.get("events") or []:
                        log_ws.append([detail["id"], mail["subject"], event.get("at", ""), event.get("label", ""), event.get("reason", "")])
            log_ws.freeze_panes = "A2"
            # 被人工删除的项目不写回主表，但要留痕，便于事后追溯/恢复
            deleted_rows = self._deleted_map(self._state())
            if deleted_rows:
                del_ws = wb.create_sheet("已删除明细")
                del_ws.append(["明细编号", "邮件主题", "客户公司", "服务项目", "来源", "删除时间"])
                for rid, info in deleted_rows.items():
                    info = info if isinstance(info, dict) else {}
                    del_ws.append([
                        rid,
                        info.get("subject", ""),
                        info.get("company", ""),
                        info.get("program", ""),
                        info.get("source", ""),
                        info.get("at", ""),
                    ])
                del_ws.freeze_panes = "A2"
            wb.save(out)
            return out


class _Handler(BaseHTTPRequestHandler):
    server_version = "ECOPVWorkbench/1.0"

    def _store(self) -> WorkbenchStore:
        return self.server.workbench_store  # type: ignore[attr-defined]

    def _send(self, payload: Any, status: int = 200, content_type: str = "application/json; charset=utf-8") -> None:
        raw = payload if isinstance(payload, bytes) else json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Cache-Control", "no-store")
        # 允许用户通过独立启动脚本打开 file:// 页面；仅接受浏览器的 null origin，
        # 不把本机数据服务开放给任意网站。
        if self.headers.get("Origin") == "null":
            self.send_header("Access-Control-Allow-Origin", "null")
        self.end_headers()
        self.wfile.write(raw)

    def _send_binary(self, payload: bytes, filename: str, content_type: str, source: str = "") -> None:
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(payload)))
        self.send_header("Content-Disposition", f"inline; filename*=UTF-8''{quote(filename)}")
        self.send_header("Cache-Control", "no-store")
        if source:
            # HTTP 头只能使用 latin-1；来源说明在界面不需要靠该头展示，
            # 这里使用 ASCII 标识避免中文来源导致响应头编码异常。
            source_code = "original" if source == "原始附件" else "reconstructed-preview"
            self.send_header("X-ECOPV-Attachment-Source", source_code)
        if self.headers.get("Origin") == "null":
            self.send_header("Access-Control-Allow-Origin", "null")
        self.end_headers()
        self.wfile.write(payload)

    def do_OPTIONS(self) -> None:  # noqa: N802
        if self.headers.get("Origin") == "null":
            self.send_response(204)
            self.send_header("Access-Control-Allow-Origin", "null")
            self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
            self.send_header("Access-Control-Allow-Headers", "Content-Type")
            self.send_header("Content-Length", "0")
            self.end_headers()
            return
        self._send({"ok": False, "error": "不允许的跨域来源"}, 403)

    def do_GET(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        if parsed.path in {"/", "/index.html"}:
            try:
                self._send(HTML_PATH.read_bytes(), content_type="text/html; charset=utf-8")
            except OSError as exc:
                self._send({"ok": False, "error": str(exc)}, 500)
            return
        if parsed.path == "/api/state":
            try:
                self._send(self._store().snapshot())
            except Exception as exc:
                self._send({"ok": False, "error": str(exc)}, 500)
            return
        if parsed.path == "/api/attachment":
            query = parse_qs(parsed.query)
            mail_id = query.get("mail_id", [""])[0]
            name = query.get("name", [""])[0]
            token = query.get("token", [""])[0]
            try:
                payload, filename, content_type, source = self._store().attachment_bytes(mail_id, name, token)
                self._send_binary(payload, filename, content_type, source)
            except FileNotFoundError as exc:
                self._send({"ok": False, "error": str(exc)}, 404)
            except Exception as exc:
                self._send({"ok": False, "error": str(exc)}, 500)
            return
        if parsed.path == "/api/health":
            self._send({"ok": True, "service": "workbench", "time": _now()})
            return
        if parsed.path == "/api/operations":
            operation_date = parse_qs(parsed.query).get("date", [""])[0]
            try:
                self._send(self._store().operations(operation_date))
            except Exception as exc:
                self._send({"ok": False, "error": str(exc)}, 500)
            return
        if parsed.path == "/api/database":
            try:
                self._send(self._store().database_info())
            except Exception as exc:
                self._send({"ok": False, "error": str(exc)}, 500)
            return
        self._send({"ok": False, "error": "not found"}, 404)

    def do_POST(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        length = int(self.headers.get("Content-Length", "0") or 0)
        try:
            payload = json.loads(self.rfile.read(length).decode("utf-8")) if length else {}
        except (ValueError, UnicodeDecodeError):
            self._send({"ok": False, "error": "请求不是有效 JSON"}, 400)
            return
        try:
            if parsed.path == "/api/action":
                self._send(self._store().action(payload))
            elif parsed.path == "/api/ingest":
                self._send(self._store().ingest(payload))
            elif parsed.path == "/api/export":
                out = self._store().export()
                self._send({"ok": True, "path": str(out), "message": f"已导出：{out.name}"})
            else:
                self._send({"ok": False, "error": "not found"}, 404)
        except Exception as exc:
            self._send({"ok": False, "error": str(exc)}, 400)

    def log_message(self, format: str, *args: Any) -> None:
        return


class WorkbenchServer:
    def __init__(
        self,
        primary_path: Optional[str] = None,
        review_path: Optional[str] = None,
        filtered_path: Optional[str] = None,
        port: int = 8765,
        database_path: Optional[str] = None,
        test_mode: bool = False,
        host: str = "127.0.0.1",
        public_host: Optional[str] = None,
    ):
        self.store = WorkbenchStore(
            primary_path,
            review_path,
            filtered_path=filtered_path,
            database_path=database_path,
            test_mode=test_mode,
        )
        self.port = int(port or 8765)
        # 默认仍只监听本机；局域网部署时由启动参数显式传入 0.0.0.0。
        self.host = str(host or "127.0.0.1").strip() or "127.0.0.1"
        self.public_host = str(public_host or "").strip()
        self.httpd: Optional[ThreadingHTTPServer] = None
        self.thread: Optional[threading.Thread] = None

    def _url_host(self) -> str:
        if self.public_host:
            return self.public_host
        if self.host not in {"0.0.0.0", "::", ""}:
            return self.host
        # 仅用于启动提示和自动打开浏览器；实际监听仍使用 0.0.0.0。
        try:
            return socket.gethostbyname(socket.gethostname())
        except OSError:
            return "127.0.0.1"

    def start(self, open_browser: bool = True) -> str:
        if self.httpd:
            url = f"http://{self._url_host()}:{self.httpd.server_address[1]}/"
            if open_browser:
                webbrowser.open(url)
            return url
        try:
            self.httpd = ThreadingHTTPServer((self.host, self.port), _Handler)
        except OSError:
            # 端口被旧实例占用时仍允许 GUI 启动，但独立 file:// 入口需要使用
            # 新页面 URL，而不是直接双击 HTML 文件。
            self.httpd = ThreadingHTTPServer((self.host, 0), _Handler)
        self.httpd.workbench_store = self.store  # type: ignore[attr-defined]
        self.thread = threading.Thread(target=self.httpd.serve_forever, name="ecopv-workbench", daemon=True)
        self.thread.start()
        url = f"http://{self._url_host()}:{self.httpd.server_address[1]}/"
        if open_browser:
            webbrowser.open(url)
        return url

    def stop(self) -> None:
        if self.httpd:
            self.httpd.shutdown()
            self.httpd.server_close()
            self.httpd = None
            self.thread = None
