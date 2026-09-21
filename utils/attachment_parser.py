"""附件解析工具 — 支持 .zip .rar .xlsx .xls .pdf .docx .csv .jpg .png .jpeg .bmp"""
import os
import re
import shutil
import stat
import subprocess
import tempfile
import zipfile
from typing import List, Dict, Optional

from utils.runtime_paths import APP_ROOT

logger = None
_ocr_engine = None


def _get_logger():
    global logger
    if logger is None:
        import logging
        logger = logging.getLogger("mail_audit")
    return logger


_ocr_init_failed = False

# 图片附件持久化目录（懒加载 OCR 用，按内容 hash 去重）
ATTACHMENT_CACHE_DIR = os.path.join(
    str(APP_ROOT), "cache", "attachments"
)


_WINDOWS_RESERVED_NAMES = {
    "CON", "PRN", "AUX", "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}


def sanitize_attachment_filename(filename: str, default: str = "image") -> str:
    """把邮件附件名转换为可在 Windows 文件系统落盘的文件名。

    邮件头是外部不可信输入，可能含 CR/LF、控制字符、路径分隔符、
    Windows 保留设备名或末尾空格/句点。这里只处理文件名本身，不允许
    它参与目录拼接；原始显示名由调用方自行保留。
    """
    value = str(filename or "")
    # CR/LF 等控制字符会直接让 Windows open() 抛出 Errno 22。
    value = re.sub(r"[\x00-\x1f\x7f]", "_", value)
    value = re.sub(r'[\\/:*?"<>|]', "_", value)
    value = re.sub(r"\s+", " ", value).strip(" .")
    if not value or value in {".", ".."}:
        value = default

    stem, ext = os.path.splitext(value)
    if stem.upper() in _WINDOWS_RESERVED_NAMES:
        stem = f"_{stem}"
        value = f"{stem}{ext}"

    # 给 hash 前缀和父目录留出余量，避免超出 Windows 单文件名/路径限制。
    max_filename_length = 180
    if len(value) > max_filename_length:
        keep = max(1, max_filename_length - len(ext))
        value = f"{stem[:keep]}{ext}"
    return value or default


def save_attachment(payload: bytes, filename: str) -> str:
    """持久化邮件附件并返回安全的本地路径。

    以前只有图片会落盘，xlsx/pdf/zip 解析完就被删除，因此工作台只能显示
    文件名，点击时没有真实文件可打开。所有顶层附件现在统一按内容 hash 去重
    保存；文件名仍经过清洗，不允许邮件头参与目录穿越。
    """
    import hashlib
    os.makedirs(ATTACHMENT_CACHE_DIR, exist_ok=True)
    h = hashlib.sha1(payload).hexdigest()[:12]
    safe = sanitize_attachment_filename(filename)
    path = os.path.join(ATTACHMENT_CACHE_DIR, f"{h}_{safe}")
    if not os.path.exists(path):
        with open(path, "wb") as f:
            f.write(payload)
    return path


def save_image_attachment(payload: bytes, filename: str) -> str:
    """兼容旧调用方：图片也走统一附件持久化。"""
    return save_attachment(payload, filename)


IMAGE_EXTS = (".jpg", ".jpeg", ".png", ".bmp", ".tiff")

# 邮件附件属于不可信输入。限制数量和解压后总量，避免路径穿越与压缩炸弹。
MAX_ARCHIVE_MEMBERS = 200
MAX_ARCHIVE_MEMBER_BYTES = 25 * 1024 * 1024
MAX_ARCHIVE_TOTAL_BYTES = 100 * 1024 * 1024

# Excel 附件不是只有一段“文本”。代理常把 18 家主体放在一个结构化表中，
# 而正文只写 15 家。若只拼成 text_content，后续规则会把正文误当成唯一来源，
# 静默少出三条。这里仅识别明确带业务实体表头的工作表，并保留每个物理数据行。
MAX_XLSX_SCAN_ROWS = 2000
XLSX_HEADER_ALIASES = {
    "customer": ("客户", "客户名称", "客户公司", "客户公司名称", "公司", "公司名称",
                 "公司英文名", "公司英文名称", "公司中文名", "公司中文名称",
                 "企业", "主体", "申请主体", "申请公司", "company", "company name",
                 "customer", "applicant"),
    "customer_code": ("客户号", "客户编号", "主体编号", "企业编号", "申请编号", "编号",
                      "序号", "company id", "customer id", "code"),
    "project": ("项目", "项目名称", "服务项目", "注册项目", "申请项目", "申报项目",
                "project", "program"),
    "country": ("国家", "注册国家", "目的国", "country"),
    "business": ("业务", "业务类型", "服务类型", "法规", "服务", "business", "service"),
    "request": ("需求", "申请类型", "事项", "动作", "办理类型", "request"),
    "agent": ("代理", "代理名称", "服务商", "服务商名称", "agent", "representative"),
    # 德国 WEEE 品类专项使用；这些字段只作为附件证据保存，不会自动
    # 把品牌/品类扩成客户公司或服务项目。
    "brand": ("品牌", "品牌名称", "品牌名", "brand", "brand name", "marke"),
    "category": (
        "品类", "品類", "类别", "類別", "产品类别", "產品類別", "商品类别",
        "商品類別", "产品分类", "產品分類", "设备类别", "設備類別", "注册类别",
        "注册品类", "申报品类", "申報品類", "category", "product category",
        "product type", "warengruppe", "produktkategorie",
    ),
}


def _repair_legacy_zip_filename(name: str, utf8_flag: bool = False) -> str:
    """修复未标 UTF-8 的 GBK/GB18030 ZIP 成员名。

    Python 的 zipfile 会把这类成员名按 CP437 解码，进而得到 "└ε╝╤..."
    一类盒线字符。仅在确有盒线/块元素乱码特征时反向编码，避免修改正常
    的英文文件名或已标记 UTF-8 的压缩包。
    """
    if not name or utf8_flag:
        return name

    has_cp437_mojibake = any(
        "\u2500" <= char <= "\u259f" or char in {"ε", "Γ", "Θ", "φ"}
        for char in name
    )
    if not has_cp437_mojibake:
        return name

    try:
        candidate = name.encode("cp437").decode("gb18030")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return name

    # 只接受确实恢复出中文的候选值；否则保留原名，绝不为了显示而猜文件名。
    if any("\u4e00" <= char <= "\u9fff" for char in candidate):
        return candidate
    return name


def _zip_member_name(info: zipfile.ZipInfo) -> str:
    """返回用于落盘/展示的 ZIP 成员名，保留 ZipInfo 本身供 zf.open 使用。"""
    return _repair_legacy_zip_filename(
        info.filename, bool(info.flag_bits & 0x800)
    )


def _safe_archive_target(destination: str, member_name: str) -> Optional[str]:
    """返回受控解压路径；绝不接受绝对路径、`..` 或盘符逃逸。"""
    name = (member_name or "").replace("\\", "/")
    if not name or name.startswith("/") or re.match(r"^[A-Za-z]:", name):
        return None
    parts = [p for p in name.split("/") if p not in ("", ".")]
    if any(p == ".." for p in parts):
        return None
    root = os.path.abspath(destination)
    target = os.path.abspath(os.path.join(root, *parts))
    try:
        if os.path.commonpath([root, target]) != root:
            return None
    except ValueError:
        return None
    return target


def _validate_archive_members(members: list, destination: str, name_getter=None) -> Optional[str]:
    if len(members) > MAX_ARCHIVE_MEMBERS:
        return f"成员数 {len(members)} 超过上限 {MAX_ARCHIVE_MEMBERS}"
    total = 0
    for member in members:
        name = (
            name_getter(member) if name_getter else
            (getattr(member, "filename", "") or getattr(member, "name", ""))
        )
        if _safe_archive_target(destination, str(name)) is None:
            return f"检测到不安全成员路径: {name!r}"
        size = int(getattr(member, "file_size", 0) or 0)
        if size > MAX_ARCHIVE_MEMBER_BYTES:
            return f"成员过大: {name!r} ({size} bytes)"
        total += size
        if total > MAX_ARCHIVE_TOTAL_BYTES:
            return f"解压总量超过上限 {MAX_ARCHIVE_TOTAL_BYTES} bytes"
    return None


def _safe_extract_zip(zf: zipfile.ZipFile, destination: str) -> None:
    members = zf.infolist()
    problem = _validate_archive_members(members, destination, _zip_member_name)
    if problem:
        raise ValueError(problem)
    for info in members:
        # Unix 模式的符号链接可把后续文件写到目标目录外，拒绝它。
        mode = info.external_attr >> 16
        if stat.S_ISLNK(mode):
            raise ValueError(f"不允许压缩包符号链接: {info.filename!r}")
        target = _safe_archive_target(destination, _zip_member_name(info))
        if info.is_dir():
            os.makedirs(target, exist_ok=True)
            continue
        os.makedirs(os.path.dirname(target), exist_ok=True)
        with zf.open(info, "r") as source, open(target, "wb") as output:
            shutil.copyfileobj(source, output, length=1024 * 1024)


def _safe_extract_rar(rf, destination: str) -> None:
    """通过 rarfile 的文件流逐个释放已验证成员，避免 `extractall`。"""
    members = rf.infolist()
    problem = _validate_archive_members(members, destination)
    if problem:
        raise ValueError(problem)
    for info in members:
        target = _safe_archive_target(destination, info.filename)
        if info.isdir():
            os.makedirs(target, exist_ok=True)
            continue
        os.makedirs(os.path.dirname(target), exist_ok=True)
        with rf.open(info) as source, open(target, "wb") as output:
            shutil.copyfileobj(source, output, length=1024 * 1024)


def _configure_rar_tool(rarfile_module) -> str:
    """让 rarfile 在 Windows 上能发现 winget 安装但尚未刷新 PATH 的 7-Zip。"""
    candidates = [shutil.which("7z"), shutil.which("7zz")]
    if os.name == "nt":
        program_files = os.environ.get("ProgramFiles", r"C:\Program Files")
        candidates.extend([
            os.path.join(program_files, "7-Zip", "7z.exe"),
            os.path.join(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"), "7-Zip", "7z.exe"),
        ])
    for candidate in candidates:
        if candidate and os.path.isfile(candidate):
            rarfile_module.SEVENZIP_TOOL = candidate
            rarfile_module.CURRENT_SETUP = None
            return candidate
    return ""

def _get_ocr():
    """懒加载 PaddleOCR 引擎（全程只初始化一次）"""
    global _ocr_engine, _ocr_init_failed
    if _ocr_init_failed:
        return None
    if _ocr_engine is None:
        try:
            # mkldnn 在部分 Windows 环境下触发段错误，必须在 paddle 导入前关闭
            os.environ.setdefault("FLAGS_use_mkldnn", "0")
            os.environ.setdefault("FLAGS_allocator_strategy", "auto_growth")
            from paddleocr import PaddleOCR
            try:
                # PaddleOCR 3.x：显式关闭 mkldnn
                _ocr_engine = PaddleOCR(lang="ch", enable_mkldnn=False)
            except TypeError:
                # PaddleOCR 2.x：不认识该参数，靠上面的环境变量兜底
                _ocr_engine = PaddleOCR(lang="ch")
            _get_logger().info("PaddleOCR 引擎初始化成功")
        except Exception as e:
            _get_logger().error(f"PaddleOCR 初始化失败: {e}")
            _ocr_init_failed = True
            return None
    return _ocr_engine


def _tesseract_path() -> str:
    """查找发布包内置或系统安装的 Tesseract。"""
    candidates = [
        os.path.join(str(APP_ROOT), "tools", "tesseract", "tesseract.exe"),
        shutil.which("tesseract"),
    ]
    if os.name == "nt":
        candidates.extend([
            os.path.join(os.environ.get("ProgramFiles", r"C:\Program Files"),
                         "Tesseract-OCR", "tesseract.exe"),
            os.path.join(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"),
                         "Tesseract-OCR", "tesseract.exe"),
        ])
    for candidate in candidates:
        if candidate and os.path.isfile(candidate):
            return candidate
    return ""


def _parse_image_with_tesseract(filepath: str) -> str:
    """使用轻量 Tesseract 作为 PaddleOCR 不可用时的图片 OCR 兜底。"""
    executable = _tesseract_path()
    if not executable:
        _get_logger().warning("OCR不可用：未找到 PaddleOCR 或 Tesseract")
        return ""
    tessdata = os.path.join(os.path.dirname(executable), "tessdata")
    command = [executable, filepath, "stdout", "-l", "chi_sim+eng", "--psm", "6"]
    if os.path.isdir(tessdata):
        command.extend(["--tessdata-dir", tessdata])
    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=45,
            creationflags=0x08000000 if os.name == "nt" else 0,
            check=False,
        )
    except Exception as exc:
        _get_logger().warning(f"Tesseract OCR 启动失败 {filepath}: {exc}")
        return ""
    if completed.returncode not in (0, 1):
        detail = (completed.stderr or "").strip()[:300]
        _get_logger().warning(
            f"Tesseract OCR 返回错误码 {completed.returncode} {filepath}: {detail}"
        )
        return ""
    text = (completed.stdout or "").strip()
    if text:
        _get_logger().info(f"Tesseract OCR 识别成功: {os.path.basename(filepath)}")
    return text


def parse_attachment(filepath: str, filename: str) -> Dict:
    """
    解析附件文件，提取文本内容
    返回: {'filename': str, 'text_content': str, 'sheets': list, 'epr_forms': list}

    epr_forms: 若该附件(或压缩包内)含「可勾选的泛欧 EPR 申请表」, 则带上勾选解析结果,
               由 field_extractor 优先据此确定项目(只查打勾的国家×业务组合)。
    """
    result = {
        "filename": filename,
        "text_content": "",
        "sheets": [],
        "epr_forms": [],
        "structured_records": [],
    }
    if not os.path.exists(filepath):
        return result

    ext = os.path.splitext(filename)[1].lower()

    try:
        if ext == ".xlsx":
            result.update(_parse_excel(filepath))
            form = _try_epr_form(filepath, filename)
            if form:
                result["epr_forms"] = [form]
        elif ext == ".xls":
            result["text_content"] = _parse_xls(filepath)
        elif ext == ".pdf":
            result["text_content"] = _parse_pdf(filepath)
        elif ext == ".docx":
            result["text_content"] = _parse_docx(filepath)
        elif ext == ".zip":
            result.update(_parse_zip(filepath))
        elif ext == ".rar":
            result.update(_parse_rar(filepath))
        elif ext == ".csv":
            result["text_content"] = _parse_csv(filepath)
        elif ext in (".txt",):
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                result["text_content"] = f.read()
        elif ext in IMAGE_EXTS:
            # 懒加载 OCR: 常规解析阶段不识别, 只登记待识别图片
            # OCR 兜底由 FieldExtractor._ocr_fallback 在字段仍缺失时触发
            result["text_content"] = ""
            result["ocr_pending"] = True
            result["filepath"] = filepath
        else:
            result["text_content"] = ""
            _get_logger().warning(f"不支持的附件格式: {filename}")
    except Exception as e:
        _get_logger().error(f"解析附件失败 {filename}: {e}")
        result["text_content"] = ""

    # 结构化行必须能回溯到实际附件；压缩包内文件会在递归解析时保留自己的
    # filename，顶层 xlsx 则在这里补上。
    for record in result.get("structured_records") or []:
        record.setdefault("attachment_name", filename)

    return result


def _try_epr_form(filepath: str, filename: str) -> Optional[Dict]:
    """尝试把 xlsx 当作 EPR 申请表解析勾选状态。

    非可勾选模板(如德国ECOPV服务信息申请表/意大利EPR申请表)返回 None,
    调用方保持原有文本规则提取, 不做任何改变。
    """
    try:
        from utils.epr_form_parser import parse_epr_form
    except ImportError:
        try:
            from .epr_form_parser import parse_epr_form
        except ImportError:
            return None
    try:
        form = parse_epr_form(filepath)
    except Exception as e:
        _get_logger().warning(f"EPR申请表解析失败 {filename}: {e}")
        return None
    if not form:
        return None

    form = dict(form)
    form["filename"] = filename
    detail = "、".join(form.get("projects") or []) or "无"
    extra = ""
    if form.get("unmatched_countries"):
        extra += f" | 国家已勾未勾业务: {form['unmatched_countries']}"
    if form.get("orphan_business"):
        extra += f" | 业务已勾未勾国家: {form['orphan_business']}"
    _get_logger().info(f"EPR申请表勾选解析: {filename} → {detail}{extra}")
    return form


def collect_epr_forms(attachments: list) -> List[Dict]:
    """汇总所有附件(含压缩包内)携带的 EPR 勾选结果"""
    forms = []
    for att in attachments or []:
        for f in att.get("epr_forms") or []:
            if f and f.get("projects") is not None:
                forms.append(f)
    return forms



def _xlsx_cell_text(value) -> str:
    """把 Excel 单元格转成稳定、可审计的文本。"""
    if value is None:
        return ""
    return str(value).replace("\u3000", " ").strip()


def _xlsx_header_role(value: str) -> str:
    """将一个表头映射到字段角色；空字符串表示不是业务表头。"""
    raw = str(value or "").strip()
    # EPR 申请表的说明页会出现“请按公司逐份填写”“请选择服务”等长段落。
    # 旧逻辑只要看到“公司”和“服务”就把该行当成明细表头，随后把法人、
    # 注册资本、签字时间等表单标签逐行当成客户。长段落不是结构化表头，
    # 必须拒绝；真正的清单表头通常是短标签，即使中英双语也不会是整段说明。
    if len(raw) > 80 and ("\n" in raw or "\r" in raw):
        return ""
    normalized = re.sub(r"\s+", "", raw).lower()
    if not normalized:
        return ""
    # 客户列只能由“客户/公司名称”这类短表头命中。旧逻辑对所有别名
    # 使用 substring 匹配，导致“必须由正规翻译公司盖章”“注意事项”等
    # 说明文字因包含“公司”而被误当成客户列，随后整张资料清单被展开成
    # 业务明细。去掉表头常见的星号/括号后仍只接受明确的表头词，拒绝
    # 任意自然语言句子。
    header_clean = re.sub(r"[\s*＊:：()（）\[\]【】]", "", normalized)
    customer_headers = {
        "客户", "客户名称", "客户公司", "客户公司名称", "公司名", "公司名称",
        "公司英文名", "公司英文名称", "公司中文名", "公司中文名称", "companyname", "companynamecn", "companynameen",
        "企业", "企业名称", "主体", "主体名称", "申请主体", "申请公司",
        "company", "companyname", "company name", "customer", "applicant",
    }
    for role, aliases in XLSX_HEADER_ALIASES.items():
        for alias in aliases:
            alias_norm = re.sub(r"\s+", "", alias).lower()
            if not alias_norm:
                continue
            if role == "customer":
                # “公司”单独出现可以是业务表单字段，不足以证明这是
                # 一行一家的客户清单；仅接受完整客户主体表头。
                matched = header_clean in customer_headers
            else:
                matched = normalized == alias_norm or alias_norm in normalized
            if matched:
                return role
    return ""


_XLSX_NON_COMPANY_VALUE_HINTS = (
    "poa", "legalrepresentative", "legalperson", "legalpositions",
    "nameoflegalperson", "placeofsignature", "signingtime",
    "registrationcapital", "companyname", "companyaddress", "companybusiness",
    "plz", "postcode", "amazonlink", "shoplink", "e-mail", "email", "tel", "phone",
    "legrepresentativename", "companyregistrationnumber", "registrationnumber", "uscc",
    "营业执照", "公司名称", "公司地址", "公司注册", "注册资本", "法人", "非中国公司",
    "公司成立日期", "成立日期", "签字", "签署", "职位", "联系信息", "联系人", "联系电话", "邮箱",
    "邮政编码", "邮编", "地址", "姓名", "身份证", "护照", "性别", "店铺链接",
    "平台信息", "服务内容", "服务的国家", "销售量", "预计销售", "说明", "备注", "请提供", "请选择", "填写",
)

# 申请表中的“注册公司信息/联系人信息”区域看起来也像一张表，
# 但它记录的是一个主体的联系人、证件、签字和联系方式，不是“一行一家公司”。
# 这些词只用于拒绝表头候选，不影响正常清单表头如“客户公司名称”。
_XLSX_FORM_HEADER_HINTS = (
    "公司注册信息", "注册公司信息", "客户注册信息", "企业注册信息",
    "联系人信息", "法人信息", "代表人信息", "联系方式",
    "companyregistrationinformation", "registrationcompanyinformation",
    "contactinformation", "legalrepresentativeinformation",
    "registrationcapital", "poasigning", "placeofsignature", "signingtime",
    "服务内容", "服务的国家", "预估销售量", "预计销售量",
    "资料列表", "文件示例", "适用国家/业务", "注意事项",
)


def _xlsx_customer_value(value: str) -> str:
    """过滤 EPR 表单字段标签、联系方式和说明文字，保留客户主体候选。"""
    text = _xlsx_cell_text(value)
    if not text:
        return ""
    normalized = re.sub(r"\s+", "", text).lower()
    if any(hint in normalized for hint in _XLSX_NON_COMPANY_VALUE_HINTS):
        return ""
    # 这些是资料清单里的说明句，不是主体名称；单独列出是为了兼容
    # “请提供”之外的表达（如“若不能提供，请客户确认好注册类别”）。
    if any(hint in normalized for hint in (
        "提供", "注意事项", "注册类别", "确认好", "翻译公司", "盖章",
        "回收公司", "适用国家", "所有国家", "产品图片或说明书", "资料列表",
        "不用提供", "保证有就可以", "否则不接单", "要求北爱公司",
    )):
        return ""
    if "@" in text or re.search(r"https?://|www\\.", text, re.I):
        return ""
    if re.fullmatch(r"[+()\-\s\d]{6,}", text):
        return ""
    # 客户编号/注册号（如 ER0651157、Z6244182）不是公司名称；编号会在
    # 申请表中和公司字段相邻，不能因为它出现在“公司”列就回填为客户。
    if re.fullmatch(r"[A-Za-z]{1,8}[-_]?\d{4,}", text):
        return ""
    # 仅由 2~4 个英文首字母大写单词组成、且没有公司后缀的候选，
    # 在申请表信息区通常是法人/联系人姓名（如 Huiming Wu）。
    # 这类值没有足够主体证据，宁可留空交人工，不冒充公司名称。
    if re.fullmatch(r"[A-Z][a-z]{1,24}(?:\s+[A-Z][a-z]{1,24}){1,3}", text):
        return ""
    if len(text) > 120:
        return ""
    return text


def _xlsx_structured_records(rows: List[List[str]], sheet_name: str) -> List[Dict]:
    """从一个工作表识别“每行一项”的业务清单。

    只接受至少含客户/编号一类实体列的表头，且数据行必须带客户名或
    “编号 + 服务字段”。这样不会把 EPR 模板的国家×业务网格或说明页当成
    真实业务记录；真正的表格行则带 row_number，可回到原附件核对。
    """
    if not rows:
        return []

    header_index = -1
    role_columns: Dict[str, List[int]] = {}
    for index, row in enumerate(rows[:60]):
        row_text = re.sub(r"\s+", "", " ".join(_xlsx_cell_text(cell) for cell in row)).lower()
        if any(hint in row_text for hint in _XLSX_FORM_HEADER_HINTS):
            # 这是 EPR 申请表的单主体信息区，不是客户明细清单表头。
            continue
        candidate: Dict[str, List[int]] = {}
        for column, cell in enumerate(row):
            role = _xlsx_header_role(cell)
            if role:
                candidate.setdefault(role, []).append(column)
        # 必须是“主体 + 业务字段”的清单，或“编号 + 业务字段”的清单。
        # 不能只因工作表中出现“公司名称”就纳入：EPR 申请表的联系人、
        # 法人、身份证等信息区同样会出现这些字样，却不代表多条业务记录。
        has_customer = bool(candidate.get("customer"))
        has_service = bool(
            candidate.get("project") or candidate.get("country")
            or candidate.get("business") or candidate.get("request")
        )
        has_customer_and_service = has_customer and has_service
        has_code_and_service = bool(candidate.get("customer_code")) and has_service
        # 德国 WEEE 附件常单独列出“品牌 | 类别”，这不是一行一个客户，
        # 但必须保留为专项证据供后续与注册工单的“品类明细”比对。
        has_weee_catalog = bool(candidate.get("brand")) and bool(candidate.get("category"))
        # 有些“注册表”是阶段一内部导出格式：项目列不在附件内，
        # 但同时提供公司、代理和客户编号，项目由邮件主题/勾选结果继承。
        # 只有三类实体表头同时出现时才放宽，避免单主体表单误入。
        has_entity_export = (
            has_customer and bool(candidate.get("customer_code"))
            and bool(candidate.get("agent"))
        )
        if has_customer_and_service or has_code_and_service or has_entity_export or has_weee_catalog:
            header_index = index
            role_columns = candidate
            # 同一张内部导出表经常同时保留“公司英文名”和“公司中文名”。
            # 中文法定名称更适合工作台和工单显示；若中文列为空，再回退到英文列。
            def customer_priority(column: int) -> tuple:
                header = _xlsx_cell_text(row[column]).lower() if column < len(row) else ""
                if "中文" in header or "chinese" in header:
                    return (0, column)
                if "英文" in header or "english" in header:
                    return (2, column)
                return (1, column)

            role_columns["customer"] = sorted(
                role_columns.get("customer", []), key=customer_priority
            )
            break
    if header_index < 0:
        return []

    records: List[Dict] = []
    blank_streak = 0
    for index, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        values = [_xlsx_cell_text(cell) for cell in row]
        if not any(values):
            blank_streak += 1
            if records and blank_streak >= 5:
                break
            continue
        blank_streak = 0

        # 重复表头、合计和说明行都不是一条业务记录。
        first_values = " ".join(values).strip()
        if all(_xlsx_header_role(value) for value in values if value):
            continue
        if re.search(r"^(合计|总计|小计|备注|说明)(?:\s|：|:|$)", first_values):
            continue

        def values_for(role: str) -> List[str]:
            return [values[col] for col in role_columns.get(role, []) if col < len(values) and values[col]]

        customers = values_for("customer")
        codes = values_for("customer_code")
        agents = values_for("agent")
        projects = values_for("project")
        countries = values_for("country")
        businesses = values_for("business")
        requests = values_for("request")
        brands = values_for("brand")
        categories = values_for("category")
        # 只有明确的客户主体才进入 customer；EPR 申请表的“法人/注册资本/
        # 签字时间”等标签即使落在同一列，也不能成为客户记录。
        customer = _xlsx_customer_value(customers[0]) if customers else ""
        # 只有序号不构成记录；编号列存在时也必须搭配至少一个业务字段。
        has_entity = bool(customer) or bool(codes and (projects or countries or businesses or requests)) or bool(brands and categories)
        if not has_entity:
            continue

        records.append({
            "sheet_name": sheet_name,
            "row_number": index,
            "cells": values[:24],
            "customer": customer,
            "customer_code": codes[0] if codes else "",
            "agent": agents[0] if agents else "",
            "project": " ".join(projects),
            "country": " ".join(countries),
            "business": " ".join(businesses),
            "request": " ".join(requests),
            "brand": " ".join(brands),
            "category": " ".join(categories),
            "record_type": "weee_catalog" if (not customer and brands and categories) else "entity",
            "raw_text": " | ".join(value for value in values if value),
        })
    return records


def _parse_excel(filepath: str) -> Dict:
    import warnings
    from openpyxl import load_workbook
    text_parts = []
    sheets = []
    structured_records = []
    try:
        # openpyxl 在 load 和 iter_rows 期间都会对外部链接/扩展发 UserWarning;
        # 写到 stderr 时若控制台处于 QuickEdit 选择模式会永久阻塞 worker 线程
        # → 全程压制 (这些警告无业务价值)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = load_workbook(filepath, read_only=True, data_only=True)
            for ws_name in wb.sheetnames:
                ws = wb[ws_name]
                rows_text = []
                rows_values = []
                for row in ws.iter_rows(max_row=MAX_XLSX_SCAN_ROWS, values_only=True):
                    cells = [_xlsx_cell_text(c) for c in row]
                    rows_values.append(cells)
                    line = " ".join(cells)
                    if line.strip():
                        rows_text.append(line)
                sheet_text = "\n".join(rows_text)
                text_parts.append(sheet_text)
                # 工作台直接展示可回溯的表格行；限制列数/单元格长度，避免把
                # 大型附件完整塞进阶段一 Excel 和浏览器响应。
                preview_rows = []
                for row_number, cells in enumerate(rows_values[:120], start=1):
                    if any(cells):
                        preview_rows.append({
                            "row_number": row_number,
                            "cells": [str(cell)[:240] for cell in cells[:24]],
                        })
                sheets.append({
                    "sheet_name": ws_name,
                    "text": sheet_text,
                    "preview_rows": preview_rows,
                })
                structured_records.extend(_xlsx_structured_records(rows_values, ws_name))
            wb.close()
    except Exception as e:
        _get_logger().error(f"解析Excel失败: {e}")
    return {
        "text_content": "\n".join(text_parts),
        "sheets": sheets,
        "structured_records": structured_records,
    }


def _parse_xls(filepath: str) -> str:
    import xlrd
    text_parts = []
    wb = xlrd.open_workbook(filepath)
    for sheet in wb.sheets():
        for row_idx in range(min(sheet.nrows, 200)):
            cells = []
            for col_idx in range(sheet.ncols):
                val = sheet.cell_value(row_idx, col_idx)
                if val:
                    cells.append(str(val).strip())
            line = " ".join(cells)
            if line.strip():
                text_parts.append(line)
    return "\n".join(text_parts)


def _parse_pdf(filepath: str) -> str:
    import pdfplumber
    text_parts = []
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages[:30]:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    return "\n".join(text_parts)


def _parse_docx(filepath: str) -> str:
    from docx import Document
    doc = Document(filepath)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _parse_image(filepath: str) -> str:
    """优先使用 PaddleOCR，失败时使用随包的轻量 Tesseract（中英文）。"""
    # 发布包内置 Tesseract 时直接走轻量后端，避免操作员电脑初始化
    # Paddle/Torch 的 DLL 和模型；开发机仍可显式设置 ECOPV_USE_PADDLE_OCR=1
    # 恢复 PaddleOCR 作为首选。
    if _tesseract_path() and os.environ.get("ECOPV_USE_PADDLE_OCR", "0") != "1":
        return _parse_image_with_tesseract(filepath)
    ocr = _get_ocr()
    if ocr is not None:
        try:
            result = ocr.ocr(filepath)
            text = _extract_ocr_texts(result)
            if text:
                return text
        except Exception as e:
            _get_logger().warning(f"PaddleOCR识别失败 {filepath}: {e}，切换 Tesseract")
    return _parse_image_with_tesseract(filepath)


def _extract_ocr_texts(result) -> str:
    """兼容 PaddleOCR 2.x / 3.x 两种返回格式，抽成纯文本"""
    lines = []
    if not result:
        return ""
    for page in result:
        if not page:
            continue
        if isinstance(page, dict):
            # PaddleOCR 3.x: {'rec_texts': [...], 'rec_scores': [...]}
            for text in page.get("rec_texts") or []:
                if text and text.strip():
                    lines.append(text.strip())
        else:
            # PaddleOCR 2.x: [[box, (text, score)], ...]
            for line_info in page:
                try:
                    text = line_info[1][0]
                except (TypeError, IndexError):
                    continue
                if isinstance(text, str) and text.strip():
                    lines.append(text.strip())
    return "\n".join(lines)


def _parse_csv(filepath: str) -> str:
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _parse_zip(filepath: str) -> Dict:
    text_parts = []
    sheets = []
    pending_images = []
    epr_forms = []
    structured_records = []
    with tempfile.TemporaryDirectory() as tmpdir:
        try:
            with zipfile.ZipFile(filepath, "r") as zf:
                _safe_extract_zip(zf, tmpdir)
        except Exception as e:
            _get_logger().error(f"解压ZIP失败: {e}")
            return {"text_content": "", "sheets": [], "pending_images": [], "epr_forms": []}

        for root, _dirs, files in os.walk(tmpdir):
            for fname in files:
                fpath = os.path.join(root, fname)
                ext = os.path.splitext(fname)[1].lower()
                if ext in (".xlsx", ".xls", ".pdf", ".docx", ".csv", ".txt"):
                    sub = parse_attachment(fpath, fname)
                    if sub["text_content"]:
                        text_parts.append(f"[{fname}] {sub['text_content']}")
                    if sub["sheets"]:
                        sheets.extend(sub["sheets"])
                    if sub.get("epr_forms"):
                        epr_forms.extend(sub["epr_forms"])
                    if sub.get("structured_records"):
                        structured_records.extend(sub["structured_records"])
                elif ext in IMAGE_EXTS:
                    # 压缩包内图片: 持久化后登记待 OCR（懒加载）
                    try:
                        with open(fpath, "rb") as f:
                            img_path = save_image_attachment(f.read(), fname)
                        pending_images.append({
                            "filename": fname, "text_content": "", "sheets": [],
                            "ocr_pending": True, "filepath": img_path,
                        })
                        text_parts.append(f"[{fname}]")
                    except Exception as e:
                        _get_logger().warning(f"压缩包内图片持久化失败 {fname}: {e}")
    return {"text_content": "\n".join(text_parts), "sheets": sheets,
            "pending_images": pending_images, "epr_forms": epr_forms,
            "structured_records": structured_records}


def _parse_rar(filepath: str) -> Dict:
    text_parts = []
    sheets = []
    pending_images = []
    epr_forms = []
    structured_records = []
    with tempfile.TemporaryDirectory() as tmpdir:
        extracted = False

        # RAR 必须有受支持的解压后端（unrar 或 7z）。使用 rarfile 的逐成员
        # 流式解压，先验证所有成员，不能对邮件附件调用 extractall。
        try:
            import rarfile
            tool_path = _configure_rar_tool(rarfile)
            if tool_path:
                _get_logger().info(f"RAR解压后端: {tool_path}")
            with rarfile.RarFile(filepath, "r") as rf:
                _safe_extract_rar(rf, tmpdir)
            extracted = True
        except Exception as e:
            _get_logger().warning(f"RAR安全解压失败: {type(e).__name__}: {e}")

        if not extracted:
            _get_logger().error(
                "RAR解压不可用，已跳过。请安装 7-Zip 或 UnRAR 并确认其可执行文件在 PATH 中: "
                f"{filepath}"
            )
            return {"text_content": "", "sheets": [], "pending_images": [], "epr_forms": [],
                    "structured_records": []}

        for root, _dirs, files in os.walk(tmpdir):
            for fname in files:
                fpath = os.path.join(root, fname)
                ext = os.path.splitext(fname)[1].lower()
                if ext in (".xlsx", ".xls", ".pdf", ".docx", ".csv", ".txt"):
                    sub = parse_attachment(fpath, fname)
                    if sub["text_content"]:
                        text_parts.append(f"[{fname}] {sub['text_content']}")
                    if sub["sheets"]:
                        sheets.extend(sub["sheets"])
                    if sub.get("epr_forms"):
                        epr_forms.extend(sub["epr_forms"])
                    if sub.get("structured_records"):
                        structured_records.extend(sub["structured_records"])
                elif ext in IMAGE_EXTS:
                    # 压缩包内图片: 持久化后登记待 OCR（懒加载）
                    try:
                        with open(fpath, "rb") as f:
                            img_path = save_image_attachment(f.read(), fname)
                        pending_images.append({
                            "filename": fname, "text_content": "", "sheets": [],
                            "ocr_pending": True, "filepath": img_path,
                        })
                        text_parts.append(f"[{fname}]")
                    except Exception as e:
                        _get_logger().warning(f"压缩包内图片持久化失败 {fname}: {e}")
    return {"text_content": "\n".join(text_parts), "sheets": sheets,
            "pending_images": pending_images, "epr_forms": epr_forms,
            "structured_records": structured_records}


def extract_emails_from_text(text: str) -> List[str]:
    """从文本中提取邮箱地址"""
    if not text:
        return []
    pattern = r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}"
    return list(set(re.findall(pattern, text)))


import re
